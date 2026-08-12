"""Unified smart-and-manual Excel worksheet mapping studio."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PySide6.QtCore import QItemSelectionModel, Qt
from PySide6.QtGui import QAction, QIcon
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QComboBox, QDialog, QFileDialog, QFrame,
    QHBoxLayout, QHeaderView, QLabel, QLineEdit, QListWidget, QMainWindow,
    QGridLayout, QMenu, QMessageBox, QProgressBar, QPushButton, QInputDialog, QStackedWidget,
    QSplitter, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget,
)

# The Qt build intentionally excludes Tk/Tcl. Import the Qt compatibility
# bootstrap first so excel_mapper sees the lightweight tkinter placeholders.
from excel_mapping_qt import bundled_asset, install_tkinter_build_stub

from excel_mapper import (
    MODE_SEQUENCE, MappingRule, best_name_match, execute_mapping_plan,
    format_cell_addresses, format_project_path, infer_mapping_mode,
    load_excel_mapping_scheme, parse_cell_addresses, resolve_project_path,
    spreadsheet_automation_provider,
    validate_mapping_plan, workbook_files_in_folder, workbook_sheet_names,
    xls_formula_risk,
)
from excel_sheet_viewer import DualSheetViewer
from smart_template import SheetPair, suggest_sheet_pairs


WORKBOOK_FILTER = "Excel 文件 (*.xls *.xlsx *.xlsm)"


def target_code_for_source_name(name: str) -> str | None:
    """Convert common three-digit source report codes to target report codes."""
    matches = re.findall(r"(?<!\d)(\d{3})(?!\d)", name)
    if not matches:
        matches = re.findall(r"(\d{3})", name)
    if not matches:
        return None
    code = matches[-1]
    return code[:2] if code.endswith("0") else f"{code[:2]}0{code[2]}"


def preferred_target_sheet(source_sheet: str, sheets: list[str]) -> str | None:
    """Prefer the bank report-code convention, then fall back to fuzzy matching."""
    expected = target_code_for_source_name(source_sheet)
    if expected:
        pattern = re.compile(rf"(?<!\d){re.escape(expected)}(?!\d)")
        matches = [sheet for sheet in sheets if pattern.search(sheet)]
        if matches:
            return best_name_match(source_sheet, matches) or matches[0]
    return best_name_match(source_sheet, sheets)


SHEET_PAIR_HEADERS = ("来源工作簿", "来源工作表", "目标工作簿", "目标工作表")
CELL_PAIR_HEADERS = (*SHEET_PAIR_HEADERS[:2], "来源单元格", *SHEET_PAIR_HEADERS[2:], "目标单元格")


def save_studio_mapping_workbook(
    path: Path, mappings: list[WorksheetMapping], base_folder: Path, template: bool = False,
) -> None:
    """Save the new one-row-per-relation format without a redundant 使用 column."""
    book = Workbook()
    sheets = book.active; sheets.title = "工作表对应"
    cells = book.create_sheet("单元格映射")
    sheets.append(SHEET_PAIR_HEADERS); cells.append(CELL_PAIR_HEADERS)
    rows = mappings or ([WorksheetMapping(
        "来源工作簿.xlsx", "来源工作表名称", "目标工作簿.xlsx", "目标工作表名称",
        ["A1", "A2"], ["A1", "A2"],
    )] if template else [])
    for mapping in rows:
        if not all((mapping.source_file, mapping.source_sheet, mapping.target_file, mapping.target_sheet)):
            continue
        base = [
            format_project_path(Path(mapping.source_file), base_folder), mapping.source_sheet,
            format_project_path(Path(mapping.target_file), base_folder), mapping.target_sheet,
        ]
        sheets.append(base)
        if mapping.source_cells or mapping.target_cells:
            target = "同位置" if mapping.source_cells == mapping.target_cells else format_cell_addresses(mapping.target_cells)
            cells.append([base[0], base[1], format_cell_addresses(mapping.source_cells), base[2], base[3], target])
    for sheet in (sheets, cells):
        header_fill = PatternFill("solid", fgColor="EEF2F6")
        edge = Side(style="thin", color="D9DEE5")
        for cell in sheet[1]:
            cell.fill = header_fill; cell.font = Font(bold=True, color="18212B")
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), 2):
            fill = PatternFill("solid", fgColor="F3F6FA" if row_number % 2 == 0 else "FFFFFF")
            for cell in row:
                cell.fill = fill; cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(48, max(16, max(len(str(item.value or "")) for item in column) + 2))
    help_sheet = book.create_sheet("填写说明")
    help_lines = (
        "Excel 工作表与单元格映射关系填写说明",
        "工作表对应：每行填写一组来源工作簿、来源工作表、目标工作簿、目标工作表。",
        "单元格映射：每行填写一组完整关系；目标单元格可写“同位置”。",
        "工作簿可填写文件名、相对路径或绝对路径；相对路径以本文件所在文件夹为基准。",
        "单元格支持 A1、A1,A3、A1-A3、A1-B3；顺序会影响一一对应关系。",
        "程序仍兼容旧版三列、四列、五列以及来源/目标两行一组格式。",
        "底色、字体和边框仅用于查看，不影响导入；可以不设置背景颜色。",
    )
    for line in help_lines: help_sheet.append([line])
    help_sheet.column_dimensions["A"].width = 110
    help_sheet["A1"].font = Font(bold=True, size=14)
    path.parent.mkdir(parents=True, exist_ok=True); book.save(path); book.close()


def load_studio_mapping_workbook(path: Path) -> list[WorksheetMapping]:
    """Load the new format, falling back to every previously supported layout."""
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        has_new = "工作表对应" in book.sheetnames or "单元格映射" in book.sheetnames
        if not has_new:
            _, _, rules = load_excel_mapping_scheme(path)
            return [WorksheetMapping(
                str(rule.source_file), rule.source_sheet, str(rule.target_file), rule.target_sheet,
                list(rule.source_cells), list(rule.target_cells),
            ) for rule in rules]
        base = path.parent.resolve(); result: list[WorksheetMapping] = []
        by_pair: dict[tuple[str, str, str, str], WorksheetMapping] = {}
        if "工作表对应" in book.sheetnames:
            sheet = book["工作表对应"]
            headers = tuple(sheet.cell(1, i).value for i in range(1, 5))
            if headers != SHEET_PAIR_HEADERS:
                raise ValueError("“工作表对应”页签表头应为：" + "、".join(SHEET_PAIR_HEADERS))
            for values in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
                if not any(values): continue
                source_file, source_sheet, target_file, target_sheet = (str(value or "").strip() for value in values)
                item = WorksheetMapping(str(resolve_project_path(source_file, base)), source_sheet,
                                        str(resolve_project_path(target_file, base)), target_sheet)
                key = (item.source_file, item.source_sheet, item.target_file, item.target_sheet)
                by_pair[key] = item; result.append(item)
        if "单元格映射" in book.sheetnames:
            sheet = book["单元格映射"]
            headers = tuple(sheet.cell(1, i).value for i in range(1, 7))
            if headers != CELL_PAIR_HEADERS:
                raise ValueError("“单元格映射”页签表头应为：" + "、".join(CELL_PAIR_HEADERS))
            for values in sheet.iter_rows(min_row=2, max_col=6, values_only=True):
                if not any(values): continue
                sf, ss, sc, tf, ts, tc = (str(value or "").strip() for value in values)
                source_path = str(resolve_project_path(sf, base)); target_path = str(resolve_project_path(tf, base))
                source_cells = parse_cell_addresses(sc)
                target_cells = list(source_cells) if tc in ("同位置", "与读取单元格相同") else parse_cell_addresses(tc)
                key = (source_path, ss, target_path, ts)
                item = by_pair.get(key)
                if item and not item.source_cells:
                    item.source_cells, item.target_cells = source_cells, target_cells
                else:
                    result.append(WorksheetMapping(source_path, ss, target_path, ts, source_cells, target_cells))
        return result
    finally:
        book.close()


@dataclass
class WorksheetMapping:
    source_file: str
    source_sheet: str
    target_file: str
    target_sheet: str
    source_cells: list[str] = field(default_factory=list)
    target_cells: list[str] = field(default_factory=list)
    score: float = 0.0
    enabled: bool = True

    @property
    def status(self) -> str:
        if not self.target_file or not self.target_sheet:
            return "尚未选择目标工作表"
        if not self.source_cells:
            return "尚未选择单元格"
        if len(self.source_cells) != len(self.target_cells):
            return f"数量不一致（{len(self.source_cells)} / {len(self.target_cells)}）"
        return f"已设置 {len(self.source_cells)} 项"


def suggest_studio_mappings(
    source_files: list[Path], target_files: list[Path]
) -> list[WorksheetMapping]:
    """Choose the best target workbook/sheet for each source worksheet."""
    best: dict[tuple[str, str], SheetPair] = {}
    for target in target_files:
        for pair in suggest_sheet_pairs(source_files, target):
            key = (pair.source_file, pair.source_sheet)
            if key not in best or pair.score > best[key].score:
                best[key] = pair
    result: list[WorksheetMapping] = []
    for pair in best.values():
        result.append(
            WorksheetMapping(
                pair.source_file, pair.source_sheet,
                pair.target_file, pair.target_sheet,
                score=pair.score,
            )
        )
    return result


def mapping_rules_from_rows(rows: list[WorksheetMapping]) -> list[MappingRule]:
    rules: list[MappingRule] = []
    for index, row in enumerate(rows, 1):
        if not row.enabled:
            continue
        if not row.source_cells or not row.target_cells:
            raise ValueError(f"第 {index} 行尚未选择需要映射的单元格。")
        if len(row.source_cells) != len(row.target_cells):
            raise ValueError(
                f"第 {index} 行来源选择了 {len(row.source_cells)} 个单元格，"
                f"目标选择了 {len(row.target_cells)} 个单元格。两边数量必须相同。"
            )
        rules.append(
            MappingRule(
                row.source_file, row.source_sheet, list(row.source_cells),
                row.target_file, row.target_sheet, list(row.target_cells),
                infer_mapping_mode(row.source_cells, row.target_cells),
            )
        )
    if not rules:
        raise ValueError("请至少设置一组工作表对应关系。")
    return rules


def mapping_relation_text(row: WorksheetMapping) -> str:
    return (
        f"{Path(row.source_file).name}/{row.source_sheet}/"
        f"{format_cell_addresses(row.source_cells)} → "
        f"{Path(row.target_file).name}/{row.target_sheet}/"
        f"{format_cell_addresses(row.target_cells)}"
    )


def validate_studio_duplicate_targets(rows: list[WorksheetMapping]) -> None:
    seen: dict[tuple[str, str, str], tuple[int, WorksheetMapping, str]] = {}
    for row_number, row in enumerate(rows, 1):
        if not row.enabled:
            continue
        for source_cell, target_cell in zip(row.source_cells, row.target_cells):
            key = (str(Path(row.target_file).resolve()), row.target_sheet, target_cell)
            previous = seen.get(key)
            if previous:
                previous_number, previous_row, previous_source_cell = previous
                raise ValueError(
                    f"第 {previous_number} 组和第 {row_number} 组都写入目标单元格：\n"
                    f"{Path(row.target_file).name}/{row.target_sheet}/{target_cell}\n\n"
                    f"第 {previous_number} 组：{mapping_relation_text(previous_row)}\n"
                    f"其中 {previous_source_cell} → {target_cell}\n\n"
                    f"第 {row_number} 组：{mapping_relation_text(row)}\n"
                    f"其中 {source_cell} → {target_cell}\n\n"
                    "请双击对应行，在双表查看器中取消其中一个重复位置。"
                )
            seen[key] = (row_number, row, source_cell)


def preserve_existing_pair_order(
    mapping: WorksheetMapping,
    selected_sources: list[str],
    selected_targets: list[str],
) -> tuple[list[str], list[str]]:
    """Keep smart/manual pair order, then append newly selected cells in view order."""
    # QItemSelectionModel normally returns unique indexes, but old smart plans
    # can contain duplicate addresses. Normalize both sides before comparing so
    # the numbers displayed in the dialog are exactly the numbers saved.
    selected_sources = list(dict.fromkeys(selected_sources))
    selected_targets = list(dict.fromkeys(selected_targets))
    if len(selected_sources) != len(selected_targets):
        raise ValueError(
            f"来源当前选中 {len(selected_sources)} 个，目标当前选中 "
            f"{len(selected_targets)} 个。两边数量必须相同。"
        )
    source_set, target_set = set(selected_sources), set(selected_targets)
    retained: list[tuple[str, str]] = []
    used_sources: set[str] = set()
    used_targets: set[str] = set()
    for source, target in zip(mapping.source_cells, mapping.target_cells):
        if (
            source in source_set and target in target_set
            and source not in used_sources and target not in used_targets
        ):
            retained.append((source, target))
            used_sources.add(source)
            used_targets.add(target)
    new_sources = [cell for cell in selected_sources if cell not in used_sources]
    new_targets = [cell for cell in selected_targets if cell not in used_targets]
    # Because retained pairs are strictly one-to-one, equal visible totals
    # guarantee equal residual totals.
    if len(new_sources) != len(new_targets):
        raise ValueError("当前选择无法组成一一对应关系，请取消全部后重新选择。")
    pairs = retained + list(zip(new_sources, new_targets))
    return [source for source, _ in pairs], [target for _, target in pairs]


class FileListPanel(QFrame):
    def __init__(self, title: str, parent=None) -> None:
        super().__init__(parent)
        self.setObjectName("filePanel")
        box = QVBoxLayout(self)
        heading = QLabel(title); heading.setObjectName("sectionTitle")
        box.addWidget(heading)
        self.list = QListWidget()
        self.list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        box.addWidget(self.list, 1)
        buttons = QHBoxLayout()
        add_files = QPushButton("添加文件")
        add_files.clicked.connect(self.choose_files)
        buttons.addWidget(add_files)
        add_folder = QPushButton("添加文件夹")
        add_folder.clicked.connect(self.choose_folder)
        buttons.addWidget(add_folder)
        remove = QPushButton("移除选中")
        remove.clicked.connect(self.remove_selected)
        buttons.addWidget(remove)
        buttons.addStretch(1)
        box.addLayout(buttons)

    def paths(self) -> list[Path]:
        return [Path(self.list.item(index).data(Qt.ItemDataRole.UserRole)) for index in range(self.list.count())]

    def selected_paths(self) -> list[Path]:
        return [Path(item.data(Qt.ItemDataRole.UserRole)) for item in self.list.selectedItems()]

    def add_paths(self, paths: list[Path]) -> None:
        existing = {str(path.resolve()) for path in self.paths()}
        for path in paths:
            resolved = str(path.resolve())
            if resolved in existing:
                continue
            self.list.addItem(path.name)
            added = self.list.item(self.list.count() - 1)
            added.setData(Qt.ItemDataRole.UserRole, resolved)
            added.setToolTip(resolved)
            existing.add(resolved)

    def choose_files(self) -> None:
        values, _ = QFileDialog.getOpenFileNames(self, "选择 Excel 文件", "", WORKBOOK_FILTER)
        self.add_paths([Path(value) for value in values])

    def choose_folder(self) -> None:
        value = QFileDialog.getExistingDirectory(self, "选择包含 Excel 文件的文件夹")
        if value:
            self.add_paths(workbook_files_in_folder(Path(value)))

    def remove_selected(self) -> None:
        for item in reversed(self.list.selectedItems()):
            self.list.takeItem(self.list.row(item))

    def take_selected(self) -> list[Path]:
        paths = self.selected_paths()
        self.remove_selected()
        return paths


class SheetMappingDialog(QDialog):
    def __init__(
        self, parent, mapping: WorksheetMapping,
        source_files: list[Path], target_files: list[Path],
    ) -> None:
        super().__init__(parent)
        self.mapping = mapping
        self.source_files = source_files
        self.target_files = target_files
        self.saved = False
        self.loading_pair = False
        self.setWindowTitle("设置工作表的数据对应关系")
        self.resize(1450, 900)
        self.setMinimumSize(1000, 700)
        root = QVBoxLayout(self)
        root.setContentsMargins(16, 14, 16, 14)
        title = QLabel("检查并修改需要映射的单元格")
        title.setObjectName("dialogTitle")
        root.addWidget(title)
        note = QLabel(
            "请在左右两张工作表中选择需要一一对应的单元格；"
            "保存时以当前选中结果为准。"
        )
        note.setObjectName("secondaryText")
        root.addWidget(note)
        selectors = QGridLayout()
        selectors.setHorizontalSpacing(18)
        selectors.setVerticalSpacing(6)
        selectors.addWidget(QLabel("来源工作簿"), 0, 0)
        selectors.addWidget(QLabel("目标工作簿"), 0, 2)
        self.source_book = QComboBox()
        for path in source_files:
            self.source_book.addItem(path.name, str(path.resolve()))
        selectors.addWidget(self.source_book, 1, 0)
        self.source_sheet = QComboBox()
        selectors.addWidget(self.source_sheet, 1, 1)
        self.target_book = QComboBox()
        for path in target_files:
            self.target_book.addItem(path.name, str(path.resolve()))
        selectors.addWidget(self.target_book, 1, 2)
        self.target_sheet = QComboBox()
        selectors.addWidget(self.target_sheet, 1, 3)
        selectors.setColumnStretch(0, 3)
        selectors.setColumnStretch(1, 2)
        selectors.setColumnStretch(2, 3)
        selectors.setColumnStretch(3, 2)
        root.addLayout(selectors)
        self._set_initial_selectors()
        self.viewer = DualSheetViewer()
        self.viewer.load_pair(
            Path(mapping.source_file), mapping.source_sheet,
            Path(mapping.target_file), mapping.target_sheet,
        )
        self.viewer.source.select_addresses(mapping.source_cells)
        self.viewer.target.select_addresses(mapping.target_cells)
        self.viewer.target.add_toolbar_button(
            "与来源单元格同位置", self.use_same_positions
        )
        root.addWidget(self.viewer, 1)

        summary = QFrame(); summary.setObjectName("summaryPanel")
        summary_box = QVBoxLayout(summary)
        self.source_summary = QLabel(); self.source_summary.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        self.target_summary = QLabel(); self.target_summary.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        summary_box.addWidget(self.source_summary)
        summary_box.addWidget(self.target_summary)
        root.addWidget(summary)
        self.viewer.source.selection_changed.connect(lambda _items: self.refresh_summary())
        self.viewer.target.selection_changed.connect(lambda _items: self.refresh_summary())
        self.source_book.currentIndexChanged.connect(self.selector_changed)
        self.target_book.currentIndexChanged.connect(self.selector_changed)
        self.source_sheet.currentIndexChanged.connect(self.selector_changed)
        self.target_sheet.currentIndexChanged.connect(self.selector_changed)
        self.refresh_summary()

        actions = QHBoxLayout()
        actions.addStretch(1)
        cancel = QPushButton("取消")
        cancel.clicked.connect(self.reject)
        actions.addWidget(cancel)
        save = QPushButton("保存这组对应关系")
        save.setObjectName("primaryButton")
        save.clicked.connect(self.save_mapping)
        actions.addWidget(save)
        root.addLayout(actions)

    def refresh_summary(self) -> None:
        source = self.viewer.source.selected_addresses()
        target = self.viewer.target.selected_addresses()
        self.source_summary.setText(
            f"来源（{len(source)}）：{format_cell_addresses(source) or '尚未选择'}"
        )
        self.target_summary.setText(
            f"目标（{len(target)}）：{format_cell_addresses(target) or '尚未选择'}"
        )

    def _set_initial_selectors(self) -> None:
        self.loading_pair = True
        try:
            source_index = self.source_book.findData(str(Path(self.mapping.source_file).resolve()))
            target_index = self.target_book.findData(str(Path(self.mapping.target_file).resolve()))
            self.source_book.setCurrentIndex(max(0, source_index))
            self.target_book.setCurrentIndex(max(0, target_index))
            self._fill_sheet_combo(
                self.source_sheet, Path(str(self.source_book.currentData())),
                self.mapping.source_sheet,
            )
            self._fill_sheet_combo(
                self.target_sheet, Path(str(self.target_book.currentData())),
                self.mapping.target_sheet,
            )
        finally:
            self.loading_pair = False

    @staticmethod
    def _fill_sheet_combo(combo: QComboBox, path: Path, selected: str = "") -> None:
        combo.clear()
        sheets = workbook_sheet_names(path)
        combo.addItems(sheets)
        if selected in sheets:
            combo.setCurrentText(selected)

    def selector_changed(self) -> None:
        if self.loading_pair or not hasattr(self, "viewer"):
            return
        self.loading_pair = True
        try:
            sender = self.sender()
            if sender is self.source_book:
                self._fill_sheet_combo(
                    self.source_sheet, Path(str(self.source_book.currentData()))
                )
            elif sender is self.target_book:
                self._fill_sheet_combo(
                    self.target_sheet, Path(str(self.target_book.currentData()))
                )
            self.mapping.source_file = str(self.source_book.currentData())
            self.mapping.target_file = str(self.target_book.currentData())
            self.mapping.source_sheet = self.source_sheet.currentText()
            self.mapping.target_sheet = self.target_sheet.currentText()
            self.mapping.source_cells = []
            self.mapping.target_cells = []
            self.viewer.load_pair(
                Path(self.mapping.source_file), self.mapping.source_sheet,
                Path(self.mapping.target_file), self.mapping.target_sheet,
            )
            self.viewer.source.select_addresses(self.mapping.source_cells)
            self.viewer.target.select_addresses(self.mapping.target_cells)
            self.refresh_summary()
        except Exception as exc:
            QMessageBox.warning(self, "无法切换工作表", str(exc))
        finally:
            self.loading_pair = False

    def use_same_positions(self) -> None:
        source = self.viewer.source.selected_addresses()
        self.viewer.target.select_addresses(source)
        self.refresh_summary()

    def save_mapping(self) -> None:
        source = self.viewer.source.selected_addresses()
        target = self.viewer.target.selected_addresses()
        if len(source) != len(target):
            QMessageBox.warning(
                self, "两边数量不一致",
                f"来源选择了 {len(source)} 个单元格，目标选择了 {len(target)} 个单元格。\n"
                "请继续选择或取消部分单元格，使两边数量相同。",
            )
            return
        if not source:
            QMessageBox.warning(self, "尚未选择", "请至少选择一组来源和目标单元格。")
            return
        try:
            source, target = preserve_existing_pair_order(
                self.mapping, source, target
            )
        except ValueError as exc:
            QMessageBox.warning(self, "无法保存对应关系", str(exc))
            return
        self.mapping.source_cells = source
        self.mapping.target_cells = target
        self.saved = True
        self.accept()

    def closeEvent(self, event) -> None:
        self.viewer.close()
        super().closeEvent(event)


class LegacyStudioWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Excel 映射工作台")
        self.resize(1280, 820)
        self.setMinimumSize(940, 650)
        self.mappings: list[WorksheetMapping] = []
        self.output_folder: Path | None = None
        self.stack = QStackedWidget()
        self.setCentralWidget(self.stack)
        self.files_page = self.build_files_page()
        self.mapping_page = self.build_mapping_page()
        self.stack.addWidget(self.files_page)
        self.stack.addWidget(self.mapping_page)

    def page_header(self, title: str, subtitle: str, layout: QVBoxLayout) -> None:
        heading = QLabel(title); heading.setObjectName("pageTitle")
        layout.addWidget(heading)
        note = QLabel(subtitle); note.setObjectName("pageSubtitle"); note.setWordWrap(True)
        layout.addWidget(note)

    def build_files_page(self) -> QWidget:
        page = QWidget(); root = QVBoxLayout(page); root.setContentsMargins(24, 20, 24, 22)
        self.page_header("1. 选择来源文件和目标文件", "可以一次添加多个文件或整个文件夹。选错身份时，选中文件后点击中间的转换按钮。", root)
        lists = QHBoxLayout()
        self.sources = FileListPanel("来源文件")
        self.targets = FileListPanel("目标文件")
        lists.addWidget(self.sources, 1)
        middle = QVBoxLayout(); middle.addStretch(1)
        transfer = QPushButton("⇄  转换为另一侧")
        transfer.clicked.connect(self.transfer_files)
        middle.addWidget(transfer)
        middle.addStretch(1)
        lists.addLayout(middle)
        lists.addWidget(self.targets, 1)
        root.addLayout(lists, 1)
        actions = QHBoxLayout(); actions.addStretch(1)
        generate = QPushButton("生成工作表对应列表  →")
        generate.setObjectName("primaryButton")
        generate.clicked.connect(self.generate_mappings)
        actions.addWidget(generate)
        root.addLayout(actions)
        return page

    def build_mapping_page(self) -> QWidget:
        page = QWidget(); root = QVBoxLayout(page); root.setContentsMargins(24, 20, 24, 22)
        top = QHBoxLayout()
        back = QPushButton("← 返回文件选择")
        back.clicked.connect(lambda: self.stack.setCurrentWidget(self.files_page))
        top.addWidget(back); top.addStretch(1)
        root.addLayout(top)
        self.page_header(
            "2. 确认工作表对应关系",
            "程序只建议工作表对应关系，不会自动选择单元格。"
            "双击一行手动设置单元格；右键可修改或删除对应关系。",
            root,
        )
        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels(
            ["使用", "来源工作簿", "来源工作表", "→", "目标工作簿", "目标工作表", "数据对应"]
        )
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.setColumnWidth(0, 54); self.table.setColumnWidth(1, 260)
        self.table.setColumnWidth(2, 150); self.table.setColumnWidth(3, 36)
        self.table.setColumnWidth(4, 260); self.table.setColumnWidth(5, 180)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.cellDoubleClicked.connect(lambda row, _column: self.open_mapping(row))
        self.table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self.mapping_menu)
        root.addWidget(self.table, 1)
        tools = QHBoxLayout()
        add = QPushButton("添加对应")
        add.clicked.connect(self.add_mapping)
        tools.addWidget(add)
        remove = QPushButton("删除对应")
        remove.clicked.connect(self.remove_mappings)
        tools.addWidget(remove)
        refresh = QPushButton("重新生成列表")
        refresh.clicked.connect(self.generate_mappings)
        tools.addWidget(refresh)
        tools.addStretch(1)
        root.addLayout(tools)
        execute_box = QFrame(); execute_box.setObjectName("summaryPanel")
        execute = QHBoxLayout(execute_box)
        self.output_edit = QLineEdit(); self.output_edit.setReadOnly(True)
        self.output_edit.setPlaceholderText("点击选择输出文件夹；不选择时使用目标文件旁的“映射结果”")
        self.output_edit.mousePressEvent = lambda _event: self.choose_output()
        execute.addWidget(QLabel("输出文件夹：")); execute.addWidget(self.output_edit, 1)
        self.progress = QProgressBar(); self.progress.setMaximumWidth(190); self.progress.setValue(0)
        execute.addWidget(self.progress)
        run = QPushButton("开始生成报表副本")
        run.setObjectName("primaryButton"); run.clicked.connect(self.run_mapping)
        execute.addWidget(run)
        root.addWidget(execute_box)
        return page

    def transfer_files(self) -> None:
        source_selected = self.sources.selected_paths()
        target_selected = self.targets.selected_paths()
        if source_selected and target_selected:
            QMessageBox.information(self, "请选择一侧", "请只在来源或目标列表中选择需要转换的文件。")
            return
        if source_selected:
            self.targets.add_paths(self.sources.take_selected())
        elif target_selected:
            self.sources.add_paths(self.targets.take_selected())
        else:
            QMessageBox.information(self, "尚未选择文件", "请先选择需要转换身份的文件。")

    def generate_mappings(self) -> None:
        sources, targets = self.sources.paths(), self.targets.paths()
        if not sources or not targets:
            QMessageBox.warning(self, "文件不完整", "请至少添加一个来源文件和一个目标文件。")
            return
        try:
            QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
            self.mappings = suggest_studio_mappings(sources, targets)
            self.refresh_table()
            self.stack.setCurrentWidget(self.mapping_page)
        except Exception as exc:
            QMessageBox.critical(self, "匹配工作表失败", str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def refresh_table(self) -> None:
        self.table.setRowCount(len(self.mappings))
        for row, mapping in enumerate(self.mappings):
            use = QTableWidgetItem("使用" if mapping.enabled else "忽略")
            use.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            values = (
                use, QTableWidgetItem(Path(mapping.source_file).name),
                QTableWidgetItem(mapping.source_sheet), QTableWidgetItem("→"),
                QTableWidgetItem(Path(mapping.target_file).name if mapping.target_file else "尚未选择"),
                QTableWidgetItem(mapping.target_sheet or "尚未选择"),
                QTableWidgetItem(mapping.status),
            )
            for column, item in enumerate(values):
                item.setToolTip(
                    mapping.source_file if column == 1 else
                    mapping.target_file if column == 4 else item.text()
                )
                self.table.setItem(row, column, item)

    def open_mapping(self, row: int) -> None:
        mapping = self.mappings[row]
        if not mapping.target_file or not mapping.target_sheet:
            QMessageBox.warning(self, "目标工作表不完整", "请先为这一行选择目标工作簿和目标工作表。")
            return
        dialog = SheetMappingDialog(
            self, mapping, self.sources.paths(), self.targets.paths()
        )
        if dialog.exec() and dialog.saved:
            self.refresh_table()

    def edit_mapping_field(self, row: int, column: int) -> None:
        mapping = self.mappings[row]
        if column in (1, 4):
            paths = self.sources.paths() if column == 1 else self.targets.paths()
            current = mapping.source_file if column == 1 else mapping.target_file
            labels = [str(path) for path in paths]
            value, ok = QInputDialog.getItem(
                self, "选择工作簿", "工作簿：", labels,
                max(0, labels.index(current) if current in labels else 0), False,
            )
            if not ok or not value:
                return
            if column == 1:
                mapping.source_file = value
                sheets = workbook_sheet_names(Path(value))
                if mapping.source_sheet not in sheets:
                    mapping.source_sheet = sheets[0]
            else:
                mapping.target_file = value
                sheets = workbook_sheet_names(Path(value))
                if mapping.target_sheet not in sheets:
                    mapping.target_sheet = sheets[0]
        else:
            path = Path(mapping.source_file if column == 2 else mapping.target_file)
            sheets = workbook_sheet_names(path)
            current = mapping.source_sheet if column == 2 else mapping.target_sheet
            value, ok = QInputDialog.getItem(
                self, "选择工作表", "工作表：", sheets,
                max(0, sheets.index(current) if current in sheets else 0), False,
            )
            if not ok or not value:
                return
            if column == 2:
                mapping.source_sheet = value
            else:
                mapping.target_sheet = value
        mapping.source_cells = []
        mapping.target_cells = []
        self.refresh_table()

    @staticmethod
    def refresh_suggestion(mapping: WorksheetMapping) -> None:
        """Changing a worksheet invalidates only its manually selected cells."""
        mapping.source_cells = []
        mapping.target_cells = []

    def mapping_menu(self, position) -> None:
        row = self.table.rowAt(position.y())
        if row < 0:
            return
        menu = QMenu(self)
        edit = menu.addAction("打开双表完整查看器")
        menu.addSeparator()
        edit_source_book = menu.addAction("修改来源工作簿")
        edit_source_sheet = menu.addAction("修改来源工作表")
        edit_target_book = menu.addAction("修改目标工作簿")
        edit_target_sheet = menu.addAction("修改目标工作表")
        menu.addSeparator()
        toggle = menu.addAction("忽略此行" if self.mappings[row].enabled else "使用此行")
        delete = menu.addAction("删除此行")
        chosen = menu.exec(self.table.viewport().mapToGlobal(position))
        if chosen == edit:
            self.open_mapping(row)
        elif chosen == edit_source_book:
            self.edit_mapping_field(row, 1)
        elif chosen == edit_source_sheet:
            self.edit_mapping_field(row, 2)
        elif chosen == edit_target_book:
            self.edit_mapping_field(row, 4)
        elif chosen == edit_target_sheet:
            self.edit_mapping_field(row, 5)
        elif chosen == toggle:
            self.mappings[row].enabled = not self.mappings[row].enabled
            self.refresh_table()
        elif chosen == delete:
            self.mappings.pop(row); self.refresh_table()

    def add_mapping(self) -> None:
        sources, targets = self.sources.paths(), self.targets.paths()
        if not sources or not targets:
            return
        chooser = QDialog(self)
        chooser.setWindowTitle("添加工作表对应关系")
        chooser.resize(620, 260)
        layout = QVBoxLayout(chooser)
        layout.addWidget(QLabel("请选择来源工作簿、来源工作表、目标工作簿和目标工作表。"))
        source_book = QComboBox(); target_book = QComboBox()
        source_sheet = QComboBox(); target_sheet = QComboBox()
        for path in sources: source_book.addItem(path.name, str(path.resolve()))
        for path in targets: target_book.addItem(path.name, str(path.resolve()))
        grid = QHBoxLayout()
        left = QVBoxLayout(); left.addWidget(QLabel("来源工作簿")); left.addWidget(source_book); left.addWidget(QLabel("来源工作表")); left.addWidget(source_sheet)
        right = QVBoxLayout(); right.addWidget(QLabel("目标工作簿")); right.addWidget(target_book); right.addWidget(QLabel("目标工作表")); right.addWidget(target_sheet)
        grid.addLayout(left, 1); grid.addLayout(right, 1); layout.addLayout(grid)
        def fill_source():
            source_sheet.clear(); source_sheet.addItems(workbook_sheet_names(Path(str(source_book.currentData()))))
        def fill_target():
            target_sheet.clear(); target_sheet.addItems(workbook_sheet_names(Path(str(target_book.currentData()))))
        source_book.currentIndexChanged.connect(fill_source)
        target_book.currentIndexChanged.connect(fill_target)
        fill_source(); fill_target()
        actions = QHBoxLayout(); actions.addStretch(1)
        cancel = QPushButton("取消"); cancel.clicked.connect(chooser.reject); actions.addWidget(cancel)
        confirm = QPushButton("添加并设置单元格"); confirm.setObjectName("primaryButton"); confirm.clicked.connect(chooser.accept); actions.addWidget(confirm)
        layout.addLayout(actions)
        if not chooser.exec():
            return
        mapping = WorksheetMapping(
            str(source_book.currentData()), source_sheet.currentText(),
            str(target_book.currentData()), target_sheet.currentText(),
        )
        self.mappings.append(mapping)
        self.refresh_table()
        self.open_mapping(len(self.mappings) - 1)

    def remove_mappings(self) -> None:
        rows = sorted({index.row() for index in self.table.selectedIndexes()}, reverse=True)
        for row in rows:
            self.mappings.pop(row)
        self.refresh_table()

    def choose_output(self) -> None:
        value = QFileDialog.getExistingDirectory(self, "选择输出文件夹", self.output_edit.text())
        if value:
            self.output_folder = Path(value); self.output_edit.setText(value)

    def run_mapping(self) -> None:
        try:
            rules = mapping_rules_from_rows(self.mappings)
            validate_studio_duplicate_targets(self.mappings)
            source_files = list(dict.fromkeys(Path(item.source_file) for item in self.mappings if item.enabled))
            target_files = list(dict.fromkeys(Path(item.target_file) for item in self.mappings if item.enabled))
            expanded, warnings = validate_mapping_plan(source_files, target_files, rules)
            if warnings and QMessageBox.question(
                self, "确认写入副本", f"目标副本中有 {len(warnings)} 个位置已有内容，是否继续？"
            ) != QMessageBox.StandardButton.Yes:
                return
            if spreadsheet_automation_provider() is None:
                source_formulas, target_formulas = xls_formula_risk(expanded, target_files)
                if (source_formulas or target_formulas) and QMessageBox.question(
                    self, "公式处理提示",
                    "未检测到 Excel/WPS。旧版 XLS 中的公式可能按当前计算结果写成数值，是否继续？"
                ) != QMessageBox.StandardButton.Yes:
                    return
            output = self.output_folder or target_files[0].parent / "映射结果"
            output.mkdir(parents=True, exist_ok=True)
            def progress(current, total, _message):
                self.progress.setValue(int(current / max(total, 1) * 100)); QApplication.processEvents()
            outputs = execute_mapping_plan(source_files, target_files, rules, output, progress)
            QMessageBox.information(self, "生成完成", "已生成报表副本：\n" + "\n".join(map(str, outputs)))
        except Exception as exc:
            QMessageBox.critical(self, "生成失败", str(exc))
        finally:
            self.progress.setValue(0)


class StudioWindow(QMainWindow):
    """Single-window worksheet and cell mapping workspace."""

    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Excel 工作表映射工具")
        self.resize(1480, 920)
        self.setMinimumSize(1050, 720)
        self.default_folder: Path | None = None
        self.default_target_file: Path | None = None
        self.output_folder: Path | None = None
        self.known_files: list[Path] = []
        self.mappings: list[WorksheetMapping] = []
        self.current_row = -1
        self.loading_viewer = False
        self.same_position_default = True
        self.summary_window = None

        central = QWidget(); self.setCentralWidget(central)
        root = QVBoxLayout(central); root.setContentsMargins(18, 14, 18, 14); root.setSpacing(10)
        heading_line = QHBoxLayout()
        heading = QLabel("Excel 工作表映射工具"); heading.setObjectName("pageTitle")
        heading_line.addWidget(heading)
        heading_line.addStretch(1)
        heading_line.addWidget(QLabel("默认工作文件夹："))
        self.folder_edit = QLineEdit(); self.folder_edit.setReadOnly(True)
        self.folder_edit.setPlaceholderText("点击选择文件夹")
        self.folder_edit.setMinimumWidth(360)
        self.folder_edit.mousePressEvent = lambda _event: self.choose_default_folder()
        heading_line.addWidget(self.folder_edit, 1)
        heading_line.addWidget(QLabel("默认目标工作簿："))
        self.target_edit = QLineEdit(); self.target_edit.setReadOnly(True)
        self.target_edit.setPlaceholderText("点击选择（可选）")
        self.target_edit.setMinimumWidth(260)
        self.target_edit.mousePressEvent = lambda _event: self.choose_default_target()
        heading_line.addWidget(self.target_edit, 1)
        root.addLayout(heading_line)
        self.build_menu_bar()

        upper = QWidget(); upper_layout = QVBoxLayout(upper)
        upper_layout.setContentsMargins(0, 0, 0, 0); upper_layout.setSpacing(6)
        map_header = QHBoxLayout()
        title = QLabel("工作表对应关系"); title.setObjectName("sectionTitle")
        map_header.addWidget(title); map_header.addStretch(1)
        remove = QPushButton("删除选中"); remove.clicked.connect(self.remove_selected); map_header.addWidget(remove)
        clear = QPushButton("清空列表"); clear.clicked.connect(self.clear_mappings); map_header.addWidget(clear)
        upper_layout.addLayout(map_header)

        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(
            ["来源工作表", "→", "目标工作表", "单元格状态"]
        )
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.setColumnWidth(0, 440); self.table.setColumnWidth(1, 38)
        self.table.setColumnWidth(2, 440)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setMinimumHeight(72)
        self.table.cellClicked.connect(self.mapping_cell_clicked)
        self.table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self.mapping_context_menu)
        self.table.verticalHeader().setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table.verticalHeader().customContextMenuRequested.connect(self.mapping_header_context_menu)
        self.table.verticalHeader().sectionClicked.connect(self.select_mapping_row)
        upper_layout.addWidget(self.table)

        lower = QWidget(); lower_layout = QVBoxLayout(lower)
        lower_layout.setContentsMargins(0, 0, 0, 0); lower_layout.setSpacing(6)
        current_line = QHBoxLayout()
        self.current_label = QLabel("请在上方选择一组完整的工作表对应关系")
        self.current_label.setObjectName("currentRelation")
        current_line.addWidget(self.current_label); current_line.addStretch(1)
        previous = QPushButton("上一组"); previous.clicked.connect(lambda: self.move_current(-1)); current_line.addWidget(previous)
        following = QPushButton("下一组"); following.clicked.connect(lambda: self.move_current(1)); current_line.addWidget(following)
        lower_layout.addLayout(current_line)

        self.viewer = DualSheetViewer(); self.viewer.setEnabled(False)
        self.viewer.source.add_toolbar_button("使用目标单元格位置", self.use_target_positions)
        self.viewer.target.add_toolbar_button("使用来源单元格位置", self.use_same_positions)
        self.viewer.source.selection_changed.connect(lambda _items: self.source_selection_changed())
        self.viewer.target.selection_changed.connect(lambda _items: self.selection_changed())
        self.viewer.controls_widget.hide()
        lower_layout.addWidget(self.viewer, 1)

        self.source_summary = QLabel("来源（0）：尚未选择")
        self.target_summary = QLabel("目标（0）：尚未选择")
        for label in (self.source_summary, self.target_summary):
            label.setObjectName("secondaryText")
            label.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
            lower_layout.addWidget(label)

        self.workspace_splitter = QSplitter(Qt.Orientation.Vertical)
        self.workspace_splitter.addWidget(upper)
        self.workspace_splitter.addWidget(lower)
        self.workspace_splitter.setChildrenCollapsible(False)
        self.workspace_splitter.setStretchFactor(0, 0)
        self.workspace_splitter.setStretchFactor(1, 1)
        self.workspace_splitter.setSizes([110, 650])
        root.addWidget(self.workspace_splitter, 1)

        execute = QFrame(); execute.setObjectName("summaryPanel")
        execute_layout = QGridLayout(execute)
        execute_layout.addWidget(QLabel("输出文件夹："), 0, 0)
        self.output_edit = QLineEdit(); self.output_edit.setReadOnly(True)
        self.output_edit.setPlaceholderText("默认在目标文件旁生成“映射结果”文件夹")
        self.output_edit.mousePressEvent = lambda _event: self.choose_output()
        execute_layout.addWidget(self.output_edit, 0, 1)
        self.status_label = QLabel("尚未开始")
        execute_layout.addWidget(self.status_label, 1, 0)
        self.progress = QProgressBar(); self.progress.setValue(0)
        execute_layout.addWidget(self.progress, 1, 1)
        run = QPushButton("生成报表副本"); run.setObjectName("primaryButton")
        run.clicked.connect(self.run_mapping); execute_layout.addWidget(run, 0, 2, 2, 1)
        root.addWidget(execute)
        self.add_blank_mapping()

    def build_menu_bar(self) -> None:
        file_menu = self.menuBar().addMenu("文件")
        choose_folder = file_menu.addAction("选择默认工作文件夹…")
        choose_folder.triggered.connect(self.choose_default_folder)
        file_menu.addSeparator()
        import_action = file_menu.addAction("导入映射关系表…")
        import_action.triggered.connect(self.import_mapping_workbook)
        export_action = file_menu.addAction("导出映射关系表…")
        export_action.triggered.connect(self.export_mapping_workbook)
        sample_action = file_menu.addAction("导出空白填写样例…")
        sample_action.triggered.connect(self.export_mapping_sample)
        file_menu.addSeparator()
        quit_action = file_menu.addAction("退出")
        quit_action.triggered.connect(self.close)

        edit_menu = self.menuBar().addMenu("编辑")
        copy_action = edit_menu.addAction("复制选中关系")
        copy_action.setShortcut("Ctrl+C"); copy_action.triggered.connect(self.copy_selected_mappings)
        delete_action = edit_menu.addAction("删除选中关系")
        delete_action.setShortcut("Delete"); delete_action.triggered.connect(self.remove_selected)
        edit_menu.addSeparator()
        clear_cells_action = edit_menu.addAction("清空选中关系的单元格")
        clear_cells_action.triggered.connect(self.clear_selected_cells)

        view_menu = self.menuBar().addMenu("视图")
        self.sync_scroll_action = view_menu.addAction("同步滚动")
        self.sync_scroll_action.setCheckable(True); self.sync_scroll_action.setChecked(True)
        self.sync_scroll_action.toggled.connect(self.viewer_link_scroll_changed)
        self.sync_zoom_action = view_menu.addAction("同步缩放")
        self.sync_zoom_action.setCheckable(True); self.sync_zoom_action.setChecked(True)
        self.sync_zoom_action.toggled.connect(self.viewer_link_zoom_changed)
        fit_action = view_menu.addAction("适应窗口")
        fit_action.triggered.connect(lambda: self.viewer.fit_both() if hasattr(self, "viewer") else None)

        settings_menu = self.menuBar().addMenu("设置")
        self.same_position_action = settings_menu.addAction("目标默认使用来源单元格位置")
        self.same_position_action.setCheckable(True); self.same_position_action.setChecked(True)
        self.same_position_action.toggled.connect(lambda checked: setattr(self, "same_position_default", checked))
        self.click_order_action = settings_menu.addAction("根据手动选中顺序一一对应")
        self.click_order_action.setCheckable(True); self.click_order_action.setChecked(False)
        self.click_order_action.toggled.connect(self.set_click_order_mode)

        # Independent utilities live at the far right of the menu bar.  They
        # open their own windows so both workflows can remain in use at once.
        tools_menu = self.menuBar().addMenu("小工具")
        summary_action = tools_menu.addAction("单元格汇总…")
        summary_action.triggered.connect(self.open_summary_window)

    def set_click_order_mode(self, checked: bool) -> None:
        if hasattr(self, "viewer"):
            self.viewer.source.set_click_order(checked)
            self.viewer.target.set_click_order(checked)
            self.save_current_selection()
            self.refresh_selection_summary()

    def open_summary_window(self) -> None:
        from excel_summary_qt import ExcelSummaryWindow
        if self.summary_window is None:
            self.summary_window = ExcelSummaryWindow()
            self.summary_window.setStyleSheet(self.styleSheet())
        self.summary_window.show(); self.summary_window.raise_(); self.summary_window.activateWindow()

    def viewer_link_scroll_changed(self, checked: bool) -> None:
        if hasattr(self, "viewer"): self.viewer.link_scroll.setChecked(checked)

    def viewer_link_zoom_changed(self, checked: bool) -> None:
        if hasattr(self, "viewer"): self.viewer.link_zoom.setChecked(checked)

    def change_zoom(self, delta: int) -> None:
        self.set_zoom(self.viewer.source.zoom_slider.value() + delta)

    def set_zoom(self, value: int) -> None:
        value = max(40, min(180, value))
        self.viewer.set_zoom(value)

    def fit_viewer(self) -> None:
        self.viewer.fit_both()

    def all_files(self) -> list[Path]:
        values = list(self.known_files)
        for mapping in self.mappings:
            if mapping.source_file:
                values.append(Path(mapping.source_file))
            if mapping.target_file:
                values.append(Path(mapping.target_file))
        return list(dict.fromkeys(path.resolve() for path in values))

    def choose_default_folder(self) -> None:
        value = QFileDialog.getExistingDirectory(
            self, "选择默认工作文件夹",
            str(self.default_folder or Path.home()),
        )
        if not value:
            return
        self.default_folder = Path(value)
        self.folder_edit.setText(str(self.default_folder))
        self.folder_edit.setToolTip(str(self.default_folder))
        self.known_files = workbook_files_in_folder(self.default_folder)
        if not self.output_folder:
            self.output_edit.setPlaceholderText(str(self.default_folder / "映射结果"))

    def choose_default_target(self) -> None:
        value, _ = QFileDialog.getOpenFileName(
            self, "选择默认目标工作簿", str(self.default_folder or Path.home()), WORKBOOK_FILTER
        )
        if not value:
            return
        self.default_target_file = Path(value).resolve()
        self.target_edit.setText(self.default_target_file.name)
        self.target_edit.setToolTip(str(self.default_target_file))
        if self.default_target_file not in self.known_files:
            self.known_files.append(self.default_target_file)

    def import_mapping_workbook(self) -> None:
        value, _ = QFileDialog.getOpenFileName(
            self, "导入映射关系表", str(self.default_folder or Path.cwd()),
            "Excel 映射关系表 (*.xlsx *.xlsm)",
        )
        if not value:
            return
        try:
            imported = load_studio_mapping_workbook(Path(value))
        except Exception as exc:
            QMessageBox.critical(self, "导入失败", str(exc)); return
        box = QMessageBox(self); box.setWindowTitle("选择导入方式")
        box.setText(f"识别到 {len(imported)} 组工作表对应关系，请选择导入方式。")
        replace = box.addButton("清空并覆盖", QMessageBox.ButtonRole.AcceptRole)
        append = box.addButton("追加到当前方案", QMessageBox.ButtonRole.ActionRole)
        box.addButton("取消", QMessageBox.ButtonRole.RejectRole); box.exec()
        self.save_current_selection()
        if box.clickedButton() == replace:
            self.mappings = imported
        elif box.clickedButton() == append:
            insert_at = len(self.mappings) - 1 if self.mappings and self.mapping_is_blank(self.mappings[-1]) else len(self.mappings)
            self.mappings[insert_at:insert_at] = imported
        else:
            return
        self.known_files.extend(
            Path(path) for mapping in imported for path in (mapping.source_file, mapping.target_file) if path
        )
        self.current_row = -1; self.ensure_blank_draft(); self.refresh_table()
        if imported:
            self.load_mapping(0 if box.clickedButton() == replace else max(0, len(self.mappings) - len(imported) - 1))
        self.status_label.setText(f"已导入 {len(imported)} 组映射关系")

    def export_mapping_workbook(self) -> None:
        self.save_current_selection()
        mappings = [item for item in self.mappings if not self.mapping_is_blank(item)]
        if not mappings:
            QMessageBox.information(self, "没有可导出的内容", "请先设置工作表对应关系。")
            return
        value, _ = QFileDialog.getSaveFileName(
            self, "导出映射关系表", str((self.default_folder or Path.cwd()) / "映射关系.xlsx"),
            "Excel 映射关系表 (*.xlsx)",
        )
        if not value: return
        path = Path(value); path = path if path.suffix.lower() == ".xlsx" else path.with_suffix(".xlsx")
        save_studio_mapping_workbook(path, mappings, path.parent.resolve())
        self.status_label.setText(f"已导出映射关系表：{path.name}")

    def export_mapping_sample(self) -> None:
        value, _ = QFileDialog.getSaveFileName(
            self, "导出空白填写样例", str((self.default_folder or Path.cwd()) / "映射关系填写样例.xlsx"),
            "Excel 映射关系表 (*.xlsx)",
        )
        if not value: return
        path = Path(value); path = path if path.suffix.lower() == ".xlsx" else path.with_suffix(".xlsx")
        save_studio_mapping_workbook(path, [], path.parent.resolve(), template=True)
        self.status_label.setText(f"已导出填写样例：{path.name}")

    def refresh_table(self) -> None:
        self.table.blockSignals(True)
        self.table.setRowCount(len(self.mappings))
        for row, mapping in enumerate(self.mappings):
            blank = self.mapping_is_blank(mapping)
            texts = (
                (
                    f"{Path(mapping.source_file).name} / {mapping.source_sheet}"
                    if mapping.source_file and mapping.source_sheet else "点击选择来源工作表"
                ),
                "→",
                (
                    f"{Path(mapping.target_file).name} / {mapping.target_sheet}"
                    if mapping.target_file and mapping.target_sheet else "点击选择目标工作表"
                ),
                mapping.status,
            )
            for column, value in enumerate(texts):
                item = QTableWidgetItem(value)
                if column == 1: item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if column == 0: item.setToolTip(mapping.source_file)
                if column == 2: item.setToolTip(mapping.target_file)
                self.table.setItem(row, column, item)
        self.table.blockSignals(False)

    @staticmethod
    def mapping_is_blank(mapping: WorksheetMapping) -> bool:
        return not any((mapping.source_file, mapping.source_sheet, mapping.target_file, mapping.target_sheet, mapping.source_cells, mapping.target_cells))

    def ensure_blank_draft(self) -> None:
        if not self.mappings or not self.mapping_is_blank(self.mappings[-1]):
            self.mappings.append(WorksheetMapping("", "", "", ""))

    def add_blank_mapping(self) -> None:
        self.save_current_selection()
        self.mappings.append(WorksheetMapping("", "", "", ""))
        self.refresh_table()
        self.table.selectRow(len(self.mappings) - 1)

    def clear_mappings(self) -> None:
        if self.mappings and QMessageBox.question(
            self, "清空确认", "确定清空所有工作表和单元格对应关系吗？"
        ) != QMessageBox.StandardButton.Yes:
            return
        self.mappings.clear(); self.current_row = -1; self.refresh_table()
        self.ensure_blank_draft(); self.refresh_table()
        self.clear_viewer_state("请添加或选择一组工作表对应关系")

    def remove_selected(self) -> None:
        rows = sorted({index.row() for index in self.table.selectedIndexes()}, reverse=True)
        if not rows: return
        self.save_current_selection()
        for row in rows: self.mappings.pop(row)
        self.current_row = -1; self.refresh_table()
        self.ensure_blank_draft(); self.refresh_table()
        if self.mappings and not self.mapping_is_blank(self.mappings[0]): self.load_mapping(min(rows[-1], len(self.mappings) - 1))
        else: self.clear_viewer_state()

    def selected_rows(self) -> list[int]:
        return sorted({index.row() for index in self.table.selectedIndexes()})

    def copy_selected_mappings(self) -> None:
        rows = self.selected_rows()
        if not rows:
            return
        insert_at = len(self.mappings) - 1 if self.mappings and self.mapping_is_blank(self.mappings[-1]) else len(self.mappings)
        copies = []
        for row in rows:
            item = self.mappings[row]
            if self.mapping_is_blank(item):
                continue
            copies.append(WorksheetMapping(
                item.source_file, item.source_sheet, item.target_file, item.target_sheet,
                list(item.source_cells), list(item.target_cells), item.score, item.enabled,
            ))
        self.mappings[insert_at:insert_at] = copies
        self.ensure_blank_draft(); self.refresh_table()

    def clear_selected_cells(self) -> None:
        for row in self.selected_rows():
            self.mappings[row].source_cells = []
            self.mappings[row].target_cells = []
        if self.current_row in self.selected_rows() and self.viewer.isEnabled():
            self.viewer.source.clear_selection(); self.viewer.target.clear_selection()
        self.refresh_table()

    def mapping_context_menu(self, position) -> None:
        row = self.table.rowAt(position.y())
        if row < 0:
            return
        self.select_row_for_context_menu(
            row,
            preserve=bool(QApplication.keyboardModifiers() & Qt.KeyboardModifier.ControlModifier),
        )
        menu = QMenu(self)
        above = menu.addAction("在上方插入一组")
        below = menu.addAction("在下方插入一组")
        copy_action = menu.addAction("复制到列表末尾")
        menu.addSeparator()
        clear_cells = menu.addAction("清空单元格选择")
        menu.addSeparator(); delete = menu.addAction("删除选中关系")
        chosen = menu.exec(self.table.viewport().mapToGlobal(position))
        if chosen == above:
            self.mappings.insert(row, WorksheetMapping("", "", "", "")); self.refresh_table()
        elif chosen == below:
            self.mappings.insert(row + 1, WorksheetMapping("", "", "", "")); self.refresh_table()
        elif chosen == copy_action:
            self.copy_selected_mappings()
        elif chosen == clear_cells:
            self.clear_selected_cells()
        elif chosen == delete:
            self.remove_selected()

    def select_mapping_row(self, row: int) -> None:
        if 0 <= row < len(self.mappings):
            self.table.selectRow(row)
            self.load_mapping(row)

    def select_row_for_context_menu(self, row: int, preserve: bool = False) -> None:
        """Make a right-click anywhere in a mapping row select that whole row."""
        if row < 0 or row >= self.table.rowCount():
            return
        model = self.table.model()
        selection = self.table.selectionModel()
        flags = QItemSelectionModel.SelectionFlag.Rows
        flags |= (
            QItemSelectionModel.SelectionFlag.Select
            if preserve
            else QItemSelectionModel.SelectionFlag.ClearAndSelect
        )
        selection.select(model.index(row, 0), flags)
        self.table.setCurrentCell(row, 0)

    def mapping_header_context_menu(self, position) -> None:
        row = self.table.verticalHeader().logicalIndexAt(position)
        if row < 0:
            return
        self.select_row_for_context_menu(
            row,
            preserve=bool(QApplication.keyboardModifiers() & Qt.KeyboardModifier.ControlModifier),
        )
        viewport_position = self.table.viewport().mapFromGlobal(
            self.table.verticalHeader().mapToGlobal(position)
        )
        self.mapping_context_menu(viewport_position)

    def mapping_cell_clicked(self, row: int, column: int) -> None:
        if column in (0, 2):
            self.choose_workbook_sheet(row, column == 0)
        else:
            self.load_mapping(row)

    def choose_workbook_sheet(self, row: int, source: bool) -> None:
        """Choose workbook and worksheet in one hierarchical menu."""
        menu = QMenu(self)
        paths = self.all_files()
        if not source and self.default_target_file in paths:
            paths.remove(self.default_target_file)
            paths.insert(0, self.default_target_file)
        if self.default_folder:
            default_paths = [path for path in paths if path.parent == self.default_folder.resolve()]
            if default_paths:
                heading = menu.addAction("默认文件夹"); heading.setEnabled(False)
                for path in default_paths:
                    book_menu = menu.addMenu(f"    {path.name}")
                    book_menu.setToolTipsVisible(True); book_menu.setToolTip(str(path))
                    for sheet in workbook_sheet_names(path):
                        action = book_menu.addAction(sheet); action.setData((str(path), sheet))
                menu.addSeparator()
        outside = [path for path in paths if not self.default_folder or path.parent != self.default_folder.resolve()]
        if outside:
            heading = menu.addAction("本次方案使用的其他文件"); heading.setEnabled(False)
            for path in outside:
                book_menu = menu.addMenu(f"    {path.name}")
                for sheet in workbook_sheet_names(path):
                    action = book_menu.addAction(sheet); action.setData((str(path), sheet))
            menu.addSeparator()
        browse = menu.addAction("选择其他文件夹中的工作表…")
        clear = menu.addAction("清空选择")
        column = 0 if source else 2
        chosen = menu.exec(self.table.viewport().mapToGlobal(self.table.visualItemRect(self.table.item(row, column)).bottomLeft()))
        if not chosen: return
        if chosen == browse:
            value, _ = QFileDialog.getOpenFileName(self, "选择 Excel 工作簿", str(self.default_folder or Path.home()), WORKBOOK_FILTER)
            if not value: return
            path = Path(value).resolve(); self.known_files.append(path)
            sheets = workbook_sheet_names(path)
            sheet, accepted = QInputDialog.getItem(self, "选择工作表", path.name, sheets, 0, False)
            if not accepted or not sheet: return
        elif chosen == clear:
            path, sheet = None, ""
        else:
            data = chosen.data()
            if not isinstance(data, tuple): return
            path, sheet = Path(data[0]), str(data[1])
        self.save_current_selection()
        mapping = self.mappings[row]
        if source:
            mapping.source_file = str(path or ""); mapping.source_sheet = sheet
        else:
            mapping.target_file = str(path or ""); mapping.target_sheet = sheet
        mapping.source_cells = []; mapping.target_cells = []
        if source and mapping.source_sheet and self.default_target_file:
            target_sheets = workbook_sheet_names(self.default_target_file)
            suggested = preferred_target_sheet(mapping.source_sheet, target_sheets)
            if suggested:
                mapping.target_file = str(self.default_target_file); mapping.target_sheet = suggested
        self.refresh_table(); self.load_mapping(row)
        self.ensure_blank_draft(); self.refresh_table()

    # Kept as compatibility shims for older tests/extensions.
    def choose_workbook(self, row: int, source: bool) -> None:
        self.choose_workbook_sheet(row, source)

    def choose_sheet(self, row: int, source: bool) -> None:
        self.choose_workbook_sheet(row, source)

    def load_mapping(self, row: int) -> None:
        if row < 0 or row >= len(self.mappings): return
        if row != self.current_row: self.save_current_selection()
        mapping = self.mappings[row]
        self.current_row = row; self.table.selectRow(row)
        if not all((mapping.source_file, mapping.source_sheet, mapping.target_file, mapping.target_sheet)):
            self.clear_viewer_state(f"第 {row + 1} / {len(self.mappings)} 组   请先选择完整的工作簿和工作表")
            return
        try:
            self.loading_viewer = True; self.viewer.setEnabled(True)
            self.viewer.load_pair(Path(mapping.source_file), mapping.source_sheet, Path(mapping.target_file), mapping.target_sheet)
            self.viewer.source.select_addresses(mapping.source_cells)
            self.viewer.target.select_addresses(mapping.target_cells)
            self.current_label.setText(
                f"第 {row + 1} / {len(self.mappings)} 组   "
                f"{Path(mapping.source_file).name} / {mapping.source_sheet}  →  "
                f"{Path(mapping.target_file).name} / {mapping.target_sheet}"
            )
            self.refresh_selection_summary()
        except Exception as exc:
            self.clear_viewer_state(); QMessageBox.warning(self, "无法读取工作表", str(exc))
        finally:
            self.loading_viewer = False

    def save_current_selection(self) -> None:
        if self.loading_viewer or self.current_row < 0 or self.current_row >= len(self.mappings): return
        mapping = self.mappings[self.current_row]
        if not self.viewer.isEnabled(): return
        sources = self.viewer.source.selected_addresses(); targets = self.viewer.target.selected_addresses()
        try:
            if len(sources) == len(targets):
                sources, targets = preserve_existing_pair_order(mapping, sources, targets)
        except ValueError:
            pass
        mapping.source_cells = sources; mapping.target_cells = targets

    def selection_changed(self) -> None:
        if self.loading_viewer: return
        self.save_current_selection(); self.refresh_selection_summary(); self.refresh_current_status()

    def refresh_current_status(self) -> None:
        """Update only the changing status cell; rebuilding the table causes visible jumps."""
        if self.current_row < 0 or self.current_row >= len(self.mappings):
            return
        item = self.table.item(self.current_row, 3)
        if item:
            item.setText(self.mappings[self.current_row].status)

    def source_selection_changed(self) -> None:
        if self.loading_viewer:
            return
        if self.same_position_default:
            self.loading_viewer = True
            try:
                self.viewer.target.select_addresses(
                    self.viewer.source.selected_addresses()
                )
            finally:
                self.loading_viewer = False
        self.selection_changed()

    def refresh_selection_summary(self) -> None:
        sources = self.viewer.source.selected_addresses(); targets = self.viewer.target.selected_addresses()
        self.source_summary.setText(
            f"来源（{len(sources)}）：{format_cell_addresses(sources) or '尚未选择'}"
        )
        self.target_summary.setText(
            f"目标（{len(targets)}）：{format_cell_addresses(targets) or '尚未选择'}"
        )

    def use_same_positions(self) -> None:
        self.viewer.target.select_addresses(self.viewer.source.selected_addresses())
        self.selection_changed()

    def use_target_positions(self) -> None:
        self.viewer.source.select_addresses(self.viewer.target.selected_addresses())
        self.selection_changed()

    def clear_viewer_state(self, message: str = "尚未选择工作表对应关系") -> None:
        self.loading_viewer = True
        try:
            self.viewer.clear_pair(); self.viewer.setEnabled(False)
            self.source_summary.setText("来源（0）：尚未选择")
            self.target_summary.setText("目标（0）：尚未选择")
            self.current_label.setText(message)
        finally:
            self.loading_viewer = False

    def move_current(self, delta: int) -> None:
        if not self.mappings: return
        self.load_mapping(max(0, min(len(self.mappings) - 1, self.current_row + delta)))

    def choose_output(self) -> None:
        value = QFileDialog.getExistingDirectory(self, "选择输出文件夹", str(self.output_folder or self.default_folder or Path.home()))
        if value: self.output_folder = Path(value); self.output_edit.setText(value)

    def run_mapping(self) -> None:
        self.save_current_selection(); self.progress.setValue(0)
        try:
            active_mappings = [item for item in self.mappings if not self.mapping_is_blank(item)]
            rules = mapping_rules_from_rows(active_mappings)
            validate_studio_duplicate_targets(active_mappings)
            source_files = list(dict.fromkeys(Path(item.source_file) for item in active_mappings if item.enabled))
            target_files = list(dict.fromkeys(Path(item.target_file) for item in active_mappings if item.enabled))
            expanded, warnings = validate_mapping_plan(source_files, target_files, rules)
            if warnings and QMessageBox.question(self, "确认写入副本", f"目标副本中有 {len(warnings)} 个位置已有内容，是否继续？") != QMessageBox.StandardButton.Yes:
                return
            if spreadsheet_automation_provider() is None:
                source_formulas, target_formulas = xls_formula_risk(expanded, target_files)
                if (source_formulas or target_formulas) and QMessageBox.question(self, "公式处理提示", "未检测到 Excel/WPS。旧版 XLS 中的公式可能按当前计算结果写成数值，是否继续？") != QMessageBox.StandardButton.Yes:
                    return
            output = self.output_folder or (self.default_folder / "映射结果" if self.default_folder else target_files[0].parent / "映射结果")
            output.mkdir(parents=True, exist_ok=True)
            def progress(current, total, message):
                self.status_label.setText(message); self.progress.setValue(int(current / max(total, 1) * 100)); QApplication.processEvents()
            outputs = execute_mapping_plan(source_files, target_files, rules, output, progress)
            self.status_label.setText("已完成")
            QMessageBox.information(self, "生成完成", "已生成报表副本：\n" + "\n".join(map(str, outputs)))
            self.table.clearSelection(); self.current_row = -1
            self.clear_viewer_state("映射已完成；请选择工作表对应关系继续查看或修改")
        except Exception as exc:
            self.status_label.setText("生成失败"); QMessageBox.critical(self, "生成失败", str(exc))
        finally:
            self.progress.setValue(0)

    def closeEvent(self, event) -> None:
        self.viewer.close(); super().closeEvent(event)


def application_style() -> str:
    return """
    QMainWindow, QWidget { background: #F5F7FA; color: #1F2937; font-family: "Microsoft YaHei UI"; font-size: 14px; }
    QLabel#pageTitle { font-size: 21px; font-weight: 600; color: #172B4D; }
    QLabel#pageSubtitle, QLabel#secondaryText, QLabel#currentRelation { color: #667085; }
    QLabel#currentRelation { font-size: 14px; font-weight: 400; }
    QLabel#sectionTitle, QLabel#dialogTitle, QLabel#sheetPaneTitle { font-size: 15px; font-weight: 600; color: #172B4D; }
    QFrame#filePanel, QFrame#summaryPanel, QFrame#sheetPane { background: white; border: 1px solid #DDE3EC; border-radius: 9px; }
    QLineEdit, QComboBox, QListWidget, QTableWidget, QTableView { background: white; border: 1px solid #CBD5E1; border-radius: 5px; selection-background-color: #DCEAFF; selection-color: #172B4D; }
    QPushButton { background: white; border: 1px solid #B8C2D1; border-radius: 5px; padding: 5px 10px; }
    QPushButton#compactButton { padding: 3px 8px; min-height: 20px; }
    QPushButton#squareButton { padding: 4px; min-width: 28px; max-width: 32px; min-height: 26px; }
    QToolButton { background: white; border: 1px solid #B8C2D1; border-radius: 5px; padding: 5px 10px; }
    QPushButton:hover { background: #F0F5FF; border-color: #7EA6E0; }
    QPushButton:checked { background: #EAF2FF; border-color: #2563EB; color: #175CD3; }
    QPushButton#primaryButton { background: #2563EB; color: white; border-color: #2563EB; font-weight: 600; }
    QPushButton#primaryButton:hover { background: #1D4ED8; }
    QProgressBar { background: #E5EAF1; border: none; border-radius: 4px; text-align: center; }
    QProgressBar::chunk { background: #2563EB; }
    """


def main() -> None:
    install_tkinter_build_stub()
    app = QApplication(sys.argv); app.setStyle("Fusion"); app.setStyleSheet(application_style())
    icon = bundled_asset("assets/excel-mapper.ico")
    if icon.exists(): app.setWindowIcon(QIcon(str(icon)))
    window = StudioWindow(); window.show(); sys.exit(app.exec())


if __name__ == "__main__":
    main()
