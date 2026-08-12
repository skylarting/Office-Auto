import os
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication

from excel_mapping_smart import (
    SmartMappingWindow, build_plan, candidate_matches, find_plan_files,
    rebind_plan,
)
from excel_sheet_viewer import SheetViewPane, WorksheetModel
from smart_template import execute_smart_plan, load_plan, save_plan


def make_book(path: Path, sheet_name: str, shifted: bool = False) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    ro, co = (2, 1) if shifted else (0, 0)
    sheet.cell(1 + ro, 2 + co, "余额")
    sheet.cell(2 + ro, 2 + co, "人民币")
    sheet.cell(2 + ro, 3 + co, "外币")
    sheet.cell(3 + ro, 1 + co, "Ⅰ.资产")
    sheet.cell(4 + ro, 1 + co, "1.现金")
    sheet.cell(4 + ro, 2 + co, 10)
    sheet.cell(4 + ro, 3 + co, 20)
    sheet.cell(5 + ro, 1 + co, "2.存款")
    sheet.cell(5 + ro, 2 + co, 30)
    sheet.cell(5 + ro, 3 + co, "=B4+C4")
    book.save(path)


class SmartWizardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_window_starts_on_beginner_home(self):
        window = SmartMappingWindow()
        self.assertIs(window.stack.currentWidget(), window.home)
        self.assertIn("智能报表助手", window.windowTitle())
        window.close()

    def test_pending_cells_are_grouped_by_sheet_pair(self):
        window = SmartMappingWindow()
        from smart_template import SmartMatch, SmartTemplatePlan
        window.plan = SmartTemplatePlan(matches=[
            SmartMatch("a.xlsx", "数据", "A1", (), (), "b.xlsx", "报表", "B1", (), (), .7, "待确认"),
            SmartMatch("a.xlsx", "数据", "A2", (), (), "b.xlsx", "报表", "B2", (), (), .7, "待确认"),
            SmartMatch("c.xlsx", "其他", "A1", (), (), "b.xlsx", "其他", "B1", (), (), .7, "待确认"),
        ])
        groups = window._pending_groups()
        self.assertEqual(len(groups), 2)
        self.assertEqual(sorted(map(len, groups)), [1, 2])
        window.close()

    def test_sheet_viewer_loads_complete_used_area_and_toggles_traits(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "viewer.xlsx"
            make_book(path, "数据")
            pane = SheetViewPane("来源数据")
            pane.load_sheet(path, "数据")
            self.assertEqual(pane.model.rows, 5)
            self.assertEqual(pane.model.columns, 3)
            pane.select_trait("formula", True)
            self.assertEqual(pane.selected_addresses(), ["C5"])
            pane.select_trait("formula", False)
            self.assertEqual(pane.selected_addresses(), [])
            pane.select_trait("number", True)
            self.assertIn("B4", pane.selected_addresses())
            self.assertNotIn("C5", pane.selected_addresses())
            pane.close()

    def test_trait_menu_state_follows_actual_cell_selection(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "viewer.xlsx"
            make_book(path, "数据")
            pane = SheetViewPane("来源数据"); pane.load_sheet(path, "数据")
            self.assertEqual(pane.trait_state("number"), "none")
            pane.select_addresses(["B4"], clear=True)
            self.assertEqual(pane.trait_state("number"), "partial")
            pane.toggle_trait("number")
            self.assertEqual(pane.trait_state("number"), "all")
            pane.toggle_trait("number")
            self.assertEqual(pane.trait_state("number"), "none")
            pane.close()

    def test_review_group_includes_auto_and_pending_items_for_same_sheet(self):
        window = SmartMappingWindow()
        from smart_template import SmartMatch, SmartTemplatePlan
        common = ("a.xlsx", "数据", "b.xlsx", "报表")
        window.plan = SmartTemplatePlan(matches=[
            SmartMatch(common[0], common[1], "A1", (), (), common[2], common[3], "B1", (), (), 1, "自动匹配"),
            SmartMatch(common[0], common[1], "A2", (), (), common[2], common[3], "B2", (), (), .7, "待确认"),
        ])
        groups = window._pending_groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual(len(groups[0]), 2)
        window.close()

    def test_learning_plan_can_save_load_and_execute_copy(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.xlsx"
            target = root / "target.xlsx"
            make_book(source, "数据")
            make_book(target, "报表", shifted=True)
            plan = build_plan([source], target)
            self.assertTrue(plan.pairs)
            self.assertGreater(candidate_matches(plan)[0], 0)
            scheme = root / "月报.smartmap.json"
            save_plan(scheme, plan)
            self.assertEqual(find_plan_files(root), [scheme])
            loaded = load_plan(scheme)
            outputs, _ = execute_smart_plan(loaded, root / "映射结果")
            self.assertEqual(len(outputs), 1)
            result = load_workbook(outputs[0], data_only=False)
            try:
                self.assertEqual(result["报表"]["C6"].value, 10)
                self.assertEqual(result["报表"]["D7"].value, "=B4+C4")
            finally:
                result.close()

    def test_monthly_rebind_uses_current_folder(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            old_folder = root / "old"; new_folder = root / "new"
            old_folder.mkdir(); new_folder.mkdir()
            old_source = old_folder / "支行数据202605.xlsx"
            new_source = new_folder / "支行数据202606.xlsx"
            old_target = root / "历史报表.xlsx"
            new_target = root / "本月模板.xlsx"
            make_book(old_source, "数据")
            make_book(new_source, "数据")
            make_book(old_target, "报表", shifted=True)
            make_book(new_target, "报表", shifted=True)
            saved = build_plan([old_source], old_target)
            rebound = rebind_plan(saved, new_folder, new_target)
            self.assertEqual(Path(rebound.pairs[0].source_file), new_source.resolve())
            self.assertEqual(Path(rebound.target_file), new_target.resolve())


if __name__ == "__main__":
    unittest.main()
