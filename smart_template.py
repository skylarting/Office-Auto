"""Smart, label-based workbook mapping for structured reporting templates."""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
import json
import re
import shutil

import xlrd
import xlwt
from xlutils.copy import copy as copy_xls_workbook
from openpyxl import load_workbook

from excel_mapper import (
    MODE_MANUAL,
    MappingRule,
    WorkbookReader,
    _write_xls_value_preserving_style,
    execute_mapping_plan,
    output_name_for_target,
    split_address,
    unique_output_path,
    workbook_sheet_names,
)


_PUNCTUATION_RE = re.compile(r"[\s\u3000.,，;；:：/\\|｜·。．、()（）\[\]【】]+")


def normalize_label(value: object) -> str:
    """Normalize presentation-only label differences without changing meaning."""
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.replace("－", "-").replace("—", "-").replace("–", "-")
    text = text.replace("％", "%").replace("＋", "+")
    return _PUNCTUATION_RE.sub("", text)


def address_for(row: int, column: int) -> str:
    letters = ""
    number = column + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row + 1}"


@dataclass
class SheetStructure:
    row_label_columns: list[int]
    column_header_rows: list[int]
    data_rows: list[int]
    data_columns: list[int]
    confidence: float = 0.0

    @property
    def summary(self) -> str:
        row_cols = ",".join(address_for(0, col)[:-1] for col in self.row_label_columns)
        header_rows = ",".join(str(row + 1) for row in self.column_header_rows)
        return f"行标题列 {row_cols or '未识别'}；列标题行 {header_rows or '未识别'}"


@dataclass
class SheetPair:
    source_file: str
    source_sheet: str
    target_file: str
    target_sheet: str
    score: float
    confirmed: bool = False
    source_structure: SheetStructure | None = None
    target_structure: SheetStructure | None = None


@dataclass
class LabelCell:
    address: str
    row_path: tuple[str, ...]
    column_path: tuple[str, ...]
    value: object
    formula: bool

    @property
    def row_text(self) -> str:
        return " / ".join(self.row_path)

    @property
    def column_text(self) -> str:
        return " / ".join(self.column_path)


@dataclass
class SmartMatch:
    source_file: str
    source_sheet: str
    source_address: str
    source_row_path: tuple[str, ...]
    source_column_path: tuple[str, ...]
    target_file: str
    target_sheet: str
    target_address: str = ""
    target_row_path: tuple[str, ...] = ()
    target_column_path: tuple[str, ...] = ()
    score: float = 0.0
    status: str = "未匹配"
    enabled: bool = True


@dataclass
class FormulaRule:
    target_sheet: str
    row_pattern: str = ""
    column_pattern: str = ""
    formula: str = ""
    overwrite: str = "保留已有公式"
    enabled: bool = True


@dataclass
class BlockRule:
    target_sheet: str
    row_pattern: str = ""
    column_pattern: str = ""
    scope: str = "交叉位置"
    reason: str = "禁填区域"
    enabled: bool = True


@dataclass
class SmartTemplatePlan:
    target_file: str = ""
    source_files: list[str] = field(default_factory=list)
    pairs: list[SheetPair] = field(default_factory=list)
    matches: list[SmartMatch] = field(default_factory=list)
    aliases: dict[str, str] = field(default_factory=dict)
    formula_rules: list[FormulaRule] = field(default_factory=list)
    block_rules: list[BlockRule] = field(default_factory=list)
    overwrite_target_formulas: bool = False


def _sheet_values(reader: WorkbookReader, sheet_name: str) -> list[list[object]]:
    rows, columns = reader.used_dimensions(sheet_name)
    rows, columns = min(rows, 500), min(columns, 200)
    return [
        [reader.read(sheet_name, address_for(row, col)).value for col in range(columns)]
        for row in range(rows)
    ]


def _merged_anchors(reader: WorkbookReader, sheet_name: str) -> dict[tuple[int, int], tuple[int, int]]:
    result: dict[tuple[int, int], tuple[int, int]] = {}
    if reader.is_xls:
        ranges = reader.xls_book.sheet_by_name(sheet_name).merged_cells
        for row_start, row_end, col_start, col_end in ranges:
            for row in range(row_start, row_end):
                for col in range(col_start, col_end):
                    result[(row, col)] = (row_start, col_start)
    else:
        sheet = reader.format_book[sheet_name]
        for merged in sheet.merged_cells.ranges:
            for row in range(merged.min_row - 1, merged.max_row):
                for col in range(merged.min_col - 1, merged.max_col):
                    result[(row, col)] = (merged.min_row - 1, merged.min_col - 1)
    return result


def _effective(values, anchors, row: int, col: int):
    anchor = anchors.get((row, col), (row, col))
    try:
        return values[anchor[0]][anchor[1]]
    except IndexError:
        return None


def detect_sheet_structure(reader: WorkbookReader, sheet_name: str) -> SheetStructure:
    """Infer label/header/data regions; the UI may override every result."""
    values = _sheet_values(reader, sheet_name)
    if not values or not values[0]:
        return SheetStructure([], [], [], [], 0.0)
    rows, columns = len(values), len(values[0])
    numeric_by_row: list[list[int]] = []
    for row in range(rows):
        numeric = []
        for col in range(columns):
            value = values[row][col]
            traits = reader.cell_traits(sheet_name, address_for(row, col))
            if "number" in traits or "formula" in traits or isinstance(value, (int, float)):
                numeric.append(col)
        numeric_by_row.append(numeric)

    provisional_rows = [row for row, cols in enumerate(numeric_by_row) if len(cols) >= 2]
    if not provisional_rows:
        provisional_rows = [row for row, cols in enumerate(numeric_by_row) if cols]
    if not provisional_rows:
        return SheetStructure([], [], [], [], 0.15)

    text_counts = [0] * columns
    for row in provisional_rows:
        for col, value in enumerate(values[row]):
            if isinstance(value, str) and normalize_label(value):
                text_counts[col] += 1
    threshold = max(2, round(len(provisional_rows) * 0.12))
    label_candidates = [col for col, count in enumerate(text_counts) if count >= threshold]
    if label_candidates:
        data_start = max(label_candidates) + 1
    else:
        data_start = min(col for row in provisional_rows for col in numeric_by_row[row])

    data_rows = [
        row for row in provisional_rows
        if any(col >= data_start for col in numeric_by_row[row])
    ]
    first_data_row = min(data_rows)
    row_label_columns = [
        col for col in range(data_start)
        if any(
            isinstance(values[row][col], str) and normalize_label(values[row][col])
            for row in data_rows
        )
    ]
    data_columns = [
        col for col in range(data_start, columns)
        if any(values[row][col] not in (None, "") for row in data_rows)
    ]
    header_candidates = [
        row for row in range(first_data_row)
        if any(values[row][col] not in (None, "") for col in data_columns)
    ]
    column_header_rows = header_candidates[-4:]
    evidence = bool(row_label_columns and column_header_rows and data_columns)
    return SheetStructure(
        row_label_columns,
        column_header_rows,
        data_rows,
        data_columns,
        0.88 if evidence else 0.55,
    )


def build_label_index(
    reader: WorkbookReader,
    sheet_name: str,
    structure: SheetStructure,
) -> list[LabelCell]:
    values = _sheet_values(reader, sheet_name)
    anchors = _merged_anchors(reader, sheet_name)
    carried: dict[int, str] = {}
    result: list[LabelCell] = []
    for row in structure.data_rows:
        row_parts: list[str] = []
        for index, col in enumerate(structure.row_label_columns):
            value = _effective(values, anchors, row, col)
            text = str(value).strip() if isinstance(value, str) else ""
            if text:
                carried[col] = text
                for deeper in structure.row_label_columns[index + 1:]:
                    carried.pop(deeper, None)
            if carried.get(col):
                row_parts.append(carried[col])
        if not row_parts:
            continue
        for col in structure.data_columns:
            col_parts = []
            for header_row in structure.column_header_rows:
                value = _effective(values, anchors, header_row, col)
                if value not in (None, ""):
                    text = str(value).strip()
                    if text and (not col_parts or col_parts[-1] != text):
                        col_parts.append(text)
            if not col_parts:
                continue
            address = address_for(row, col)
            result.append(
                LabelCell(
                    address,
                    tuple(row_parts),
                    tuple(col_parts),
                    values[row][col],
                    "formula" in reader.cell_traits(sheet_name, address),
                )
            )
    return result


def _path_key(parts: tuple[str, ...], aliases: dict[str, str]) -> str:
    normalized = "/".join(filter(None, (normalize_label(part) for part in parts)))
    return normalize_label(aliases.get(normalized, normalized))


def _label_score(source: LabelCell, target: LabelCell, aliases: dict[str, str]) -> float:
    source_row = _path_key(source.row_path, aliases)
    source_col = _path_key(source.column_path, aliases)
    target_row = _path_key(target.row_path, aliases)
    target_col = _path_key(target.column_path, aliases)
    row_score = SequenceMatcher(None, source_row, target_row).ratio()
    col_score = SequenceMatcher(None, source_col, target_col).ratio()
    return row_score * 0.62 + col_score * 0.38


def match_pair(pair: SheetPair, aliases: dict[str, str] | None = None) -> list[SmartMatch]:
    aliases = aliases or {}
    source_reader = WorkbookReader(Path(pair.source_file))
    target_reader = WorkbookReader(Path(pair.target_file))
    try:
        source_structure = pair.source_structure or detect_sheet_structure(source_reader, pair.source_sheet)
        target_structure = pair.target_structure or detect_sheet_structure(target_reader, pair.target_sheet)
        pair.source_structure, pair.target_structure = source_structure, target_structure
        sources = build_label_index(source_reader, pair.source_sheet, source_structure)
        targets = build_label_index(target_reader, pair.target_sheet, target_structure)
        target_by_key: dict[tuple[str, str], list[LabelCell]] = {}
        for target in targets:
            key = (_path_key(target.row_path, aliases), _path_key(target.column_path, aliases))
            target_by_key.setdefault(key, []).append(target)

        matches: list[SmartMatch] = []
        for source in sources:
            if source.formula or source.value in (None, ""):
                continue
            exact = target_by_key.get(
                (_path_key(source.row_path, aliases), _path_key(source.column_path, aliases)),
                [],
            )
            candidate = exact[0] if len(exact) == 1 else None
            score = 1.0 if candidate else 0.0
            status = "自动匹配" if candidate else "未匹配"
            if candidate is None and targets:
                ranked = sorted(
                    ((_label_score(source, target, aliases), target) for target in targets),
                    key=lambda item: item[0],
                    reverse=True,
                )
                best_score, best = ranked[0]
                margin = best_score - (ranked[1][0] if len(ranked) > 1 else 0)
                if best_score >= 0.92 and margin >= 0.04:
                    candidate, score, status = best, best_score, "相似匹配"
                elif best_score >= 0.45:
                    candidate, score, status = best, best_score, "待确认"
            matches.append(
                SmartMatch(
                    pair.source_file,
                    pair.source_sheet,
                    source.address,
                    source.row_path,
                    source.column_path,
                    pair.target_file,
                    pair.target_sheet,
                    candidate.address if candidate else "",
                    candidate.row_path if candidate else (),
                    candidate.column_path if candidate else (),
                    score,
                    status,
                    bool(candidate and status != "待确认"),
                )
            )
        return matches
    finally:
        source_reader.close()
        target_reader.close()


def _sheet_signature(path: Path, sheet_name: str) -> set[str]:
    reader = WorkbookReader(path)
    try:
        values = _sheet_values(reader, sheet_name)
        labels = {
            normalize_label(value)
            for row in values[:80]
            for value in row[:30]
            if isinstance(value, str) and len(normalize_label(value)) >= 2
        }
        return set(sorted(labels, key=len, reverse=True)[:120])
    finally:
        reader.close()


def suggest_sheet_pairs(source_files: list[Path], target_file: Path) -> list[SheetPair]:
    target_sheets = workbook_sheet_names(target_file)
    target_signatures = {
        sheet: _sheet_signature(target_file, sheet) for sheet in target_sheets
    }
    result: list[SheetPair] = []
    used: set[str] = set()
    for source_file in source_files:
        for source_sheet in workbook_sheet_names(source_file):
            source_signature = _sheet_signature(source_file, source_sheet)
            ranked = []
            for target_sheet in target_sheets:
                name_score = SequenceMatcher(
                    None,
                    normalize_label(source_sheet),
                    normalize_label(target_sheet),
                ).ratio()
                union = source_signature | target_signatures[target_sheet]
                overlap = (
                    len(source_signature & target_signatures[target_sheet]) / len(union)
                    if union else 0.0
                )
                unused_bonus = 0.03 if target_sheet not in used else 0.0
                ranked.append((name_score * 0.45 + overlap * 0.52 + unused_bonus, target_sheet))
            score, target_sheet = max(ranked, default=(0.0, ""))
            if target_sheet:
                used.add(target_sheet)
            result.append(
                SheetPair(
                    str(source_file.resolve()),
                    source_sheet,
                    str(target_file.resolve()),
                    target_sheet,
                    min(score, 1.0),
                )
            )
    return result


def _pattern_matches(pattern: str, parts: tuple[str, ...]) -> bool:
    if not pattern.strip():
        return True
    wanted = normalize_label(pattern)
    return wanted in normalize_label("/".join(parts))


def match_is_blocked(match: SmartMatch, rules: list[BlockRule]) -> str | None:
    for rule in rules:
        if not rule.enabled or rule.target_sheet not in ("*", match.target_sheet):
            continue
        row_match = _pattern_matches(rule.row_pattern, match.target_row_path)
        col_match = _pattern_matches(rule.column_pattern, match.target_column_path)
        applies = (
            row_match and col_match if rule.scope == "交叉位置"
            else row_match if rule.scope == "整行"
            else col_match
        )
        if applies:
            return rule.reason or "禁填区域"
    return None


def matches_to_mapping_rules(plan: SmartTemplatePlan) -> tuple[list[MappingRule], list[str]]:
    rules: list[MappingRule] = []
    skipped: list[str] = []
    target_readers: dict[str, WorkbookReader] = {}
    try:
        for match in plan.matches:
            if not match.enabled or not match.target_address:
                continue
            blocked = match_is_blocked(match, plan.block_rules)
            if blocked:
                skipped.append(f"{match.target_sheet}!{match.target_address}：{blocked}")
                continue
            reader = target_readers.get(match.target_file)
            if reader is None:
                reader = WorkbookReader(Path(match.target_file))
                target_readers[match.target_file] = reader
            if (
                "formula" in reader.cell_traits(match.target_sheet, match.target_address)
                and not plan.overwrite_target_formulas
            ):
                skipped.append(f"{match.target_sheet}!{match.target_address}：保留已有公式")
                continue
            rules.append(
                MappingRule(
                    match.source_file,
                    match.source_sheet,
                    [match.source_address],
                    match.target_file,
                    match.target_sheet,
                    [match.target_address],
                    MODE_MANUAL,
                )
            )
    finally:
        for reader in target_readers.values():
            reader.close()
    return rules, skipped


def _render_formula(template: str, address: str) -> str:
    row, column = split_address(address)
    column_text = address_for(0, column)[:-1]
    rendered = template.strip().format(
        cell=address,
        row=row + 1,
        col=column_text,
    )
    return rendered if rendered.startswith("=") else "=" + rendered


def resolve_formula_targets(plan: SmartTemplatePlan) -> list[tuple[str, str, str, str, str]]:
    """Return target file, sheet, address, rendered formula and policy."""
    result: list[tuple[str, str, str, str, str]] = []
    seen_sheets: set[tuple[str, str]] = set()
    for pair in plan.pairs:
        key = (pair.target_file, pair.target_sheet)
        if key in seen_sheets:
            continue
        seen_sheets.add(key)
        reader = WorkbookReader(Path(pair.target_file))
        try:
            structure = pair.target_structure or detect_sheet_structure(reader, pair.target_sheet)
            for cell in build_label_index(reader, pair.target_sheet, structure):
                probe = SmartMatch(
                    "", "", "", (), (), pair.target_file, pair.target_sheet,
                    cell.address, cell.row_path, cell.column_path,
                )
                if match_is_blocked(probe, plan.block_rules):
                    continue
                for rule in plan.formula_rules:
                    if not rule.enabled or rule.target_sheet not in ("*", pair.target_sheet):
                        continue
                    if not _pattern_matches(rule.row_pattern, cell.row_path):
                        continue
                    if not _pattern_matches(rule.column_pattern, cell.column_path):
                        continue
                    traits = reader.cell_traits(pair.target_sheet, cell.address)
                    if rule.overwrite == "仅填空白" and "nonempty" in traits:
                        continue
                    if rule.overwrite == "覆盖数值但保留已有公式" and "formula" in traits:
                        continue
                    result.append(
                        (
                            pair.target_file,
                            pair.target_sheet,
                            cell.address,
                            _render_formula(rule.formula, cell.address),
                            rule.overwrite,
                        )
                    )
                    break
        finally:
            reader.close()
    return result


def _apply_formulas_to_output(
    original: Path,
    output: Path,
    formulas: list[tuple[str, str, str]],
) -> None:
    if not formulas:
        return
    if output.suffix.lower() == ".xls":
        read_book = xlrd.open_workbook(output, formatting_info=True, on_demand=False)
        write_book = copy_xls_workbook(read_book)
        try:
            for sheet, address, formula in formulas:
                _write_xls_value_preserving_style(
                    read_book,
                    write_book,
                    sheet,
                    address,
                    xlwt.Formula(formula.lstrip("=")),
                )
            write_book.save(str(output))
        finally:
            read_book.release_resources()
        return
    book = load_workbook(output, keep_vba=output.suffix.lower() == ".xlsm")
    try:
        for sheet, address, formula in formulas:
            book[sheet][address].value = formula
        book.save(output)
    finally:
        book.close()


def execute_smart_plan(
    plan: SmartTemplatePlan,
    output_folder: Path,
    progress=None,
) -> tuple[list[Path], list[str]]:
    value_rules, skipped = matches_to_mapping_rules(plan)
    source_files = list(dict.fromkeys(Path(rule.source_file) for rule in value_rules))
    target_files = list(
        dict.fromkeys(
            [Path(plan.target_file)]
            + [Path(pair.target_file) for pair in plan.pairs if pair.target_file]
        )
    )
    if value_rules:
        outputs = execute_mapping_plan(
            source_files,
            target_files,
            value_rules,
            output_folder,
            progress,
        )
    else:
        output_folder.mkdir(parents=True, exist_ok=True)
        outputs = []
        for target in target_files:
            output = unique_output_path(output_folder / output_name_for_target(target))
            shutil.copy2(target, output)
            outputs.append(output)

    output_by_target = {
        str(target.resolve()): output for target, output in zip(target_files, outputs)
    }
    grouped: dict[str, list[tuple[str, str, str]]] = {}
    for target_file, sheet, address, formula, _policy in resolve_formula_targets(plan):
        grouped.setdefault(str(Path(target_file).resolve()), []).append((sheet, address, formula))
    for target_key, formulas in grouped.items():
        output = output_by_target.get(target_key)
        if output is not None:
            _apply_formulas_to_output(Path(target_key), output, formulas)
    return outputs, skipped


def save_plan(path: Path, plan: SmartTemplatePlan) -> None:
    path.write_text(json.dumps(asdict(plan), ensure_ascii=False, indent=2), encoding="utf-8")


def load_plan(path: Path) -> SmartTemplatePlan:
    raw = json.loads(path.read_text(encoding="utf-8"))
    raw["pairs"] = [
        SheetPair(
            **{
                **item,
                "source_structure": SheetStructure(**item["source_structure"])
                if item.get("source_structure") else None,
                "target_structure": SheetStructure(**item["target_structure"])
                if item.get("target_structure") else None,
            }
        )
        for item in raw.get("pairs", [])
    ]
    raw["matches"] = [SmartMatch(**item) for item in raw.get("matches", [])]
    raw["formula_rules"] = [FormulaRule(**item) for item in raw.get("formula_rules", [])]
    raw["block_rules"] = [BlockRule(**item) for item in raw.get("block_rules", [])]
    return SmartTemplatePlan(**raw)
