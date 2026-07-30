"""GUI tool for extracting configurable Excel cells into summary sheets."""

from __future__ import annotations

from argparse import ArgumentParser
from dataclasses import dataclass
from pathlib import Path
import re
import sys
from tkinter import (
    BOTH,
    END,
    LEFT,
    RIGHT,
    W,
    X,
    Canvas,
    IntVar,
    StringVar,
    Text,
    Tk,
    Toplevel,
    filedialog,
    messagebox,
    ttk,
)
from typing import Callable, Iterable

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


DEFAULT_CELL_ADDRESSES = [
    "A2", "C2", "A3", "C3",
]
VALID_CELL_RE = re.compile(r"^[A-Z]{1,3}[1-9][0-9]*$")
CELL_RANGE_RE = re.compile(
    r"^([A-Z]{1,3}[1-9][0-9]*)-([A-Z]{1,3}[1-9][0-9]*)$"
)
WORKBOOK_SUFFIXES = {".xls", ".xlsx", ".xlsm"}


@dataclass
class SummaryRecord:
    file_name: str
    sheet_name: str
    values: list[object]
    number_formats: list[str]


def parse_cell_addresses(raw_text: str) -> list[str]:
    """Parse comma, semicolon, whitespace, or Chinese punctuation separators."""
    parts = re.split(r"[\s,，;；]+", raw_text.strip().upper())
    addresses: list[str] = []
    invalid: list[str] = []

    for part in parts:
        if not part:
            continue
        range_match = CELL_RANGE_RE.fullmatch(part)
        if range_match:
            try:
                expanded = expand_cell_range(*range_match.groups())
            except ValueError:
                invalid.append(part)
                continue
            for address in expanded:
                if address not in addresses:
                    addresses.append(address)
        elif not VALID_CELL_RE.fullmatch(part):
            invalid.append(part)
        elif part not in addresses:
            addresses.append(part)

    if invalid:
        raise ValueError("以下单元格地址不合法：" + "、".join(invalid))
    if not addresses:
        raise ValueError("请至少填写一个需要提取的单元格地址。")
    return addresses


def expand_cell_range(start: str, end: str) -> list[str]:
    start_column_text = "".join(char for char in start if char.isalpha())
    start_row = int("".join(char for char in start if char.isdigit()))
    end_column_text = "".join(char for char in end if char.isalpha())
    end_row = int("".join(char for char in end if char.isdigit()))
    start_column = column_letters_to_number(start_column_text)
    end_column = column_letters_to_number(end_column_text)
    first_column, last_column = sorted((start_column, end_column))
    first_row, last_row = sorted((start_row, end_row))
    count = (last_column - first_column + 1) * (last_row - first_row + 1)
    if count > 10000:
        raise ValueError("单元格范围过大。")
    return [
        f"{column_number_to_letters(column)}{row}"
        for row in range(first_row, last_row + 1)
        for column in range(first_column, last_column + 1)
    ]


def is_source_workbook(path: Path, summary_name: str = "汇总") -> bool:
    name = path.name
    if name.startswith("~$") or name.startswith(".~"):
        return False
    if path.suffix.lower() not in WORKBOOK_SUFFIXES:
        return False
    excluded_stems = {
        "SheetSummary",
        "批量汇总表",
        f"{summary_name}",
    }
    if path.stem in excluded_stems:
        return False
    if path.stem.endswith("_汇总表") or path.stem.endswith("_汇总表_Python"):
        return False
    return True


def find_workbooks(path: Path, source_mode: str, summary_name: str) -> list[Path]:
    if source_mode == "file":
        if not path.is_file() or not is_source_workbook(path, summary_name):
            raise ValueError("请选择有效的 .xls、.xlsx 或 .xlsm 文件。")
        return [path]

    if not path.is_dir():
        raise ValueError("请选择一个有效的文件夹。")

    files = sorted(
        item for item in path.iterdir()
        if item.is_file() and is_source_workbook(item, summary_name)
    )
    if not files:
        raise ValueError(
            "所选文件夹中没有可处理的 .xls、.xlsx 或 .xlsm 文件。"
        )
    return files


def split_cell_address(address: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)", address)
    if not match:
        raise ValueError(f"单元格地址不合法：{address}")
    return int(match.group(2)) - 1, column_letters_to_number(match.group(1)) - 1


def xls_cell_value(book, cell):
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    return cell.value


def xls_number_format(book, cell) -> str:
    try:
        xf = book.xf_list[cell.xf_index]
        return book.format_map[xf.format_key].format_str or "General"
    except (AttributeError, IndexError, KeyError):
        return "General"


def read_xls_records(
    source_path: Path,
    addresses: list[str],
    summary_name: str,
) -> list[SummaryRecord]:
    book = xlrd.open_workbook(
        source_path,
        formatting_info=True,
        on_demand=True,
    )
    records: list[SummaryRecord] = []
    try:
        for sheet in book.sheets():
            if sheet.name == summary_name:
                continue
            values: list[object] = []
            number_formats: list[str] = []
            for address in addresses:
                row, column = split_cell_address(address)
                if row >= sheet.nrows or column >= sheet.ncols:
                    values.append(None)
                    number_formats.append("General")
                    continue
                cell = sheet.cell(row, column)
                values.append(xls_cell_value(book, cell))
                number_formats.append(xls_number_format(book, cell))
            records.append(
                SummaryRecord(
                    file_name=source_path.name,
                    sheet_name=sheet.name,
                    values=values,
                    number_formats=number_formats,
                )
            )
    finally:
        book.release_resources()
    return records


def read_records(
    source_path: Path,
    addresses: list[str],
    summary_name: str,
) -> list[SummaryRecord]:
    if source_path.suffix.lower() == ".xls":
        return read_xls_records(source_path, addresses, summary_name)

    keep_vba = source_path.suffix.lower() == ".xlsm"
    value_book = load_workbook(
        source_path,
        data_only=True,
        keep_vba=keep_vba,
    )
    format_book = load_workbook(
        source_path,
        data_only=False,
        keep_vba=keep_vba,
    )
    records: list[SummaryRecord] = []

    try:
        for sheet_name in value_book.sheetnames:
            if sheet_name == summary_name:
                continue
            value_sheet = value_book[sheet_name]
            format_sheet = format_book[sheet_name]
            records.append(
                SummaryRecord(
                    file_name=source_path.name,
                    sheet_name=sheet_name,
                    values=[value_sheet[address].value for address in addresses],
                    number_formats=[
                        format_sheet[address].number_format
                        for address in addresses
                    ],
                )
            )
    finally:
        value_book.close()
        format_book.close()

    return records


def format_summary_sheet(
    sheet,
    addresses: list[str],
    records: list[SummaryRecord],
    include_file_name: bool,
) -> None:
    headers = (
        ["文件名称", "Sheet名称", *addresses]
        if include_file_name
        else ["Sheet名称", *addresses]
    )
    sheet.append(headers)

    value_start_column = 3 if include_file_name else 2
    for row_number, record in enumerate(records, start=2):
        if include_file_name:
            sheet.cell(row=row_number, column=1, value=record.file_name)
            sheet.cell(row=row_number, column=2, value=record.sheet_name)
        else:
            sheet.cell(row=row_number, column=1, value=record.sheet_name)

        for offset, (value, number_format) in enumerate(
            zip(record.values, record.number_formats),
        ):
            target = sheet.cell(
                row=row_number,
                column=value_start_column + offset,
                value=value,
            )
            target.number_format = number_format

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 24 if include_file_name else 18
    if include_file_name:
        sheet.column_dimensions["B"].width = 18
    for column in range(value_start_column, len(headers) + 1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=column).column_letter
        ].width = 14

    thin_side = Side(style="thin", color="808080")
    all_borders = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    for row in sheet.iter_rows(
        min_row=1,
        max_row=sheet.max_row,
        min_col=1,
        max_col=sheet.max_column,
    ):
        for cell in row:
            cell.border = all_borders


def create_new_summary_workbook(
    source_files: list[Path],
    addresses: list[str],
    output_path: Path,
    summary_name: str,
    progress: Callable[[int, int, str], None] | None = None,
) -> None:
    all_records: list[SummaryRecord] = []
    total = len(source_files)

    for index, source_path in enumerate(source_files, start=1):
        if progress:
            progress(index - 1, total, f"正在读取：{source_path.name}")
        all_records.extend(read_records(source_path, addresses, summary_name))

    output_book = Workbook()
    output_sheet = output_book.active
    output_sheet.title = summary_name
    format_summary_sheet(
        output_sheet,
        addresses,
        all_records,
        include_file_name=len(source_files) > 1,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_book.save(output_path)
    output_book.close()

    if progress:
        progress(total, total, f"已生成：{output_path.name}")


def unique_sheet_name(workbook, preferred_name: str) -> str:
    if preferred_name not in workbook.sheetnames:
        return preferred_name

    index = 2
    while True:
        suffix = f" ({index})"
        candidate = f"{preferred_name[:31 - len(suffix)]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
        index += 1


def unique_output_path(path: Path) -> Path:
    if not path.exists():
        return path

    index = 2
    while True:
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate
        index += 1


def copy_xls_to_xlsx(source_path: Path):
    source_book = xlrd.open_workbook(
        source_path,
        formatting_info=True,
        on_demand=True,
    )
    output_book = Workbook()
    output_book.remove(output_book.active)
    try:
        for source_sheet in source_book.sheets():
            output_sheet = output_book.create_sheet(source_sheet.name)
            for row_index in range(source_sheet.nrows):
                for column_index in range(source_sheet.ncols):
                    source_cell = source_sheet.cell(row_index, column_index)
                    target_cell = output_sheet.cell(
                        row=row_index + 1,
                        column=column_index + 1,
                        value=xls_cell_value(source_book, source_cell),
                    )
                    target_cell.number_format = xls_number_format(
                        source_book,
                        source_cell,
                    )
    finally:
        source_book.release_resources()
    return output_book


def create_summary_copy(
    source_path: Path,
    addresses: list[str],
    summary_name: str,
    output_path: Path,
) -> Path:
    records = read_records(source_path, addresses, summary_name)
    if source_path.suffix.lower() == ".xls":
        workbook = copy_xls_to_xlsx(source_path)
    else:
        keep_vba = source_path.suffix.lower() == ".xlsm"
        workbook = load_workbook(source_path, keep_vba=keep_vba)

    actual_name = unique_sheet_name(workbook, summary_name)
    summary_sheet = workbook.create_sheet(actual_name, 0)
    format_summary_sheet(
        summary_sheet,
        addresses,
        records,
        include_file_name=False,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path


def create_summary_copies(
    source_files: list[Path],
    addresses: list[str],
    summary_name: str,
    output_target: Path,
    progress: Callable[[int, int, str], None] | None = None,
) -> list[Path]:
    total = len(source_files)
    output_paths: list[Path] = []
    for index, source_path in enumerate(source_files, start=1):
        if progress:
            progress(index - 1, total, f"正在处理：{source_path.name}")
        copy_name = (
            f"{source_path.stem}.xlsx"
            if source_path.suffix.lower() == ".xls"
            else source_path.name
        )
        proposed = (
            output_target / copy_name
            if len(source_files) > 1
            else output_target
        )
        output_path = unique_output_path(proposed)
        create_summary_copy(
            source_path,
            addresses,
            summary_name,
            output_path,
        )
        output_paths.append(output_path)
        if progress:
            progress(index, total, f"已完成：{source_path.name}")
    return output_paths


def default_output_path(source_path: Path, source_mode: str) -> Path:
    if source_mode == "file":
        return source_path.with_name(f"{source_path.stem}_汇总表.xlsx")
    return source_path / "批量汇总表.xlsx"


def default_copy_target(source_path: Path, source_mode: str) -> Path:
    if source_mode == "file":
        suffix = (
            ".xlsx"
            if source_path.suffix.lower() == ".xls"
            else source_path.suffix
        )
        return source_path.with_name(
            f"{source_path.stem}_已汇总{suffix}"
        )
    return source_path / "汇总结果"


def column_letters_to_number(letters: str) -> int:
    value = 0
    for character in letters.upper():
        if character < "A" or character > "Z":
            raise ValueError("列名必须由英文字母组成，例如 Z 或 AZ。")
        value = value * 26 + ord(character) - ord("A") + 1
    return value


def column_number_to_letters(number: int) -> str:
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def cell_sort_key(address: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)", address)
    if not match:
        return (sys.maxsize, sys.maxsize)
    return (int(match.group(2)), column_letters_to_number(match.group(1)))


class CellGridPicker(ttk.Frame):
    """Scrollable canvas-based cell picker kept in sync with manual input."""

    CELL_WIDTH = 48
    CELL_HEIGHT = 28
    ROW_HEADER_WIDTH = 48
    COLUMN_HEADER_HEIGHT = 28
    MAX_ROWS = 200
    MAX_COLUMNS = 52

    def __init__(
        self,
        parent,
        on_change: Callable[[list[str]], None],
    ) -> None:
        super().__init__(parent)
        self.on_change = on_change
        self.row_count = 30
        self.column_count = 26
        self.selected: set[str] = set()

        controls = ttk.Frame(self)
        controls.pack(fill=X, pady=(0, 8))
        ttk.Label(controls, text="显示行数：").pack(side=LEFT)
        self.rows_var = IntVar(value=self.row_count)
        ttk.Spinbox(
            controls,
            from_=1,
            to=self.MAX_ROWS,
            width=6,
            textvariable=self.rows_var,
        ).pack(side=LEFT)
        ttk.Label(controls, text="显示至列：").pack(side=LEFT, padx=(14, 0))
        self.end_column_var = StringVar(value="Z")
        ttk.Entry(
            controls,
            width=6,
            textvariable=self.end_column_var,
        ).pack(side=LEFT)
        ttk.Button(
            controls,
            text="更新网格",
            command=self.update_dimensions,
        ).pack(side=LEFT, padx=(8, 0))
        ttk.Label(
            controls,
            text="点击单元格可选中或取消",
            foreground="#666666",
        ).pack(side=LEFT, padx=(16, 0))

        grid_frame = ttk.Frame(self)
        grid_frame.pack(fill=BOTH, expand=True)
        self.canvas = Canvas(
            grid_frame,
            height=220,
            background="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#B8B8B8",
        )
        x_scroll = ttk.Scrollbar(
            grid_frame,
            orient="horizontal",
            command=self.canvas.xview,
        )
        y_scroll = ttk.Scrollbar(
            grid_frame,
            orient="vertical",
            command=self.canvas.yview,
        )
        self.canvas.configure(
            xscrollcommand=x_scroll.set,
            yscrollcommand=y_scroll.set,
        )
        self.canvas.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        grid_frame.rowconfigure(0, weight=1)
        grid_frame.columnconfigure(0, weight=1)
        self.canvas.bind("<Button-1>", self._handle_click)
        self._draw()

    def update_dimensions(self) -> None:
        try:
            rows = int(self.rows_var.get())
            columns = column_letters_to_number(
                self.end_column_var.get().strip()
            )
            if not 1 <= rows <= self.MAX_ROWS:
                raise ValueError(
                    f"显示行数应在 1 至 {self.MAX_ROWS} 之间。"
                )
            if not 1 <= columns <= self.MAX_COLUMNS:
                raise ValueError(
                    "可视选择器最多显示至 AZ 列；更远地址可用手动输入。"
                )
            self.row_count = rows
            self.column_count = columns
            self.end_column_var.set(column_number_to_letters(columns))
            self._draw()
        except (TypeError, ValueError) as exc:
            messagebox.showerror("网格范围错误", str(exc), parent=self)

    def set_selected(self, addresses: Iterable[str]) -> None:
        self.selected = set(addresses)
        self._draw()

    def _handle_click(self, event) -> None:
        x = self.canvas.canvasx(event.x)
        y = self.canvas.canvasy(event.y)
        if x < self.ROW_HEADER_WIDTH or y < self.COLUMN_HEADER_HEIGHT:
            return

        column = int((x - self.ROW_HEADER_WIDTH) // self.CELL_WIDTH) + 1
        row = int((y - self.COLUMN_HEADER_HEIGHT) // self.CELL_HEIGHT) + 1
        if not 1 <= column <= self.column_count or not 1 <= row <= self.row_count:
            return

        address = f"{column_number_to_letters(column)}{row}"
        if address in self.selected:
            self.selected.remove(address)
        else:
            self.selected.add(address)
        ordered = sorted(self.selected, key=cell_sort_key)
        self._draw()
        self.on_change(ordered)

    def _draw(self) -> None:
        self.canvas.delete("all")
        total_width = (
            self.ROW_HEADER_WIDTH + self.column_count * self.CELL_WIDTH
        )
        total_height = (
            self.COLUMN_HEADER_HEIGHT + self.row_count * self.CELL_HEIGHT
        )
        self.canvas.configure(scrollregion=(0, 0, total_width, total_height))

        self.canvas.create_rectangle(
            0,
            0,
            self.ROW_HEADER_WIDTH,
            self.COLUMN_HEADER_HEIGHT,
            fill="#E9EDF3",
            outline="#C5CBD3",
        )
        for column in range(1, self.column_count + 1):
            x1 = self.ROW_HEADER_WIDTH + (column - 1) * self.CELL_WIDTH
            x2 = x1 + self.CELL_WIDTH
            self.canvas.create_rectangle(
                x1,
                0,
                x2,
                self.COLUMN_HEADER_HEIGHT,
                fill="#E9EDF3",
                outline="#C5CBD3",
            )
            self.canvas.create_text(
                (x1 + x2) / 2,
                self.COLUMN_HEADER_HEIGHT / 2,
                text=column_number_to_letters(column),
                fill="#303640",
            )

        for row in range(1, self.row_count + 1):
            y1 = self.COLUMN_HEADER_HEIGHT + (row - 1) * self.CELL_HEIGHT
            y2 = y1 + self.CELL_HEIGHT
            self.canvas.create_rectangle(
                0,
                y1,
                self.ROW_HEADER_WIDTH,
                y2,
                fill="#E9EDF3",
                outline="#C5CBD3",
            )
            self.canvas.create_text(
                self.ROW_HEADER_WIDTH / 2,
                (y1 + y2) / 2,
                text=str(row),
                fill="#303640",
            )

            for column in range(1, self.column_count + 1):
                x1 = self.ROW_HEADER_WIDTH + (column - 1) * self.CELL_WIDTH
                x2 = x1 + self.CELL_WIDTH
                address = f"{column_number_to_letters(column)}{row}"
                selected = address in self.selected
                self.canvas.create_rectangle(
                    x1,
                    y1,
                    x2,
                    y2,
                    fill="#2F75B5" if selected else "#FFFFFF",
                    outline="#D8DDE5",
                )
                if selected:
                    self.canvas.create_text(
                        (x1 + x2) / 2,
                        (y1 + y2) / 2,
                        text=address,
                        fill="#FFFFFF",
                    )


class SourceSelectionDialog:
    """One browser that can return either a workbook or a directory."""

    def __init__(self, parent: Tk, initial_path: Path | None = None) -> None:
        self.result: Path | None = None
        self.current_dir = self._initial_directory(initial_path)
        self.items: dict[str, Path] = {}

        self.window = Toplevel(parent)
        self.window.title("选择表格文件或文件夹")
        self.window.geometry("720x480")
        self.window.minsize(600, 400)
        self.window.transient(parent)
        self.window.grab_set()

        container = ttk.Frame(self.window, padding=14)
        container.pack(fill=BOTH, expand=True)

        navigation = ttk.Frame(container)
        navigation.pack(fill=X, pady=(0, 8))
        ttk.Button(
            navigation,
            text="上一级",
            command=self._go_up,
        ).pack(side=LEFT)
        self.path_var = StringVar(value=str(self.current_dir))
        path_entry = ttk.Entry(navigation, textvariable=self.path_var)
        path_entry.pack(side=LEFT, fill=X, expand=True, padx=(8, 0))
        path_entry.bind("<Return>", self._go_to_entered_path)

        ttk.Label(
            container,
            text="选择一个 Excel 文件，或选择文件夹以处理其中的全部 Excel 文件。",
        ).pack(anchor=W, pady=(0, 8))

        tree_frame = ttk.Frame(container)
        tree_frame.pack(fill=BOTH, expand=True)
        self.tree = ttk.Treeview(
            tree_frame,
            columns=("type",),
            show="tree headings",
            selectmode="browse",
        )
        self.tree.heading("#0", text="名称")
        self.tree.heading("type", text="类型")
        self.tree.column("#0", width=480, anchor=W)
        self.tree.column("type", width=120, anchor=W)
        scroll = ttk.Scrollbar(
            tree_frame,
            orient="vertical",
            command=self.tree.yview,
        )
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.pack(side=LEFT, fill=BOTH, expand=True)
        scroll.pack(side=RIGHT, fill="y")
        self.tree.bind("<Double-1>", self._double_click)

        actions = ttk.Frame(container)
        actions.pack(fill=X, pady=(10, 0))
        ttk.Button(
            actions,
            text="确定",
            command=self._confirm,
        ).pack(side=LEFT)
        ttk.Button(
            actions,
            text="取消",
            command=self.window.destroy,
        ).pack(side=LEFT, padx=(8, 0))

        self._refresh()
        self.window.protocol("WM_DELETE_WINDOW", self.window.destroy)
        parent.wait_window(self.window)

    @staticmethod
    def _initial_directory(initial_path: Path | None) -> Path:
        if initial_path:
            candidate = initial_path if initial_path.is_dir() else initial_path.parent
            if candidate.is_dir():
                return candidate
        return Path.home()

    def _refresh(self) -> None:
        self.tree.delete(*self.tree.get_children())
        self.items.clear()
        self.path_var.set(str(self.current_dir))

        current_id = self.tree.insert(
            "",
            END,
            text="[选择当前文件夹]",
            values=("文件夹",),
        )
        self.items[current_id] = self.current_dir

        try:
            entries = sorted(
                self.current_dir.iterdir(),
                key=lambda path: (not path.is_dir(), path.name.lower()),
            )
        except OSError as exc:
            messagebox.showerror(
                "无法打开文件夹",
                str(exc),
                parent=self.window,
            )
            return

        for path in entries:
            if path.name.startswith("."):
                continue
            if not path.is_dir() and path.suffix.lower() not in WORKBOOK_SUFFIXES:
                continue
            item_id = self.tree.insert(
                "",
                END,
                text=path.name,
                values=("文件夹" if path.is_dir() else "Excel 文件",),
            )
            self.items[item_id] = path

    def _go_up(self) -> None:
        parent = self.current_dir.parent
        if parent != self.current_dir:
            self.current_dir = parent
            self._refresh()

    def _go_to_entered_path(self, _event=None) -> None:
        path = Path(self.path_var.get().strip())
        if path.is_file():
            self.result = path
            self.window.destroy()
        elif path.is_dir():
            self.current_dir = path
            self._refresh()
        else:
            messagebox.showerror(
                "路径无效",
                "请输入有效的文件夹或 Excel 文件路径。",
                parent=self.window,
            )

    def _double_click(self, _event=None) -> None:
        selection = self.tree.selection()
        if not selection:
            return
        path = self.items[selection[0]]
        if path.is_dir() and self.tree.item(selection[0], "text") != "[选择当前文件夹]":
            self.current_dir = path
            self._refresh()
        elif path.is_file():
            self.result = path
            self.window.destroy()

    def _confirm(self) -> None:
        selection = self.tree.selection()
        if not selection:
            messagebox.showinfo(
                "请选择数据来源",
                "请先选中一个 Excel 文件或文件夹。",
                parent=self.window,
            )
            return
        self.result = self.items[selection[0]]
        self.window.destroy()


class SummaryApp:
    def __init__(self, root: Tk) -> None:
        self.root = root
        self.root.title("Excel 单元格汇总工具")
        self.root.geometry("920x660")
        self.root.minsize(820, 560)

        self.source_mode = ""
        self.source_choice = StringVar(value="选择一个表格文件")
        self.source_path = StringVar()
        self.source_summary = StringVar(value="尚未选择数据来源")
        self.output_choice = StringVar(
            value="创建新的汇总文件，集中提取指定单元格数据"
        )
        self.output_path = StringVar()
        self.summary_name = StringVar(value="汇总")
        self.selection_summary = StringVar()
        self.selection_details = StringVar()
        self.validation_text = StringVar()
        self.status_text = StringVar(value="等待开始")
        self.progress_percent = StringVar(value="0%")
        self.selected_addresses = list(DEFAULT_CELL_ADDRESSES)
        self.cells_valid = True
        self._syncing_cells = False

        self._build_ui()
        self._set_selected_addresses(DEFAULT_CELL_ADDRESSES)
        self._show_output_fields()
        self._reset_progress()

    def _build_ui(self) -> None:
        container = ttk.Frame(self.root, padding=18)
        container.pack(fill=BOTH, expand=True)

        action_row = ttk.Frame(container)
        action_row.pack(side="bottom", fill=X)
        self.start_button = ttk.Button(
            action_row,
            text="开始提取",
            command=self._run,
            state="disabled",
        )
        self.start_button.pack(side=LEFT)
        ttk.Button(
            action_row,
            text="退出",
            command=self.root.destroy,
        ).pack(side=LEFT, padx=(8, 0))

        progress_frame = ttk.LabelFrame(
            container,
            text="3. 处理进度",
            padding=10,
        )
        progress_frame.pack(side="bottom", fill=X, pady=(10, 10))
        progress_title = ttk.Frame(progress_frame)
        progress_title.pack(fill=X, pady=(0, 5))
        ttk.Label(progress_title, text="处理进度").pack(side=LEFT)
        ttk.Label(
            progress_title,
            textvariable=self.progress_percent,
        ).pack(side=RIGHT)
        self.progress = ttk.Progressbar(progress_frame, mode="determinate")
        self.progress.pack(fill=X)
        ttk.Label(
            progress_frame,
            textvariable=self.status_text,
        ).pack(anchor=W, pady=(5, 0))

        top_row = ttk.Frame(container)
        top_row.pack(fill=X, pady=(0, 10))
        top_row.columnconfigure(0, weight=1, uniform="top")
        top_row.columnconfigure(1, weight=1, uniform="top")

        source_frame = ttk.LabelFrame(
            top_row,
            text="1. 选择数据来源",
            padding=10,
        )
        source_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 5))

        source_actions = ttk.Frame(source_frame)
        source_actions.pack(fill=X)
        source_combo = ttk.Combobox(
            source_actions,
            textvariable=self.source_choice,
            values=("选择一个表格文件", "选择文件夹中的所有表格"),
            state="readonly",
            width=24,
        )
        source_combo.pack(side=LEFT, fill=X, expand=True)
        source_combo.bind(
            "<<ComboboxSelected>>",
            self._source_choice_changed,
        )
        ttk.Button(
            source_actions,
            text="浏览...",
            command=self._browse_selected_source,
        ).pack(side=LEFT, padx=(8, 0))
        ttk.Label(
            source_frame,
            textvariable=self.source_summary,
            foreground="#3A3A3A",
            wraplength=390,
        ).pack(anchor=W, pady=(8, 0))

        output_frame = ttk.LabelFrame(
            top_row,
            text="选择输出方式",
            padding=10,
        )
        output_frame.grid(row=0, column=1, sticky="nsew", padx=(5, 0))
        output_selector = ttk.Frame(output_frame)
        output_selector.pack(fill=X)
        ttk.Label(output_selector, text="输出方式：").pack(side=LEFT)
        output_combo = ttk.Combobox(
            output_selector,
            textvariable=self.output_choice,
            values=(
                "创建新的汇总文件，集中提取指定单元格数据",
                "另存为工作簿副本，在首页插入汇总工作表",
            ),
            state="readonly",
            width=30,
        )
        output_combo.pack(side=LEFT, fill=X, expand=True)
        output_combo.bind("<<ComboboxSelected>>", self._output_changed)
        self.output_fields = ttk.Frame(output_frame)
        self.output_fields.pack(fill=X, pady=(8, 0))

        cell_section = ttk.Frame(container)
        cell_section.pack(fill=BOTH, expand=True)
        cell_header = ttk.Frame(cell_section)
        cell_header.pack(fill=X, pady=(0, 5))
        ttk.Label(
            cell_header,
            text="2. 选择需要提取的单元格",
        ).pack(side=LEFT)

        cell_frame = ttk.LabelFrame(cell_section, padding=10)
        cell_frame.pack(fill=BOTH, expand=True)
        self.cell_notebook = ttk.Notebook(cell_frame)
        self.cell_notebook.pack(fill=BOTH, expand=True)

        manual_tab = ttk.Frame(self.cell_notebook, padding=10)
        visual_tab = ttk.Frame(self.cell_notebook, padding=10)
        self.cell_notebook.add(manual_tab, text="手动输入")
        self.cell_notebook.add(visual_tab, text="点击选择")

        manual_toolbar = ttk.Frame(manual_tab)
        manual_toolbar.pack(fill=X)
        ttk.Label(
            manual_toolbar,
            text="可用逗号、空格或换行分隔，例如：A2, C2, A3, C3",
        ).pack(side=LEFT)
        ttk.Button(
            manual_toolbar,
            text="恢复默认",
            command=self._restore_default_cells,
        ).pack(side=RIGHT)
        ttk.Button(
            manual_toolbar,
            text="清空",
            command=self._clear_cells,
        ).pack(side=RIGHT, padx=(0, 6))
        self.cell_text = Text(manual_tab, height=4, wrap="word")
        self.cell_text.pack(fill=X, pady=(8, 6))
        self.cell_text.bind("<KeyRelease>", self._manual_cells_changed)
        ttk.Label(
            manual_tab,
            textvariable=self.validation_text,
            foreground="#C62828",
        ).pack(anchor=W)
        ttk.Label(
            manual_tab,
            textvariable=self.selection_summary,
        ).pack(anchor=W, pady=(8, 0))

        visual_toolbar = ttk.Frame(visual_tab)
        visual_toolbar.pack(fill=X, pady=(0, 8))
        ttk.Label(
            visual_toolbar,
            textvariable=self.selection_details,
            wraplength=610,
            justify=LEFT,
        ).pack(side=LEFT, fill=X, expand=True)
        ttk.Button(
            visual_toolbar,
            text="恢复默认",
            command=self._restore_default_cells,
        ).pack(side=RIGHT)
        ttk.Button(
            visual_toolbar,
            text="清空",
            command=self._clear_cells,
        ).pack(side=RIGHT, padx=(0, 6))
        self.cell_grid = CellGridPicker(
            visual_tab,
            self._grid_cells_changed,
        )
        self.cell_grid.pack(fill=BOTH, expand=True)


    def _browse_source(self) -> None:
        initial = Path(self.source_path.get()) if self.source_path.get() else None
        dialog = SourceSelectionDialog(self.root, initial)
        selected = dialog.result
        if not selected:
            return

        if selected.is_file():
            if not is_source_workbook(selected, self.summary_name.get()):
                messagebox.showerror(
                    "文件无效",
                    "请选择有效的 .xls、.xlsx 或 .xlsm 文件。",
                    parent=self.root,
                )
                return
            self.source_mode = "file"
            self.source_path.set(str(selected))
            self.source_summary.set(f"已选择文件：{selected}")
            self._source_changed()
            return

        try:
            files = find_workbooks(
                selected,
                "folder",
                self.summary_name.get(),
            )
        except ValueError as exc:
            messagebox.showerror("无法选择文件夹", str(exc), parent=self.root)
            return

        confirmed = messagebox.askyesno(
            "确认批量提取",
            f"该文件夹中发现 {len(files)} 个可处理的 Excel 文件。\n\n"
            "是否提取该文件夹下的所有表格文件？",
            parent=self.root,
        )
        if not confirmed:
            return

        self.source_mode = "folder"
        self.source_path.set(str(selected))
        self.source_summary.set(
            f"已选择文件夹：{selected}（共 {len(files)} 个 Excel 文件）"
        )
        self._source_changed()

    def _browse_selected_source(self) -> None:
        if self.source_choice.get() == "选择文件夹中的所有表格":
            self._choose_folder()
        else:
            self._choose_file()

    def _source_choice_changed(self, _event=None) -> None:
        self.source_mode = ""
        self.source_path.set("")
        self.output_path.set("")
        self.source_summary.set("尚未选择数据来源")
        self._show_output_fields()
        self._configuration_changed()

    def _choose_file(self) -> None:
        selected = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[
                ("Excel 文件", "*.xls *.xlsx *.xlsm"),
                ("所有文件", "*.*"),
            ],
        )
        if selected:
            path = Path(selected)
            self.source_mode = "file"
            self.source_path.set(str(path))
            self.source_summary.set(f"已选择文件：{path}")
            self._source_changed()

    def _choose_folder(self) -> None:
        selected = filedialog.askdirectory(
            title="选择包含 Excel 文件的文件夹"
        )
        if not selected:
            return

        path = Path(selected)
        try:
            files = find_workbooks(path, "folder", self.summary_name.get())
        except ValueError as exc:
            messagebox.showerror("无法选择文件夹", str(exc), parent=self.root)
            return

        confirmed = messagebox.askyesno(
            "确认批量提取",
            f"该文件夹中发现 {len(files)} 个可处理的 Excel 文件。\n\n"
            "是否提取该文件夹下的所有表格文件？",
            parent=self.root,
        )
        if not confirmed:
            return

        self.source_mode = "folder"
        self.source_path.set(str(path))
        self.source_summary.set(
            f"已选择文件夹：{path}（共 {len(files)} 个 Excel 文件）"
        )
        self._source_changed()

    def _source_changed(self) -> None:
        source_path = Path(self.source_path.get())
        default_path = (
            default_output_path(source_path, self.source_mode)
            if self.output_choice.get()
            == "创建新的汇总文件，集中提取指定单元格数据"
            else default_copy_target(source_path, self.source_mode)
        )
        self.output_path.set(str(default_path))
        self._show_output_fields()
        self._configuration_changed()

    def _browse_output(self) -> None:
        if (
            self.output_choice.get()
            == "另存为工作簿副本，在首页插入汇总工作表"
            and self.source_mode == "folder"
        ):
            selected = filedialog.askdirectory(
                title="选择副本输出文件夹"
            )
        else:
            current = Path(self.output_path.get() or "汇总表.xlsx")
            suffix = (
                (
                    ".xlsx"
                    if Path(self.source_path.get()).suffix.lower() == ".xls"
                    else Path(self.source_path.get()).suffix
                )
                if self.output_choice.get()
                == "另存为工作簿副本，在首页插入汇总工作表"
                else ".xlsx"
            )
            selected = filedialog.asksaveasfilename(
                title="选择输出文件位置",
                defaultextension=suffix,
                filetypes=[
                    ("Excel 工作簿", f"*{suffix}"),
                    ("所有文件", "*.*"),
                ],
                initialfile=current.name,
            )
        if selected:
            self.output_path.set(selected)
            self._configuration_changed()

    def _output_changed(self, _event=None) -> None:
        if self.source_path.get():
            source_path = Path(self.source_path.get())
            default_path = (
                default_output_path(source_path, self.source_mode)
                if self.output_choice.get()
                == "创建新的汇总文件，集中提取指定单元格数据"
                else default_copy_target(source_path, self.source_mode)
            )
            self.output_path.set(str(default_path))
        self._show_output_fields()
        self._configuration_changed()

    def _show_output_fields(self) -> None:
        for child in self.output_fields.winfo_children():
            child.destroy()

        location_row = ttk.Frame(self.output_fields)
        location_row.pack(fill=X)
        location_label = (
            "输出文件夹："
            if (
                self.output_choice.get()
                == "另存为工作簿副本，在首页插入汇总工作表"
                and self.source_mode == "folder"
            )
            else "输出文件："
        )
        ttk.Label(location_row, text=location_label).pack(side=LEFT)
        output_entry = ttk.Entry(
            location_row,
            textvariable=self.output_path,
        )
        output_entry.pack(side=LEFT, fill=X, expand=True)
        output_entry.bind("<KeyRelease>", self._configuration_changed)
        ttk.Button(
            location_row,
            text="浏览...",
            command=self._browse_output,
        ).pack(side=LEFT, padx=(6, 0))

        if (
            self.output_choice.get()
            == "另存为工作簿副本，在首页插入汇总工作表"
        ):
            name_row = ttk.Frame(self.output_fields)
            name_row.pack(fill=X, pady=(6, 0))
            ttk.Label(name_row, text="表单名称：").pack(side=LEFT)
            name_entry = ttk.Entry(
                name_row,
                textvariable=self.summary_name,
                width=14,
            )
            name_entry.pack(side=LEFT)
            name_entry.bind("<KeyRelease>", self._configuration_changed)
            ttk.Label(
                name_row,
                text="只写入新副本；原文件不会修改",
                foreground="#666666",
            ).pack(side=LEFT, padx=(8, 0))

    def _manual_cells_changed(self, _event=None) -> None:
        if self._syncing_cells:
            return
        try:
            addresses = parse_cell_addresses(self.cell_text.get("1.0", END))
            self.validation_text.set("")
            self._set_selected_addresses(addresses, update_text=False)
        except ValueError as exc:
            self.cells_valid = False
            self.validation_text.set(str(exc))
            self.start_button.configure(state="disabled")
            self._reset_progress()

    def _grid_cells_changed(self, addresses: list[str]) -> None:
        self._set_selected_addresses(addresses)

    def _set_selected_addresses(
        self,
        addresses: Iterable[str],
        update_text: bool = True,
    ) -> None:
        ordered = sorted(set(addresses), key=cell_sort_key)
        self.selected_addresses = ordered
        self.cells_valid = bool(ordered)
        self._syncing_cells = True
        try:
            if update_text:
                self.cell_text.delete("1.0", END)
                self.cell_text.insert("1.0", ", ".join(ordered))
            self.cell_grid.set_selected(ordered)
        finally:
            self._syncing_cells = False

        preview = ", ".join(ordered[:8])
        if len(ordered) > 8:
            preview += "..."
        self.selection_summary.set(
            f"已选择 {len(ordered)} 个单元格"
            + (f"：{preview}" if preview else "")
        )
        self.selection_details.set(
            f"已选择 {len(ordered)} 个单元格："
            + (", ".join(ordered) if ordered else "无")
        )
        self.validation_text.set(
            "" if ordered else "请至少选择一个需要提取的单元格。"
        )
        self._configuration_changed()

    def _restore_default_cells(self) -> None:
        self._set_selected_addresses(DEFAULT_CELL_ADDRESSES)

    def _clear_cells(self) -> None:
        self._set_selected_addresses([])

    def _configuration_changed(self, _event=None) -> None:
        self._reset_progress()
        valid = bool(
            self.source_path.get()
            and self.selected_addresses
            and self.cells_valid
        )
        valid = valid and bool(self.output_path.get().strip())
        if (
            self.output_choice.get()
            == "另存为工作簿副本，在首页插入汇总工作表"
        ):
            valid = valid and bool(self.summary_name.get().strip())
        self.start_button.configure(state="normal" if valid else "disabled")

    def _reset_progress(self) -> None:
        self.progress.configure(maximum=100, value=0, mode="determinate")
        self.progress_percent.set("0%")
        self.status_text.set("等待开始")

    def _progress(self, current: int, total: int, message: str) -> None:
        maximum = max(total, 1)
        self.progress.configure(maximum=maximum, value=current)
        percentage = int(current / maximum * 100)
        self.progress_percent.set(f"{percentage}%")
        self.status_text.set(message)
        self.root.update_idletasks()

    def _run(self) -> None:
        try:
            self._reset_progress()
            summary_name = (
                "汇总"
                if self.output_choice.get()
                == "创建新的汇总文件，集中提取指定单元格数据"
                else self.summary_name.get().strip()
            )
            if not summary_name:
                raise ValueError("请填写汇总表名称。")
            if any(char in summary_name for char in r"[]:*?/\\"):
                raise ValueError("汇总表名称包含 Excel 不允许的字符。")
            if len(summary_name) > 31:
                raise ValueError("汇总表名称不能超过 31 个字符。")

            addresses = list(self.selected_addresses)
            source_path = Path(self.source_path.get().strip())
            source_files = find_workbooks(
                source_path,
                self.source_mode,
                summary_name,
            )

            output_text = self.output_path.get().strip()
            if not output_text:
                raise ValueError("请选择输出位置。")
            output_target = Path(output_text)

            if (
                self.output_choice.get()
                == "创建新的汇总文件，集中提取指定单元格数据"
            ):
                output_path = output_target
                if output_path.resolve() in {
                    path.resolve() for path in source_files
                }:
                    raise ValueError("输出文件不能覆盖源文件。")
                create_new_summary_workbook(
                    source_files,
                    addresses,
                    output_path,
                    summary_name,
                    self._progress,
                )
                result_message = f"汇总文件已生成：\n{output_path}"
            else:
                if (
                    len(source_files) == 1
                    and output_target.resolve() == source_files[0].resolve()
                ):
                    raise ValueError("副本输出文件不能与源文件相同。")
                output_paths = create_summary_copies(
                    source_files,
                    addresses,
                    summary_name,
                    output_target,
                    self._progress,
                )
                result_message = (
                    f"已生成 {len(output_paths)} 个新副本；原文件未修改。\n\n"
                    f"输出位置：{output_target}"
                )

            self.progress.configure(value=self.progress.cget("maximum"))
            self.progress_percent.set("100%")
            self.status_text.set("处理完成")
            messagebox.showinfo("处理完成", result_message, parent=self.root)
        except Exception as exc:
            self.status_text.set("处理失败。")
            messagebox.showerror("处理失败", str(exc), parent=self.root)


def run_command_line() -> None:
    parser = ArgumentParser(description="汇总 Excel 工作簿中的指定单元格。")
    parser.add_argument("source", type=Path, help="源工作簿路径")
    parser.add_argument("-o", "--output", type=Path, help="输出工作簿路径")
    parser.add_argument(
        "-c",
        "--cells",
        default=",".join(DEFAULT_CELL_ADDRESSES),
        help="需要提取的单元格，以逗号分隔",
    )
    args = parser.parse_args()

    addresses = parse_cell_addresses(args.cells)
    output_path = args.output or args.source.with_name(
        f"{args.source.stem}_汇总表.xlsx"
    )
    create_new_summary_workbook(
        [args.source],
        addresses,
        output_path,
        "汇总",
    )
    print(f"汇总表已生成：{output_path.resolve()}")


def main() -> None:
    if len(sys.argv) > 1:
        run_command_line()
        return
    root = Tk()
    SummaryApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
