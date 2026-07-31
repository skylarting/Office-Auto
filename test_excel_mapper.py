from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import xlrd
import xlwt

import excel_mapper


def make_xlsx(path: Path, sheet_name: str, values: dict[str, object]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for address, value in values.items():
        sheet[address] = value
    workbook.save(path)
    workbook.close()


def make_xls(path: Path, sheet_name: str) -> None:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet(sheet_name)
    style = xlwt.easyxf(num_format_str="0")
    sheet.write(1, 0, 12.6, style)
    sheet.write(1, 2, 34.4, style)
    target_style = xlwt.easyxf(
        "font: bold on, colour white;"
        "pattern: pattern solid, fore_colour dark_blue;"
        "borders: left thin, right thin, top thin, bottom thin;",
        num_format_str="#,##0.00",
    )
    sheet.write(3, 3, None, target_style)
    sheet.write(3, 4, None, target_style)
    workbook.save(str(path))


class ExcelMapperTests(unittest.TestCase):
    def test_workbook_files_in_folder(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            (folder / "甲.xlsx").touch()
            (folder / "乙.xls").touch()
            (folder / "~$临时.xlsx").touch()
            (folder / "说明.txt").touch()
            self.assertEqual(
                [path.name for path in excel_mapper.workbook_files_in_folder(folder)],
                ["乙.xls", "甲.xlsx"],
            )

    def test_workbook_reader_exposes_cell_fill_colors(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            xlsx_path = folder / "颜色.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "数据"
            sheet["A1"] = "表头"
            sheet["A1"].fill = PatternFill(
                fill_type="solid",
                fgColor="FF336699",
            )
            workbook.save(xlsx_path)
            workbook.close()

            reader = excel_mapper.WorkbookReader(xlsx_path)
            self.assertEqual(reader.fill_color("数据", "A1"), "#336699")
            self.assertIsNone(reader.fill_color("数据", "B1"))
            reader.close()

            xls_path = folder / "颜色.xls"
            legacy = xlwt.Workbook()
            legacy_sheet = legacy.add_sheet("数据")
            style = xlwt.easyxf(
                "pattern: pattern solid, fore_colour yellow;"
            )
            legacy_sheet.write(0, 0, "表头", style)
            legacy.save(str(xls_path))

            legacy_reader = excel_mapper.WorkbookReader(xls_path)
            self.assertIsNotNone(legacy_reader.fill_color("数据", "A1"))
            legacy_reader.close()

    def test_workbook_reader_classifies_cell_types(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            xlsx_path = folder / "类型.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "数据"
            sheet["A1"] = 12.5
            sheet["A2"] = "文字"
            sheet["A3"] = "=A1*2"
            sheet["A4"].fill = PatternFill(
                fill_type="solid",
                fgColor="FFFFFF00",
            )
            workbook.save(xlsx_path)
            workbook.close()
            reader = excel_mapper.WorkbookReader(xlsx_path)
            self.assertIn("number", reader.cell_traits("数据", "A1"))
            self.assertIn("text", reader.cell_traits("数据", "A2"))
            self.assertIn("formula", reader.cell_traits("数据", "A3"))
            self.assertIn("fill", reader.cell_traits("数据", "A4"))
            reader.close()

            xls_path = folder / "类型.xls"
            legacy = xlwt.Workbook()
            legacy_sheet = legacy.add_sheet("数据")
            legacy_sheet.write(0, 0, 12.5)
            legacy_sheet.write(1, 0, "文字")
            legacy_sheet.write(2, 0, xlwt.Formula("A1*2"))
            legacy.save(str(xls_path))
            legacy_reader = excel_mapper.WorkbookReader(xls_path)
            self.assertIn("number", legacy_reader.cell_traits("数据", "A1"))
            self.assertIn("text", legacy_reader.cell_traits("数据", "A2"))
            self.assertIn("formula", legacy_reader.cell_traits("数据", "A3"))
            legacy_reader.close()

    def test_expand_sequence_and_one_to_many(self) -> None:
        sequence = excel_mapper.MappingRule(
            "source.xlsx",
            "数据",
            ["A2", "C2"],
            "target.xlsx",
            "模板",
            ["B5", "D5"],
            excel_mapper.MODE_SEQUENCE,
        )
        expanded = excel_mapper.expand_rule(sequence)
        self.assertEqual(
            [(item.source_cell, item.target_cell) for item in expanded],
            [("A2", "B5"), ("C2", "D5")],
        )

        one_to_many = excel_mapper.MappingRule(
            "source.xlsx",
            "数据",
            ["A2"],
            "target.xlsx",
            "模板",
            ["B5", "D5", "F5"],
            excel_mapper.MODE_ONE_TO_MANY,
        )
        expanded = excel_mapper.expand_rule(one_to_many)
        self.assertEqual(
            [item.target_cell for item in expanded],
            ["B5", "D5", "F5"],
        )

    def test_invalid_mapping_counts(self) -> None:
        rule = excel_mapper.MappingRule(
            "source.xlsx",
            "数据",
            ["A2", "A3"],
            "target.xlsx",
            "模板",
            ["B5"],
            excel_mapper.MODE_SEQUENCE,
        )
        with self.assertRaises(ValueError):
            excel_mapper.expand_rule(rule)

    def test_execute_multiple_targets_without_modifying_originals(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            source = folder / "来源.xlsx"
            target_one = folder / "目标一.xlsx"
            target_two = folder / "目标二.xlsx"
            output = folder / "输出"
            make_xlsx(source, "数据", {"A2": 12.6, "C2": 34.4})
            make_xlsx(target_one, "模板", {"B5": "原值"})
            make_xlsx(target_two, "报表", {})

            source_book = load_workbook(source)
            source_book["数据"]["A2"].number_format = "0"
            source_book["数据"]["C2"].number_format = "0"
            source_book.save(source)
            source_book.close()
            target_book = load_workbook(target_one)
            target_cell = target_book["模板"]["B5"]
            target_cell.fill = PatternFill("solid", fgColor="F4B183")
            target_cell.number_format = "0.00"
            target_book.save(target_one)
            target_book.close()

            rules = [
                excel_mapper.MappingRule(
                    str(source),
                    "数据",
                    ["A2", "C2"],
                    str(target_one),
                    "模板",
                    ["B5", "D5"],
                    excel_mapper.MODE_SEQUENCE,
                ),
                excel_mapper.MappingRule(
                    str(source),
                    "数据",
                    ["A2"],
                    str(target_two),
                    "报表",
                    ["C3", "E3"],
                    excel_mapper.MODE_ONE_TO_MANY,
                ),
            ]
            outputs = excel_mapper.execute_mapping_plan(
                [source],
                [target_one, target_two],
                rules,
                output,
            )
            self.assertEqual(len(outputs), 2)

            original = load_workbook(target_one)
            self.assertEqual(original["模板"]["B5"].value, "原值")
            original.close()

            mapped_one = load_workbook(output / "目标一_已映射.xlsx")
            self.assertEqual(mapped_one["模板"]["B5"].value, 12.6)
            self.assertEqual(mapped_one["模板"]["D5"].value, 34.4)
            self.assertEqual(mapped_one["模板"]["B5"].number_format, "0.00")
            self.assertEqual(
                mapped_one["模板"]["B5"].fill.fgColor.rgb,
                "00F4B183",
            )
            mapped_one.close()

            mapped_two = load_workbook(output / "目标二_已映射.xlsx")
            self.assertEqual(mapped_two["报表"]["C3"].value, 12.6)
            self.assertEqual(mapped_two["报表"]["E3"].value, 12.6)
            mapped_two.close()

    def test_xls_source_and_target_are_supported(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            source = folder / "来源.xls"
            target = folder / "目标.xls"
            output = folder / "输出"
            make_xls(source, "数据")
            make_xls(target, "模板")
            rule = excel_mapper.MappingRule(
                str(source),
                "数据",
                ["A2", "C2"],
                str(target),
                "模板",
                ["D4", "E4"],
                excel_mapper.MODE_SEQUENCE,
            )

            outputs = excel_mapper.execute_mapping_plan(
                [source],
                [target],
                [rule],
                output,
            )
            self.assertEqual(outputs[0].suffix, ".xls")
            copied = xlrd.open_workbook(
                outputs[0],
                formatting_info=True,
            )
            copied_sheet = copied.sheet_by_name("模板")
            self.assertEqual(copied_sheet.cell_value(3, 3), 12.6)
            self.assertEqual(copied_sheet.cell_value(3, 4), 34.4)
            for column in (3, 4):
                xf = copied.xf_list[
                    copied_sheet.cell(3, column).xf_index
                ]
                font = copied.font_list[xf.font_index]
                self.assertTrue(font.bold)
                self.assertEqual(xf.background.fill_pattern, 1)
                self.assertEqual(xf.border.left_line_style, 1)
                self.assertEqual(copied.format_map[xf.format_key].format_str, "#,##0.00")
            copied.release_resources()

    def test_project_round_trip(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            project = folder / "方案.txt"
            source = folder / "来源.xlsx"
            target = folder / "目标.xlsx"
            rule = excel_mapper.MappingRule(
                str(source),
                "数据",
                ["A2"],
                str(target),
                "模板",
                ["B5"],
                excel_mapper.MODE_MANUAL,
            )
            excel_mapper.save_mapping_project(
                project,
                [rule],
            )
            sources, targets, rules, output = (
                excel_mapper.load_mapping_project(project)
            )
            self.assertEqual(sources, [source.resolve()])
            self.assertEqual(targets, [target.resolve()])
            self.assertEqual(Path(rules[0].source_file), source.resolve())
            self.assertEqual(Path(rules[0].target_file), target.resolve())
            self.assertEqual(rules[0].source_cells, ["A2"])
            self.assertEqual(rules[0].target_cells, ["B5"])
            self.assertEqual(output, (folder / "映射结果").resolve())
            text = project.read_text(encoding="utf-8-sig")
            self.assertEqual(
                text,
                "【来源.xlsx；数据；A2 → 目标.xlsx；模板；B5】\n",
            )

    def test_excel_scheme_uses_paired_source_and_target_rows(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            scheme = folder / "映射方案.xlsx"
            source = folder / "来源.xlsx"
            target = folder / "目标.xlsx"
            rules = [
                excel_mapper.MappingRule(
                    str(source),
                    "数据",
                    ["A1", "A2", "B1", "B2"],
                    str(target),
                    "模板",
                    ["A1", "A2", "B1", "B2"],
                    excel_mapper.MODE_SEQUENCE,
                )
            ]

            excel_mapper.save_excel_mapping_scheme(
                scheme,
                rules,
                folder,
            )
            workbook = load_workbook(scheme)
            sheet = workbook["映射方案"]
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                list(excel_mapper.SCHEME_HEADERS),
            )
            self.assertEqual(
                [cell.value for cell in sheet[2]],
                [1, "来源", "来源.xlsx", "数据", "A1-B2"],
            )
            self.assertEqual(
                [cell.value for cell in sheet[3]],
                [1, "目标", "目标.xlsx", "模板", "同位置"],
            )
            self.assertNotEqual(
                sheet["A2"].fill.fgColor.rgb,
                sheet["A3"].fill.fgColor.rgb,
            )
            workbook.close()

            sources, targets, loaded = (
                excel_mapper.load_excel_mapping_scheme(scheme)
            )
            self.assertEqual(sources, [source.resolve()])
            self.assertEqual(targets, [target.resolve()])
            self.assertEqual(loaded[0].source_cells, ["A1", "A2", "B1", "B2"])
            self.assertEqual(loaded[0].target_cells, loaded[0].source_cells)

    def test_compact_excel_scheme_infers_alternating_source_target_rows(
        self,
    ) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            scheme = folder / "精简映射方案.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["工作簿", "工作表", "单元格"])
            sheet.append(["来源一.xlsx", "数据", "A1-A3"])
            sheet.append(["目标一.xlsx", "模板", "同位置"])
            sheet.append(["来源二.xlsx", "汇总", "C2"])
            sheet.append(["目标二.xlsx", "报表", "D5"])
            workbook.save(scheme)
            workbook.close()

            sources, targets, rules = (
                excel_mapper.load_excel_mapping_scheme(scheme)
            )

            self.assertEqual(len(rules), 2)
            self.assertEqual(sources[0], (folder / "来源一.xlsx").resolve())
            self.assertEqual(targets[1], (folder / "目标二.xlsx").resolve())
            self.assertEqual(rules[0].source_cells, ["A1", "A2", "A3"])
            self.assertEqual(rules[0].target_cells, ["A1", "A2", "A3"])
            self.assertEqual(rules[1].target_cells, ["D5"])

    def test_compact_scheme_keeps_source_without_target_as_editable_draft(
        self,
    ) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            scheme = folder / "未完成方案.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["工作簿", "工作表", "单元格"])
            sheet.append(["来源一.xlsx", "数据", "A1-A3"])
            workbook.save(scheme)
            workbook.close()

            sources, targets, rules = (
                excel_mapper.load_excel_mapping_scheme(scheme)
            )

            self.assertEqual(len(rules), 1)
            self.assertEqual(len(sources), 1)
            self.assertEqual(targets, [])
            self.assertEqual(rules[0].target_file, "")
            self.assertEqual(rules[0].target_sheet, "")
            self.assertEqual(rules[0].source_row, 2)
            self.assertIsNone(rules[0].target_row)

    def test_four_column_scheme_infers_groups_from_source_target_types(
        self,
    ) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            scheme = folder / "四列映射方案.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["类型", "工作簿", "工作表", "单元格"])
            sheet.append(["来源", "来源一.xlsx", "数据", "A1-A3"])
            sheet.append(["目标", "目标一.xlsx", "模板", "同位置"])
            sheet.append(["来源", "来源二.xlsx", "汇总", "C2"])
            workbook.save(scheme)
            workbook.close()

            sources, targets, rules = (
                excel_mapper.load_excel_mapping_scheme(scheme)
            )

            self.assertEqual(len(rules), 2)
            self.assertEqual(len(sources), 2)
            self.assertEqual(len(targets), 1)
            self.assertEqual(rules[0].target_cells, rules[0].source_cells)
            self.assertEqual(rules[1].target_file, "")
            self.assertEqual(rules[1].source_row, 4)
            self.assertIsNone(rules[1].target_row)

    def test_project_text_can_be_serialized_and_parsed_in_memory(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            source = folder / "来源.xlsx"
            target = folder / "目标.xlsx"
            rule = excel_mapper.MappingRule(
                str(source),
                "数据",
                ["B3", "B4", "B5", "B6"],
                str(target),
                "模板",
                ["A1", "B1", "C1", "D1"],
                excel_mapper.MODE_SEQUENCE,
            )

            text = excel_mapper.serialize_mapping_project([rule], folder)
            sources, targets, rules = excel_mapper.parse_mapping_project_text(
                text,
                folder,
            )

            self.assertEqual(sources, [source.resolve()])
            self.assertEqual(targets, [target.resolve()])
            self.assertEqual(rules[0].source_cells, ["B3", "B4", "B5", "B6"])
            self.assertEqual(rules[0].target_cells, ["A1", "B1", "C1", "D1"])

    def test_txt_example_starts_with_field_explanation(self) -> None:
        self.assertTrue(
            excel_mapper.TXT_EXAMPLE.startswith(
                "【来源工作簿.xlsx；来源工作表名称；A1,A2,A3 → "
                "目标工作簿.xlsx；目标工作表名称；A1,A2,A3】"
            )
        )
        self.assertNotIn("每行填写一条映射", excel_mapper.TXT_EXAMPLE)
        self.assertIn("每行填写一条映射", excel_mapper.TXT_INSTRUCTIONS)

    def test_txt_project_can_be_written_by_hand(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            project = folder / "手工方案.txt"
            project.write_text(
                "{来源.xlsx;数据;A1 => 目标.xlsx;模板;B2,C2}\n"
                "（来源.xlsx；数据；C3-C5 >> 目标.xlsx；模板；同位置）\n",
                encoding="utf-8-sig",
            )

            sources, targets, rules, output = (
                excel_mapper.load_mapping_project(project)
            )

            self.assertEqual(sources, [(folder / "来源.xlsx").resolve()])
            self.assertEqual(targets, [(folder / "目标.xlsx").resolve()])
            self.assertEqual(rules[0].mode, excel_mapper.MODE_ONE_TO_MANY)
            self.assertEqual(rules[1].source_cells, ["C3", "C4", "C5"])
            self.assertEqual(rules[1].target_cells, ["C3", "C4", "C5"])
            self.assertEqual(output, (folder / "映射结果").resolve())

    def test_txt_project_rejects_unmatched_brackets(self) -> None:
        with TemporaryDirectory() as temp_dir:
            project = Path(temp_dir) / "错误方案.txt"
            project.write_text(
                "【来源.xlsx；数据；A1 → 目标.xlsx；模板；A1]\n",
                encoding="utf-8-sig",
            )
            with self.assertRaisesRegex(ValueError, "括号不匹配"):
                excel_mapper.load_mapping_project(project)

    def test_mapping_mode_is_inferred(self) -> None:
        self.assertEqual(
            excel_mapper.infer_mapping_mode(["A1"], ["B1"]),
            excel_mapper.MODE_MANUAL,
        )
        self.assertEqual(
            excel_mapper.infer_mapping_mode(["A1"], ["B1", "C1"]),
            excel_mapper.MODE_ONE_TO_MANY,
        )
        self.assertEqual(
            excel_mapper.infer_mapping_mode(["A1", "A2"], ["B1", "B2"]),
            excel_mapper.MODE_SEQUENCE,
        )
        with self.assertRaises(ValueError):
            excel_mapper.infer_mapping_mode(
                ["A1", "A2"],
                ["B1", "B2", "B3"],
            )

    def test_cell_addresses_are_compacted_for_display(self) -> None:
        self.assertEqual(
            excel_mapper.format_cell_addresses(
                ["A1", "A2", "A3", "B1", "B2", "B3"]
            ),
            "A1-B3",
        )
        self.assertEqual(
            excel_mapper.format_cell_addresses(["A1", "A2", "A3"]),
            "A1-A3",
        )
        self.assertEqual(
            excel_mapper.format_cell_addresses(["A1", "C3"]),
            "A1,C3",
        )
        self.assertEqual(
            excel_mapper.format_cell_addresses(
                [
                    "A1", "A2", "A3",
                    "B1", "B2", "B3",
                    "D1", "D2", "D3",
                ]
            ),
            "A1-B3,D1-D3",
        )

    def test_best_name_match_prefers_same_or_common_prefix(self) -> None:
        self.assertEqual(
            excel_mapper.best_name_match(
                "测试数据 - 副本.xlsx",
                ["空白表.xlsx", "测试数据1.xlsx", "测试数据.xlsx"],
            ),
            "测试数据.xlsx",
        )
        self.assertEqual(
            excel_mapper.best_name_match(
                "江苏分行统计",
                ["浙江分行统计", "江苏分行汇总", "其他"],
            ),
            "江苏分行汇总",
        )
        self.assertIsNone(
            excel_mapper.best_name_match("甲", ["乙", "丙"])
        )


if __name__ == "__main__":
    unittest.main()
