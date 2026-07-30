"""GUI for mapping cells between multiple Excel workbooks."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import re
import shutil
from tkinter import (
    BOTH,
    END,
    LEFT,
    RIGHT,
    W,
    X,
    Canvas,
    StringVar,
    Text,
    Tk,
    Toplevel,
    filedialog,
    messagebox,
    ttk,
)
from typing import Callable

import xlrd
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

from export_summary import (
    WORKBOOK_SUFFIXES,
    column_letters_to_number,
    copy_xls_to_xlsx,
    parse_cell_addresses,
    unique_output_path,
    xls_cell_value,
    xls_number_format,
)


MODE_SEQUENCE = "按顺序一一对应"
MODE_ONE_TO_MANY = "一个来源写入多个目标"
MODE_MANUAL = "逐条手动设置"
PROJECT_BRACKETS = {
    "【": "】",
    "[": "]",
    "{": "}",
    "(": ")",
    "（": "）",
    "〔": "〕",
}
PROJECT_DIRECTION_RE = re.compile(r"\s*(?:→|->|=>|>>|》)\s*")
PROJECT_FIELD_RE = re.compile(r"\s*[；;]\s*")
TXT_EXAMPLE = """【来源工作簿.xlsx；来源工作表名称；A1,A2,A3 → 目标工作簿.xlsx；目标工作表名称；A1,A2,A3】
【测试数据.xlsx；汇总；A1,B3,D5 → 目标.xlsx；Sheet1；同位置】
【测试数据.xlsx；数据；B3-B6 → 目标.xlsx；Sheet2；A1-D1】
【来源.xlsx；数据；A1-B3 → 目标.xlsx；模板；D1-E3】
【C:\\业务资料\\来源.xlsx；统计表；C3-C8 → D:\\报表模板\\目标.xlsx；汇总；E3-E8】
"""

TXT_INSTRUCTIONS = """说明：
1. 每行填写一条映射，每一侧依次填写：工作簿；工作表；单元格。
2. 字段可使用中文分号“；”或英文分号“;”。
3. B3-B6、A1-D1 表示连续范围；A1-B3 表示左上角到右下角的矩形区域。
4. “同位置”表示写入地址与读取地址相同。
5. 只写文件名时，工作簿应位于“方案基准文件夹”中。
6. 外层括号支持：【】、[]、{}、()、（）和〔〕。
7. 方向箭头支持：→、->、=>、>> 和 》。
8. 输出默认保存在方案基准文件夹的“映射结果”文件夹。
"""


@dataclass
class CellValue:
    value: object
    number_format: str


@dataclass
class MappingRule:
    source_file: str
    source_sheet: str
    source_cells: list[str]
    target_file: str
    target_sheet: str
    target_cells: list[str]
    mode: str


@dataclass
class ExpandedMapping:
    source_file: str
    source_sheet: str
    source_cell: str
    target_file: str
    target_sheet: str
    target_cell: str


def workbook_files_in_folder(folder: Path) -> list[Path]:
    return sorted(
        path
        for path in folder.iterdir()
        if (
            path.is_file()
            and path.suffix.lower() in WORKBOOK_SUFFIXES
            and not path.name.startswith(("~$", ".~"))
        )
    )


def workbook_sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    book = load_workbook(
        path,
        read_only=True,
        keep_vba=path.suffix.lower() == ".xlsm",
    )
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def expand_rule(rule: MappingRule) -> list[ExpandedMapping]:
    sources = rule.source_cells
    targets = rule.target_cells
    if not sources or not targets:
        raise ValueError("来源单元格和目标单元格都不能为空。")

    if rule.mode == MODE_ONE_TO_MANY:
        if len(sources) != 1:
            raise ValueError("“一个来源写入多个目标”只能填写一个来源单元格。")
        pairs = [(sources[0], target) for target in targets]
    elif rule.mode == MODE_MANUAL:
        if len(sources) != 1 or len(targets) != 1:
            raise ValueError("“逐条手动设置”每条规则只能填写一个来源和一个目标。")
        pairs = [(sources[0], targets[0])]
    elif rule.mode == MODE_SEQUENCE:
        if len(sources) != len(targets):
            raise ValueError(
                "“按顺序一一对应”的来源和目标单元格数量必须相同。"
            )
        pairs = list(zip(sources, targets))
    else:
        raise ValueError(f"不支持的映射方式：{rule.mode}")

    return [
        ExpandedMapping(
            source_file=rule.source_file,
            source_sheet=rule.source_sheet,
            source_cell=source,
            target_file=rule.target_file,
            target_sheet=rule.target_sheet,
            target_cell=target,
        )
        for source, target in pairs
    ]


def infer_mapping_mode(
    source_cells: list[str],
    target_cells: list[str],
) -> str:
    if not source_cells or not target_cells:
        raise ValueError("来源单元格和目标单元格都不能为空。")
    if len(source_cells) == len(target_cells):
        return MODE_MANUAL if len(source_cells) == 1 else MODE_SEQUENCE
    if len(source_cells) == 1:
        return MODE_ONE_TO_MANY
    raise ValueError(
        "无法自动判断映射关系：多个来源单元格只能对应相同数量的目标单元格；"
        "一个来源单元格可以对应多个目标单元格。"
    )


def expand_rules(rules: list[MappingRule]) -> list[ExpandedMapping]:
    expanded: list[ExpandedMapping] = []
    for index, rule in enumerate(rules, start=1):
        try:
            expanded.extend(expand_rule(rule))
        except ValueError as exc:
            raise ValueError(f"第 {index} 条映射：{exc}") from exc
    return expanded


class WorkbookReader:
    DEFAULT_THEME_COLORS = (
        "FFFFFF",
        "000000",
        "EEECE1",
        "1F497D",
        "4F81BD",
        "C0504D",
        "9BBB59",
        "8064A2",
        "4BACC6",
        "F79646",
        "0000FF",
        "800080",
    )

    def __init__(self, path: Path) -> None:
        self.path = path
        self.is_xls = path.suffix.lower() == ".xls"
        if self.is_xls:
            self.xls_book = xlrd.open_workbook(
                path,
                formatting_info=True,
                on_demand=True,
            )
            self.value_book = None
            self.format_book = None
        else:
            keep_vba = path.suffix.lower() == ".xlsm"
            self.xls_book = None
            self.value_book = load_workbook(
                path,
                data_only=True,
                read_only=False,
                keep_vba=keep_vba,
            )
            self.format_book = load_workbook(
                path,
                data_only=False,
                read_only=False,
                keep_vba=keep_vba,
            )

    def read(self, sheet_name: str, address: str) -> CellValue:
        if self.is_xls:
            if sheet_name not in self.xls_book.sheet_names():
                raise ValueError(f"{self.path.name} 中不存在工作表“{sheet_name}”。")
            sheet = self.xls_book.sheet_by_name(sheet_name)
            row, column = split_address(address)
            if row >= sheet.nrows or column >= sheet.ncols:
                return CellValue(None, "General")
            cell = sheet.cell(row, column)
            return CellValue(
                xls_cell_value(self.xls_book, cell),
                xls_number_format(self.xls_book, cell),
            )

        if sheet_name not in self.value_book.sheetnames:
            raise ValueError(f"{self.path.name} 中不存在工作表“{sheet_name}”。")
        return CellValue(
            self.value_book[sheet_name][address].value,
            self.format_book[sheet_name][address].number_format,
        )

    @staticmethod
    def _apply_tint(rgb: str, tint: float) -> str:
        channels = [int(rgb[index:index + 2], 16) for index in (0, 2, 4)]
        adjusted: list[int] = []
        for channel in channels:
            value = (
                channel * (1 + tint)
                if tint < 0
                else channel * (1 - tint) + 255 * tint
            )
            adjusted.append(max(0, min(255, round(value))))
        return "".join(f"{value:02X}" for value in adjusted)

    def fill_color(self, sheet_name: str, address: str) -> str | None:
        if self.is_xls:
            sheet = self.xls_book.sheet_by_name(sheet_name)
            row, column = split_address(address)
            if row >= sheet.nrows or column >= sheet.ncols:
                return None
            cell = sheet.cell(row, column)
            xf = self.xls_book.xf_list[cell.xf_index]
            background = xf.background
            if not background.fill_pattern:
                return None
            rgb = self.xls_book.colour_map.get(
                background.pattern_colour_index
            )
            return (
                f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
                if rgb is not None
                else None
            )

        cell = self.format_book[sheet_name][address]
        if not cell.fill.fill_type:
            return None
        color = cell.fill.fgColor
        rgb: str | None = None
        if color.type == "rgb" and color.rgb:
            rgb = str(color.rgb)[-6:]
        elif color.type == "indexed" and color.indexed is not None:
            index = int(color.indexed)
            if 0 <= index < len(COLOR_INDEX):
                rgb = COLOR_INDEX[index][-6:]
        elif color.type == "theme" and color.theme is not None:
            index = int(color.theme)
            if 0 <= index < len(self.DEFAULT_THEME_COLORS):
                rgb = self.DEFAULT_THEME_COLORS[index]
        if rgb is None or not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
            return None
        tint = float(color.tint or 0)
        if tint:
            rgb = self._apply_tint(rgb, tint)
        return f"#{rgb.upper()}"

    def close(self) -> None:
        if self.xls_book is not None:
            self.xls_book.release_resources()
        if self.value_book is not None:
            self.value_book.close()
        if self.format_book is not None:
            self.format_book.close()


def split_address(address: str) -> tuple[int, int]:
    addresses = parse_cell_addresses(address)
    if len(addresses) != 1:
        raise ValueError(f"单元格地址不合法：{address}")
    normalized = addresses[0]
    column_text = "".join(char for char in normalized if char.isalpha())
    row_text = "".join(char for char in normalized if char.isdigit())
    return int(row_text) - 1, column_letters_to_number(column_text) - 1


def validate_mapping_plan(
    source_files: list[Path],
    target_files: list[Path],
    rules: list[MappingRule],
) -> tuple[list[ExpandedMapping], list[str]]:
    if not source_files:
        raise ValueError("请至少添加一个来源工作簿。")
    if not target_files:
        raise ValueError("请至少添加一个目标工作簿。")
    if not rules:
        raise ValueError("请至少添加一条映射规则。")

    source_set = {str(path.resolve()) for path in source_files}
    target_set = {str(path.resolve()) for path in target_files}
    for path in [*source_files, *target_files]:
        if not path.is_file() or path.suffix.lower() not in WORKBOOK_SUFFIXES:
            raise ValueError(f"文件不存在或格式不支持：{path}")

    expanded = expand_rules(rules)
    seen_targets: dict[tuple[str, str, str], int] = {}
    warnings: list[str] = []
    readers: dict[str, WorkbookReader] = {}
    try:
        for index, mapping in enumerate(expanded, start=1):
            source_key = str(Path(mapping.source_file).resolve())
            target_key = str(Path(mapping.target_file).resolve())
            if source_key not in source_set:
                raise ValueError(f"第 {index} 项使用了未添加的来源文件。")
            if target_key not in target_set:
                raise ValueError(f"第 {index} 项使用了未添加的目标文件。")

            source_reader = readers.get(source_key)
            if source_reader is None:
                source_reader = WorkbookReader(Path(mapping.source_file))
                readers[source_key] = source_reader
            source_reader.read(mapping.source_sheet, mapping.source_cell)

            target_reader = readers.get(target_key)
            if target_reader is None:
                target_reader = WorkbookReader(Path(mapping.target_file))
                readers[target_key] = target_reader
            existing = target_reader.read(
                mapping.target_sheet,
                mapping.target_cell,
            ).value
            if existing not in (None, ""):
                warnings.append(
                    f"{Path(mapping.target_file).name}/"
                    f"{mapping.target_sheet}/{mapping.target_cell} 已有内容"
                )

            destination = (
                target_key,
                mapping.target_sheet,
                mapping.target_cell,
            )
            if destination in seen_targets:
                previous = seen_targets[destination]
                raise ValueError(
                    f"第 {previous} 项和第 {index} 项写入同一个目标单元格："
                    f"{Path(mapping.target_file).name}/"
                    f"{mapping.target_sheet}/{mapping.target_cell}"
                )
            seen_targets[destination] = index
    finally:
        for reader in readers.values():
            reader.close()
    return expanded, warnings


def output_name_for_target(path: Path) -> str:
    suffix = ".xlsx" if path.suffix.lower() == ".xls" else path.suffix
    return f"{path.stem}_已映射{suffix}"


def execute_mapping_plan(
    source_files: list[Path],
    target_files: list[Path],
    rules: list[MappingRule],
    output_folder: Path,
    progress: Callable[[int, int, str], None] | None = None,
) -> list[Path]:
    expanded, _warnings = validate_mapping_plan(
        source_files,
        target_files,
        rules,
    )
    output_folder.mkdir(parents=True, exist_ok=True)
    readers: dict[str, WorkbookReader] = {}
    target_books: dict[str, object] = {}
    output_paths: dict[str, Path] = {}
    try:
        for target_path in target_files:
            key = str(target_path.resolve())
            output_path = unique_output_path(
                output_folder / output_name_for_target(target_path)
            )
            if target_path.suffix.lower() == ".xls":
                workbook = copy_xls_to_xlsx(target_path)
            else:
                shutil.copy2(target_path, output_path)
                workbook = load_workbook(
                    output_path,
                    keep_vba=target_path.suffix.lower() == ".xlsm",
                )
            target_books[key] = workbook
            output_paths[key] = output_path

        total = len(expanded)
        for index, mapping in enumerate(expanded, start=1):
            if progress:
                progress(
                    index - 1,
                    total,
                    f"正在映射：{Path(mapping.source_file).name}/"
                    f"{mapping.source_cell} → "
                    f"{Path(mapping.target_file).name}/{mapping.target_cell}",
                )
            source_key = str(Path(mapping.source_file).resolve())
            reader = readers.get(source_key)
            if reader is None:
                reader = WorkbookReader(Path(mapping.source_file))
                readers[source_key] = reader
            cell_data = reader.read(mapping.source_sheet, mapping.source_cell)

            target_key = str(Path(mapping.target_file).resolve())
            target_book = target_books[target_key]
            if mapping.target_sheet not in target_book.sheetnames:
                raise ValueError(
                    f"{Path(mapping.target_file).name} 中不存在工作表"
                    f"“{mapping.target_sheet}”。"
                )
            target_cell = target_book[mapping.target_sheet][mapping.target_cell]
            target_cell.value = cell_data.value
            target_cell.number_format = cell_data.number_format
            if progress:
                progress(index, total, f"已完成 {index}/{total} 项映射")

        for key, workbook in target_books.items():
            workbook.save(output_paths[key])
            workbook.close()
        target_books.clear()
        return list(output_paths.values())
    finally:
        for reader in readers.values():
            reader.close()
        for workbook in target_books.values():
            workbook.close()


def save_mapping_project(
    path: Path,
    rules: list[MappingRule],
) -> None:
    text = serialize_mapping_project(rules, path.parent)
    path.write_text(text, encoding="utf-8-sig")


def serialize_mapping_project(
    rules: list[MappingRule],
    base_folder: Path,
) -> str:
    lines: list[str] = []
    for rule in rules:
        source_file = format_project_path(Path(rule.source_file), base_folder)
        target_file = format_project_path(Path(rule.target_file), base_folder)
        target_cells = (
            "同位置"
            if rule.source_cells == rule.target_cells
            else format_cell_addresses(rule.target_cells)
        )
        lines.append(
            (
                f"【{source_file}；{rule.source_sheet}；"
                f"{format_cell_addresses(rule.source_cells)} → "
                f"{target_file}；{rule.target_sheet}；{target_cells}】"
            )
        )
    return "\n".join(lines) + ("\n" if lines else "")


def load_mapping_project(
    path: Path,
) -> tuple[list[Path], list[Path], list[MappingRule], Path]:
    sources, targets, rules = parse_mapping_project_text(
        path.read_text(encoding="utf-8-sig"),
        path.parent,
    )
    return sources, targets, rules, (path.parent / "映射结果").resolve()


def parse_mapping_project_text(
    text: str,
    base_folder: Path,
) -> tuple[list[Path], list[Path], list[MappingRule]]:
    sources: list[Path] = []
    targets: list[Path] = []
    rules: list[MappingRule] = []

    for line_number, raw_line in enumerate(
        text.lstrip("\ufeff").splitlines(),
        start=1,
    ):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        content = extract_bracketed_text(line, line_number)
        direction_parts = PROJECT_DIRECTION_RE.split(content, maxsplit=1)
        if len(direction_parts) != 2:
            raise ValueError(
                f"第 {line_number} 行缺少读取和写入之间的箭头 →。"
            )
        source_fields = PROJECT_FIELD_RE.split(direction_parts[0], maxsplit=2)
        target_fields = PROJECT_FIELD_RE.split(direction_parts[1], maxsplit=2)
        if len(source_fields) != 3 or len(target_fields) != 3:
            raise ValueError(
                f"第 {line_number} 行格式错误。正确格式为："
                "【工作簿；工作表；单元格 → 工作簿；工作表；单元格】"
            )
        try:
            source_path = resolve_project_path(source_fields[0], base_folder)
            target_path = resolve_project_path(target_fields[0], base_folder)
            source_cells = parse_cell_addresses(source_fields[2])
            target_text = target_fields[2].strip()
            target_cells = (
                list(source_cells)
                if target_text in ("", "同位置", "与读取单元格相同")
                else parse_cell_addresses(target_text)
            )
            mode = infer_mapping_mode(source_cells, target_cells)
        except ValueError as exc:
            raise ValueError(f"第 {line_number} 行：{exc}") from exc
        rule = MappingRule(
            source_file=str(source_path),
            source_sheet=source_fields[1].strip(),
            source_cells=source_cells,
            target_file=str(target_path),
            target_sheet=target_fields[1].strip(),
            target_cells=target_cells,
            mode=mode,
        )
        rules.append(rule)
        if source_path not in sources:
            sources.append(source_path)
        if target_path not in targets:
            targets.append(target_path)
    if not rules:
        raise ValueError("方案文件中没有映射关系。")
    return sources, targets, rules


def resolve_project_path(text: str, base_folder: Path) -> Path:
    path = Path(text.strip())
    return path if path.is_absolute() else (base_folder / path).resolve()


def format_project_path(path: Path, base_folder: Path) -> str:
    path = path.resolve()
    base_folder = base_folder.resolve()
    if path.parent == base_folder:
        return path.name
    return str(path)


def extract_bracketed_text(text: str, line_number: int) -> str:
    text = text.strip()
    if not text or text[0] not in PROJECT_BRACKETS:
        raise ValueError(f"第 {line_number} 行需要使用括号包住内容。")
    expected = PROJECT_BRACKETS[text[0]]
    if not text.endswith(expected):
        raise ValueError(
            f"第 {line_number} 行括号不匹配，应以 {expected} 结束。"
        )
    return text[1:-1].strip()


def column_number_to_letters(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def cell_sort_key(address: str) -> tuple[int, int]:
    row, column = split_address(address)
    return row, column


def cell_column_sort_key(address: str) -> tuple[int, int]:
    row, column = split_address(address)
    return column, row


def format_cell_addresses(addresses: list[str]) -> str:
    if not addresses:
        return ""
    positions = [split_address(address) for address in addresses]
    rows = [row for row, _column in positions]
    columns = [column for _row, column in positions]
    first_row, last_row = min(rows), max(rows)
    first_column, last_column = min(columns), max(columns)
    expected = [
        (row, column)
        for column in range(first_column, last_column + 1)
        for row in range(first_row, last_row + 1)
    ]
    if positions == expected:
        start = (
            f"{column_number_to_letters(first_column + 1)}{first_row + 1}"
        )
        end = f"{column_number_to_letters(last_column + 1)}{last_row + 1}"
        return start if start == end else f"{start}-{end}"
    return ",".join(addresses)


def normalized_match_name(name: str) -> str:
    text = Path(name).stem if Path(name).suffix else name
    text = re.sub(
        r"(?:\s*[-_（(]?\s*副本\s*[）)]?)$",
        "",
        text,
        flags=re.IGNORECASE,
    )
    return re.sub(r"[\s_\-]+", "", text).casefold()


def best_name_match(source_name: str, candidates: list[str]) -> str | None:
    source = normalized_match_name(source_name)
    best: str | None = None
    best_score = 0
    for candidate in candidates:
        candidate_name = normalized_match_name(candidate)
        if not source or not candidate_name:
            continue
        if source == candidate_name:
            score = 10000 + len(source)
        else:
            common = 0
            for left, right in zip(source, candidate_name):
                if left != right:
                    break
                common += 1
            score = 1000 + common if common >= 2 else 0
        if score > best_score:
            best = candidate
            best_score = score
    return best


class CellPickerDialog:
    ROW_HEIGHT = 28
    COLUMN_WIDTH = 105
    HEADER_HEIGHT = 30
    ROW_HEADER_WIDTH = 48

    def __init__(
        self,
        parent,
        workbook_path: Path,
        sheet_name: str,
        initial_cells: list[str],
    ) -> None:
        self.result: list[str] | None = None
        self.path = workbook_path
        self.sheet_name = sheet_name
        self.selected = set(initial_cells)
        self.anchor: tuple[int, int] | None = None
        self.cell_display_cache: dict[str, tuple[object, str | None]] = {}
        self.rows = max(
            50,
            max((cell_sort_key(item)[0] + 1 for item in initial_cells), default=0),
        )
        self.columns = max(
            16,
            max((cell_sort_key(item)[1] + 1 for item in initial_cells), default=0),
        )
        self.reader = WorkbookReader(workbook_path)

        self.window = Toplevel(parent)
        self.window.title(f"选择单元格 — {workbook_path.name} / {sheet_name}")
        self.window.geometry("900x620")
        self.window.minsize(720, 500)
        self.window.transient(parent)
        self.window.grab_set()
        self.window.protocol("WM_DELETE_WINDOW", self._cancel)

        self.selected_text = StringVar()
        self.jump_text = StringVar()
        container = ttk.Frame(self.window, padding=12)
        container.pack(fill=BOTH, expand=True)

        top = ttk.Frame(container)
        top.pack(fill=X)
        ttk.Label(top, text="已选择：").pack(side=LEFT)
        ttk.Label(top, textvariable=self.selected_text).pack(
            side=LEFT,
            fill=X,
            expand=True,
        )
        ttk.Button(top, text="清空", command=self._clear).pack(side=RIGHT)
        ttk.Button(top, text="跳转", command=self._jump).pack(
            side=RIGHT,
            padx=(6, 0),
        )
        jump_entry = ttk.Entry(top, textvariable=self.jump_text, width=9)
        jump_entry.pack(side=RIGHT)
        ttk.Label(top, text="定位到：").pack(side=RIGHT, padx=(10, 4))
        jump_entry.bind("<Return>", lambda _event: self._jump())

        ttk.Label(
            container,
            text="单击选择或取消；按住鼠标拖动可选择连续区域。",
        ).pack(anchor=W, pady=(6, 8))

        table = ttk.Frame(container)
        table.pack(fill=BOTH, expand=True)
        table.rowconfigure(1, weight=1)
        table.columnconfigure(1, weight=1)
        self.corner = Canvas(
            table,
            width=self.ROW_HEADER_WIDTH,
            height=self.HEADER_HEIGHT,
            highlightthickness=1,
        )
        self.corner.grid(row=0, column=0, sticky="nsew")
        self.column_header = Canvas(
            table,
            height=self.HEADER_HEIGHT,
            highlightthickness=1,
        )
        self.column_header.grid(row=0, column=1, sticky="ew")
        self.row_header = Canvas(
            table,
            width=self.ROW_HEADER_WIDTH,
            highlightthickness=1,
        )
        self.row_header.grid(row=1, column=0, sticky="ns")
        self.grid_canvas = Canvas(table, background="white", highlightthickness=1)
        self.grid_canvas.grid(row=1, column=1, sticky="nsew")

        horizontal = ttk.Scrollbar(
            table,
            orient="horizontal",
            command=self._xview,
        )
        horizontal.grid(row=2, column=1, sticky="ew")
        vertical = ttk.Scrollbar(
            table,
            orient="vertical",
            command=self._yview,
        )
        vertical.grid(row=1, column=2, sticky="ns")
        self.grid_canvas.configure(
            xscrollcommand=horizontal.set,
            yscrollcommand=vertical.set,
        )
        self.grid_canvas.bind("<Button-1>", self._press)
        self.grid_canvas.bind("<ButtonRelease-1>", self._release)
        self.grid_canvas.bind("<MouseWheel>", self._mousewheel)
        self.grid_canvas.bind("<Shift-MouseWheel>", self._shift_mousewheel)

        actions = ttk.Frame(container)
        actions.pack(fill=X, pady=(10, 0))
        ttk.Button(actions, text="确定选择", command=self._confirm).pack(
            side=LEFT,
        )
        ttk.Button(actions, text="取消", command=self._cancel).pack(
            side=LEFT,
            padx=(8, 0),
        )

        self._draw()
        self._update_selected_text()
        parent.wait_window(self.window)

    def _draw(self) -> None:
        self.column_header.delete("all")
        self.row_header.delete("all")
        self.grid_canvas.delete("all")
        width = self.columns * self.COLUMN_WIDTH
        height = self.rows * self.ROW_HEIGHT
        self.column_header.configure(scrollregion=(0, 0, width, self.HEADER_HEIGHT))
        self.row_header.configure(scrollregion=(0, 0, self.ROW_HEADER_WIDTH, height))
        self.grid_canvas.configure(scrollregion=(0, 0, width, height))

        for column in range(self.columns):
            x1 = column * self.COLUMN_WIDTH
            x2 = x1 + self.COLUMN_WIDTH
            self.column_header.create_rectangle(
                x1,
                0,
                x2,
                self.HEADER_HEIGHT,
                fill="#edf2f7",
                outline="#b8c2cc",
            )
            self.column_header.create_text(
                (x1 + x2) / 2,
                self.HEADER_HEIGHT / 2,
                text=column_number_to_letters(column + 1),
            )
        for row in range(self.rows):
            y1 = row * self.ROW_HEIGHT
            y2 = y1 + self.ROW_HEIGHT
            self.row_header.create_rectangle(
                0,
                y1,
                self.ROW_HEADER_WIDTH,
                y2,
                fill="#edf2f7",
                outline="#b8c2cc",
            )
            self.row_header.create_text(
                self.ROW_HEADER_WIDTH / 2,
                (y1 + y2) / 2,
                text=str(row + 1),
            )
            for column in range(self.columns):
                self._draw_cell(row, column)

    def _draw_cell(self, row: int, column: int) -> None:
        address = f"{column_number_to_letters(column + 1)}{row + 1}"
        x1 = column * self.COLUMN_WIDTH
        y1 = row * self.ROW_HEIGHT
        if address not in self.cell_display_cache:
            try:
                value = self.reader.read(self.sheet_name, address).value
                original_fill = self.reader.fill_color(
                    self.sheet_name,
                    address,
                )
            except Exception:
                value = ""
                original_fill = None
            self.cell_display_cache[address] = (value, original_fill)
        value, original_fill = self.cell_display_cache[address]
        selected = address in self.selected
        self.grid_canvas.create_rectangle(
            x1,
            y1,
            x1 + self.COLUMN_WIDTH,
            y1 + self.ROW_HEIGHT,
            fill=original_fill or "white",
            outline="#1683e2" if selected else "#d6dce2",
            width=3 if selected else 1,
            tags=(f"cell-{address}",),
        )
        text = "" if value is None else str(value).replace("\n", " ")
        if len(text) > 14:
            text = text[:13] + "…"
        text_color = "black"
        if original_fill:
            red, green, blue = (
                int(original_fill[index:index + 2], 16)
                for index in (1, 3, 5)
            )
            if red * 299 + green * 587 + blue * 114 < 128000:
                text_color = "white"
        self.grid_canvas.create_text(
            x1 + 5,
            y1 + self.ROW_HEIGHT / 2,
            text=text,
            fill=text_color,
            anchor=W,
            tags=(f"cell-{address}",),
        )

    def _refresh_cell(self, row: int, column: int) -> None:
        address = f"{column_number_to_letters(column + 1)}{row + 1}"
        self.grid_canvas.delete(f"cell-{address}")
        self._draw_cell(row, column)

    def _event_cell(self, event) -> tuple[int, int]:
        x = self.grid_canvas.canvasx(event.x)
        y = self.grid_canvas.canvasy(event.y)
        column = max(0, min(self.columns - 1, int(x // self.COLUMN_WIDTH)))
        row = max(0, min(self.rows - 1, int(y // self.ROW_HEIGHT)))
        return row, column

    def _press(self, event) -> None:
        self.anchor = self._event_cell(event)

    def _release(self, event) -> None:
        if self.anchor is None:
            return
        end_row, end_column = self._event_cell(event)
        start_row, start_column = self.anchor
        cells = [
            (row, column)
            for row in range(min(start_row, end_row), max(start_row, end_row) + 1)
            for column in range(
                min(start_column, end_column),
                max(start_column, end_column) + 1,
            )
        ]
        addresses = [
            f"{column_number_to_letters(column + 1)}{row + 1}"
            for row, column in cells
        ]
        if len(addresses) == 1 and addresses[0] in self.selected:
            self.selected.remove(addresses[0])
        else:
            self.selected.update(addresses)
        for row, column in cells:
            self._refresh_cell(row, column)
        self.anchor = None
        self._update_selected_text()

    def _xview(self, *args) -> None:
        self.grid_canvas.xview(*args)
        self.column_header.xview(*args)

    def _yview(self, *args) -> None:
        self.grid_canvas.yview(*args)
        self.row_header.yview(*args)

    def _mousewheel(self, event) -> None:
        units = -1 if event.delta > 0 else 1
        self._yview("scroll", units * 3, "units")

    def _shift_mousewheel(self, event) -> None:
        units = -1 if event.delta > 0 else 1
        self._xview("scroll", units * 3, "units")

    def _jump(self) -> None:
        try:
            address = parse_cell_addresses(self.jump_text.get())[0]
            row, column = split_address(address)
        except ValueError as exc:
            messagebox.showerror("无法跳转", str(exc), parent=self.window)
            return
        needs_redraw = row >= self.rows or column >= self.columns
        self.rows = max(self.rows, row + 20)
        self.columns = max(self.columns, column + 8)
        if needs_redraw:
            self._draw()
        width = max(self.columns * self.COLUMN_WIDTH, 1)
        height = max(self.rows * self.ROW_HEIGHT, 1)
        self._xview("moveto", max(0, column * self.COLUMN_WIDTH / width))
        self._yview("moveto", max(0, row * self.ROW_HEIGHT / height))

    def _clear(self) -> None:
        previous = list(self.selected)
        self.selected.clear()
        for address in previous:
            row, column = cell_sort_key(address)
            if row < self.rows and column < self.columns:
                self._refresh_cell(row, column)
        self._update_selected_text()

    def _update_selected_text(self) -> None:
        ordered = sorted(self.selected, key=cell_column_sort_key)
        if not ordered:
            text = "尚未选择"
        elif len(ordered) <= 50:
            text = format_cell_addresses(ordered)
        else:
            text = (
                format_cell_addresses(ordered[:12])
                + f"……（共 {len(ordered)} 个）"
            )
        self.selected_text.set(text)

    def _confirm(self) -> None:
        if not self.selected:
            messagebox.showinfo(
                "尚未选择",
                "请至少选择一个单元格。",
                parent=self.window,
            )
            return
        self.result = sorted(self.selected, key=cell_column_sort_key)
        self.reader.close()
        self.window.destroy()

    def _cancel(self) -> None:
        self.reader.close()
        self.window.destroy()


class MappingDialog:
    def __init__(
        self,
        parent,
        source_files: list[Path],
        target_files: list[Path],
        initial: MappingRule | None = None,
        defaults: MappingRule | None = None,
    ) -> None:
        self.result: MappingRule | None = None
        self.source_files = source_files
        self.target_files = target_files
        self.window = Toplevel(parent)
        self.window.title("设置映射关系")
        self.window.geometry("760x410")
        self.window.resizable(False, False)
        self.window.transient(parent)
        self.window.grab_set()

        base = initial or defaults
        self.source_file = StringVar(
            value=base.source_file if base else str(source_files[0])
        )
        self.source_sheet = StringVar(
            value=base.source_sheet if base else ""
        )
        self.source_cells = StringVar(
            value=format_cell_addresses(initial.source_cells) if initial else ""
        )
        self.target_file = StringVar(
            value=base.target_file if base else str(target_files[0])
        )
        self.target_sheet = StringVar(
            value=base.target_sheet if base else ""
        )
        self.target_cells = StringVar(
            value=format_cell_addresses(initial.target_cells) if initial else ""
        )
        self._sync_target_cells = initial is None
        self._updating_target_cells = False
        self._target_file_manually_selected = initial is not None
        self._target_sheet_manually_selected = initial is not None
        container = ttk.Frame(self.window, padding=16)
        container.pack(fill=BOTH, expand=True)
        (
            self.source_file_combo,
            self.source_sheet_combo,
            self.source_cells_entry,
        ) = self._location_group(
            container,
            "读取位置",
            self.source_file,
            self.source_sheet,
            self.source_cells,
            source_files,
            0,
            True,
        )
        (
            self.target_file_combo,
            self.target_sheet_combo,
            self.target_cells_entry,
        ) = self._location_group(
            container,
            "写入位置",
            self.target_file,
            self.target_sheet,
            self.target_cells,
            target_files,
            1,
            False,
        )

        actions = ttk.Frame(container)
        actions.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(14, 0))
        ttk.Button(actions, text="保存映射", command=self._save).pack(side=LEFT)
        ttk.Button(
            actions,
            text="取消",
            command=self.window.destroy,
        ).pack(side=LEFT, padx=(8, 0))
        container.columnconfigure(0, weight=1)
        container.columnconfigure(1, weight=1)

        self._refresh_source_sheets()
        if initial is None:
            self._match_target_file()
        self._refresh_target_sheets()
        if initial is None:
            self._match_target_sheet()
        self.source_cells.trace_add("write", self._copy_source_cells_to_target)
        self.target_cells.trace_add("write", self._target_cells_changed)
        parent.wait_window(self.window)

    def _location_group(
        self,
        parent,
        title: str,
        file_var: StringVar,
        sheet_var: StringVar,
        cells_var: StringVar,
        files: list[Path],
        column: int,
        is_source: bool,
    ):
        frame = ttk.LabelFrame(parent, text=title, padding=10)
        frame.grid(
            row=0,
            column=column,
            sticky="nsew",
            padx=(0, 5) if column == 0 else (5, 0),
        )
        ttk.Label(frame, text="工作簿：").pack(anchor=W)
        file_combo = ttk.Combobox(
            frame,
            textvariable=file_var,
            values=[str(path) for path in files],
            state="readonly",
        )
        file_combo.pack(fill=X, pady=(2, 8))
        ttk.Label(frame, text="工作表：").pack(anchor=W)
        sheet_combo = ttk.Combobox(
            frame,
            textvariable=sheet_var,
            state="readonly",
        )
        sheet_combo.pack(fill=X, pady=(2, 8))
        cells_title = ttk.Frame(frame)
        cells_title.pack(fill=X)
        ttk.Label(
            cells_title,
            text="读取单元格：" if is_source else "写入单元格：",
        ).pack(side=LEFT)
        ttk.Button(
            cells_title,
            text="选择单元格…",
            command=lambda: self._pick_cells(is_source),
        ).pack(side=LEFT, padx=(6, 0))
        cells_row = ttk.Frame(frame)
        cells_row.pack(fill=X, pady=(2, 0))
        cells_entry = Text(
            cells_row,
            height=3,
            wrap="word",
            font=("TkDefaultFont", 10),
        )
        cells_entry.pack(side=LEFT, fill=BOTH, expand=True)
        cells_scroll = ttk.Scrollbar(
            cells_row,
            orient="vertical",
            command=cells_entry.yview,
        )
        cells_scroll.pack(side=LEFT, fill="y")
        cells_entry.configure(yscrollcommand=cells_scroll.set)
        self._bind_text_to_variable(cells_entry, cells_var)
        if column == 0:
            file_combo.bind(
                "<<ComboboxSelected>>",
                lambda _event: self._source_file_selected(),
            )
            sheet_combo.bind(
                "<<ComboboxSelected>>",
                lambda _event: self._source_sheet_selected(),
            )
        else:
            file_combo.bind(
                "<<ComboboxSelected>>",
                lambda _event: self._target_file_selected(),
            )
            sheet_combo.bind(
                "<<ComboboxSelected>>",
                lambda _event: self._target_sheet_selected(),
            )
        return file_combo, sheet_combo, cells_entry

    @staticmethod
    def _bind_text_to_variable(text_widget: Text, variable: StringVar) -> None:
        changing = {"value": False}

        def variable_changed(*_args) -> None:
            if changing["value"]:
                return
            value = variable.get()
            current = text_widget.get("1.0", "end-1c")
            if current == value:
                return
            changing["value"] = True
            try:
                text_widget.delete("1.0", END)
                text_widget.insert("1.0", value)
            finally:
                changing["value"] = False

        def text_changed(_event=None) -> None:
            if changing["value"]:
                return
            changing["value"] = True
            try:
                variable.set(text_widget.get("1.0", "end-1c"))
            finally:
                changing["value"] = False

        variable_changed()
        variable.trace_add("write", variable_changed)
        text_widget.bind("<KeyRelease>", text_changed)
        text_widget.bind("<FocusOut>", text_changed)

    def _pick_cells(self, is_source: bool) -> None:
        file_var = self.source_file if is_source else self.target_file
        sheet_var = self.source_sheet if is_source else self.target_sheet
        cells_var = self.source_cells if is_source else self.target_cells
        try:
            initial = (
                parse_cell_addresses(cells_var.get())
                if cells_var.get().strip()
                else []
            )
            dialog = CellPickerDialog(
                self.window,
                Path(file_var.get()),
                sheet_var.get(),
                initial,
            )
            if dialog.result:
                cells_var.set(", ".join(dialog.result))
        except Exception as exc:
            messagebox.showerror(
                "无法打开单元格选择器",
                str(exc),
                parent=self.window,
            )

    def _source_file_selected(self) -> None:
        self._refresh_source_sheets()
        if not self._target_file_manually_selected:
            self._match_target_file()
            self._target_sheet_manually_selected = False
            self._refresh_target_sheets()
        self._source_sheet_selected()

    def _source_sheet_selected(self) -> None:
        if not self._target_sheet_manually_selected:
            self._match_target_sheet()

    def _target_file_selected(self) -> None:
        self._target_file_manually_selected = True
        self._target_sheet_manually_selected = False
        self._refresh_target_sheets()
        self._match_target_sheet()

    def _target_sheet_selected(self) -> None:
        self._target_sheet_manually_selected = True

    def _match_target_file(self) -> None:
        source_name = Path(self.source_file.get()).name
        target_paths = list(self.target_file_combo.cget("values"))
        target_names = [Path(item).name for item in target_paths]
        matched_name = best_name_match(source_name, target_names)
        if matched_name:
            for path in target_paths:
                if Path(path).name == matched_name:
                    self.target_file.set(path)
                    break

    def _match_target_sheet(self) -> None:
        source_name = self.source_sheet.get()
        target_names = list(self.target_sheet_combo.cget("values"))
        matched_name = best_name_match(source_name, target_names)
        if matched_name:
            self.target_sheet.set(matched_name)

    def _copy_source_cells_to_target(self, *_args) -> None:
        if not self._sync_target_cells:
            return
        self._updating_target_cells = True
        try:
            self.target_cells.set(self.source_cells.get())
        finally:
            self._updating_target_cells = False

    def _target_cells_changed(self, *_args) -> None:
        if not self._updating_target_cells:
            self._sync_target_cells = False

    def _refresh_source_sheets(self) -> None:
        self._refresh_sheets(
            Path(self.source_file.get()),
            self.source_sheet,
            self.source_sheet_combo,
        )

    def _refresh_target_sheets(self) -> None:
        self._refresh_sheets(
            Path(self.target_file.get()),
            self.target_sheet,
            self.target_sheet_combo,
        )

    @staticmethod
    def _refresh_sheets(path: Path, variable: StringVar, combo) -> None:
        try:
            names = workbook_sheet_names(path)
            combo.configure(values=names)
            if variable.get() not in names:
                variable.set(names[0] if names else "")
        except Exception as exc:
            messagebox.showerror("无法读取工作表", str(exc), parent=combo)

    def _save(self) -> None:
        try:
            source_cells = parse_cell_addresses(self.source_cells.get())
            target_cells = parse_cell_addresses(self.target_cells.get())
            rule = MappingRule(
                source_file=self.source_file.get(),
                source_sheet=self.source_sheet.get(),
                source_cells=source_cells,
                target_file=self.target_file.get(),
                target_sheet=self.target_sheet.get(),
                target_cells=target_cells,
                mode=infer_mapping_mode(source_cells, target_cells),
            )
            expand_rule(rule)
            self.result = rule
            self.window.destroy()
        except ValueError as exc:
            messagebox.showerror("映射设置错误", str(exc), parent=self.window)


class MapperApp:
    def __init__(self, root: Tk) -> None:
        self.root = root
        self.root.title("Excel 单元格映射工具")
        self.root.geometry("1180x800")
        self.root.minsize(900, 600)
        self.source_files: list[Path] = []
        self.target_files: list[Path] = []
        self.rules: list[MappingRule] = []
        self.output_folder = StringVar()
        self.status = StringVar(value="请添加来源工作簿和目标工作簿。")
        self.progress_text = StringVar(value="0%")
        self.txt_base_folder = StringVar()
        self.txt_status = StringVar(value="尚未生成或导入 TXT 方案。")
        self._active_workflow_tab = 0
        self._switching_workflow_tab = False
        self._updating_txt_editor = False
        self._txt_dirty = False
        self._configure_styles()
        self._build_ui()

    def _configure_styles(self) -> None:
        style = ttk.Style(self.root)
        style.configure("Toolbar.TButton", padding=(10, 4))
        style.configure(
            "Primary.TButton",
            padding=(18, 7),
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        style.configure(
            "TLabelframe.Label",
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        style.configure("TNotebook.Tab", padding=(12, 5))
        style.configure(
            "Treeview",
            rowheight=24,
            font=("Microsoft YaHei UI", 9),
        )
        style.configure(
            "Treeview.Heading",
            font=("Microsoft YaHei UI", 9, "bold"),
        )
        style.configure(
            "Section.TLabel",
            font=("Microsoft YaHei UI", 10, "bold"),
        )
        style.configure(
            "Muted.TLabel",
            foreground="#5f6368",
        )

    def _build_ui(self) -> None:
        container = ttk.Frame(self.root, padding=16)
        container.pack(fill=BOTH, expand=True)

        fixed_bottom = ttk.Frame(container)
        fixed_bottom.pack(side="bottom", fill=X)
        content = ttk.Frame(container)
        content.pack(fill=BOTH, expand=True)

        self.workflow_notebook = ttk.Notebook(content)
        self.workflow_notebook.pack(fill=BOTH, expand=True)
        manual_tab = ttk.Frame(self.workflow_notebook, padding=10)
        txt_tab = ttk.Frame(self.workflow_notebook, padding=12)
        self.workflow_notebook.add(manual_tab, text="手动设置")
        self.workflow_notebook.add(txt_tab, text="TXT 方案")
        self.workflow_notebook.bind(
            "<<NotebookTabChanged>>",
            self._workflow_tab_changed,
        )
        self.clear_all_button = ttk.Button(
            content,
            text="清空所有内容",
            command=self._clear_all_workflow,
            style="Toolbar.TButton",
        )
        self.clear_all_button.place(
            relx=1.0,
            x=-8,
            y=3,
            anchor="ne",
        )
        self.clear_all_button.lift()

        self.manual_panes = ttk.Panedwindow(manual_tab, orient="vertical")
        self.manual_panes.pack(fill=BOTH, expand=True)

        file_row = ttk.Frame(self.manual_panes)
        self.manual_panes.add(file_row, weight=1)
        file_row.columnconfigure(0, weight=1, uniform="files")
        file_row.columnconfigure(2, weight=1, uniform="files")
        self.source_tree = self._file_panel(
            file_row,
            "1. 来源工作簿",
            0,
            self._add_sources,
            self._add_source_folder,
            self._remove_sources,
        )
        transfer = ttk.Frame(file_row, padding=(8, 42))
        transfer.grid(row=0, column=1, sticky="ns")
        ttk.Button(
            transfer,
            text="移到目标 →",
            command=self._move_sources_to_targets,
        ).pack(fill=X, pady=(0, 8))
        ttk.Button(
            transfer,
            text="← 移到来源",
            command=self._move_targets_to_sources,
        ).pack(fill=X)
        self.target_tree = self._file_panel(
            file_row,
            "2. 目标工作簿",
            2,
            self._add_targets,
            self._add_target_folder,
            self._remove_targets,
        )

        mapping_frame = ttk.LabelFrame(
            self.manual_panes,
            text="3. 映射关系",
            padding=10,
        )
        self.manual_panes.add(mapping_frame, weight=2)
        mapping_frame.columnconfigure(0, weight=1)
        mapping_frame.rowconfigure(0, weight=1)

        mapping_table = ttk.Frame(mapping_frame)
        mapping_table.grid(row=0, column=0, sticky="nsew")
        mapping_table.columnconfigure(0, weight=1)
        mapping_table.rowconfigure(0, weight=1)
        self.mapping_tree = ttk.Treeview(
            mapping_table,
            columns=("source", "target"),
            show="headings",
            height=6,
        )
        self.mapping_tree.heading("source", text="来源位置")
        self.mapping_tree.heading("target", text="目标位置")
        self.mapping_tree.column("source", width=450)
        self.mapping_tree.column("target", width=450)
        self.mapping_tree.grid(row=0, column=0, sticky="nsew")
        mapping_vertical = ttk.Scrollbar(
            mapping_table,
            orient="vertical",
            command=self.mapping_tree.yview,
        )
        mapping_vertical.grid(row=0, column=1, sticky="ns")
        mapping_horizontal = ttk.Scrollbar(
            mapping_table,
            orient="horizontal",
            command=self.mapping_tree.xview,
        )
        mapping_horizontal.grid(row=1, column=0, sticky="ew")
        self.mapping_tree.configure(
            yscrollcommand=mapping_vertical.set,
            xscrollcommand=mapping_horizontal.set,
        )
        self.mapping_tree.bind(
            "<Double-1>",
            lambda _event: self._edit_rule(),
        )
        self.mapping_actions = ttk.Frame(mapping_frame)
        self.mapping_actions.grid(
            row=1,
            column=0,
            sticky="ew",
            pady=(8, 0),
        )
        for text, command in (
            ("添加映射", self._add_rule),
            ("编辑选中", self._edit_rule),
            ("复制选中", self._duplicate_rule),
            ("删除选中", self._delete_rule),
            ("清空映射", self._clear_rules),
        ):
            ttk.Button(
                self.mapping_actions,
                text=text,
                command=command,
            ).pack(side=LEFT, padx=(0, 6))
        ttk.Button(
            self.mapping_actions,
            text="预检查",
            command=self._precheck,
        ).pack(side=RIGHT)
        ttk.Button(
            self.mapping_actions,
            text="展开预览",
            command=self._preview_expanded,
        ).pack(side=RIGHT, padx=(0, 6))
        plan_actions = ttk.Frame(txt_tab)
        plan_actions.pack(fill=X, pady=(0, 10))
        ttk.Label(
            plan_actions,
            text="方案操作：",
            style="Section.TLabel",
        ).pack(side=LEFT, padx=(0, 8))
        for text, command in (
            ("导入 TXT 方案", self._load_project),
            ("保存 TXT 方案", self._save_project),
        ):
            ttk.Button(
                plan_actions,
                text=text,
                command=command,
                style="Toolbar.TButton",
            ).pack(side=LEFT, padx=(0, 6))

        base_group = ttk.LabelFrame(
            txt_tab,
            text="路径设置",
            padding=8,
        )
        base_group.pack(fill=X, pady=(0, 10))
        base_row = ttk.Frame(base_group)
        base_row.pack(fill=X)
        ttk.Label(base_row, text="方案基准文件夹：").pack(side=LEFT)
        ttk.Entry(
            base_row,
            textvariable=self.txt_base_folder,
        ).pack(side=LEFT, fill=X, expand=True)
        ttk.Button(
            base_row,
            text="选择文件夹…",
            command=self._browse_txt_base,
        ).pack(side=LEFT, padx=(8, 0))

        content_panes = ttk.Panedwindow(txt_tab, orient="horizontal")
        content_panes.pack(fill=BOTH, expand=True)

        editor_group = ttk.LabelFrame(
            content_panes,
            text="方案内容",
            padding=8,
        )
        example_group = ttk.LabelFrame(
            content_panes,
            text="格式示例与填写说明",
            padding=8,
        )
        content_panes.add(editor_group, weight=3)
        content_panes.add(example_group, weight=2)

        txt_editor_frame = ttk.Frame(editor_group)
        txt_editor_frame.pack(fill=BOTH, expand=True, pady=(3, 0))
        self.txt_editor = Text(
            txt_editor_frame,
            wrap="none",
            font=("TkFixedFont", 10),
            undo=True,
            width=72,
        )
        self.txt_editor.pack(side=LEFT, fill=BOTH, expand=True)
        txt_vertical = ttk.Scrollbar(
            txt_editor_frame,
            orient="vertical",
            command=self.txt_editor.yview,
        )
        txt_vertical.pack(side=RIGHT, fill="y")
        txt_horizontal = ttk.Scrollbar(
            editor_group,
            orient="horizontal",
            command=self.txt_editor.xview,
        )
        txt_horizontal.pack(fill=X)
        self.txt_editor.configure(
            yscrollcommand=txt_vertical.set,
            xscrollcommand=txt_horizontal.set,
        )
        self.txt_editor.bind("<<Modified>>", self._txt_editor_modified)

        example_header = ttk.Frame(example_group)
        example_header.pack(fill=X, pady=(0, 5))
        ttk.Label(
            example_header,
            text="可直接复制后修改：",
            style="Muted.TLabel",
        ).pack(side=LEFT)
        ttk.Button(
            example_header,
            text="复制示例",
            command=self._copy_txt_example,
        ).pack(side=RIGHT)
        self.txt_example_view = Text(
            example_group,
            wrap="char",
            font=("TkFixedFont", 9),
            height=8,
            width=42,
        )
        self.txt_example_view.insert("1.0", TXT_EXAMPLE)
        self.txt_example_view.configure(state="disabled")
        self.txt_example_view.pack(fill=X)
        ttk.Separator(example_group).pack(fill=X, pady=8)
        instructions = Text(
            example_group,
            wrap="word",
            font=("TkDefaultFont", 9),
            height=8,
            width=42,
        )
        instructions.insert("1.0", TXT_INSTRUCTIONS)
        instructions.configure(state="disabled")
        instructions.pack(fill=BOTH, expand=True)

        ttk.Label(
            txt_tab,
            textvariable=self.txt_status,
            style="Muted.TLabel",
        ).pack(anchor=W, pady=(8, 0))

        output_frame = ttk.LabelFrame(
            fixed_bottom,
            text="4. 输出与执行",
            padding=10,
        )
        output_frame.pack(fill=X, pady=(10, 0))
        output_content = ttk.Frame(output_frame)
        output_content.pack(fill=X, expand=True)

        location_row = ttk.Frame(output_content)
        location_row.pack(fill=X)
        ttk.Label(location_row, text="副本输出文件夹：").pack(side=LEFT)
        ttk.Entry(
            location_row,
            textvariable=self.output_folder,
        ).pack(side=LEFT, fill=X, expand=True)
        ttk.Button(
            location_row,
            text="浏览...",
            command=self._browse_output,
        ).pack(side=LEFT, padx=(8, 0))

        progress_row = ttk.Frame(output_content)
        progress_row.pack(fill=X, pady=(8, 0))
        self.progress = ttk.Progressbar(progress_row, mode="determinate")
        self.progress.pack(side=LEFT, fill=X, expand=True)
        ttk.Label(
            progress_row,
            textvariable=self.progress_text,
        ).pack(side=RIGHT, padx=(8, 0))

        status_actions = ttk.Frame(output_content)
        status_actions.pack(fill=X, pady=(7, 0))
        ttk.Label(status_actions, textvariable=self.status).pack(
            side=LEFT,
            anchor=W,
            fill=X,
            expand=True,
        )
        ttk.Button(
            status_actions,
            text="退出",
            command=self.root.destroy,
            width=12,
        ).pack(side=RIGHT)
        ttk.Button(
            status_actions,
            text="开始映射",
            command=self._run,
            width=12,
            style="Primary.TButton",
        ).pack(side=RIGHT, padx=(0, 8))

    def _file_panel(
        self,
        parent,
        title: str,
        column: int,
        add_command,
        folder_command,
        remove_command,
    ):
        frame = ttk.LabelFrame(parent, text=title, padding=10)
        frame.grid(
            row=0,
            column=column,
            sticky="nsew",
            padx=(0, 5) if column == 0 else (5, 0),
        )
        list_area = ttk.Frame(frame)
        list_area.pack(fill=BOTH, expand=True)
        list_area.columnconfigure(0, weight=1)
        list_area.rowconfigure(0, weight=1)
        tree = ttk.Treeview(
            list_area,
            columns=("path",),
            show="headings",
            height=4,
        )
        tree.heading("path", text="文件路径")
        tree.column("path", width=800, minwidth=430, stretch=False)
        tree.grid(row=0, column=0, sticky="nsew")
        vertical = ttk.Scrollbar(
            list_area,
            orient="vertical",
            command=tree.yview,
        )
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal = ttk.Scrollbar(
            list_area,
            orient="horizontal",
            command=tree.xview,
        )
        horizontal.grid(row=1, column=0, sticky="ew")
        tree.configure(
            yscrollcommand=vertical.set,
            xscrollcommand=horizontal.set,
        )
        actions = ttk.Frame(frame)
        actions.pack(fill=X, pady=(6, 0))
        ttk.Button(actions, text="添加文件", command=add_command).pack(side=LEFT)
        ttk.Button(
            actions,
            text="添加文件夹",
            command=folder_command,
        ).pack(side=LEFT, padx=(6, 0))
        ttk.Button(
            actions,
            text="删除选中",
            command=remove_command,
        ).pack(side=LEFT, padx=(6, 0))
        return tree

    @staticmethod
    def _choose_workbooks() -> list[Path]:
        selected = filedialog.askopenfilenames(
            title="选择 Excel 工作簿",
            filetypes=[
                ("Excel 文件", "*.xls *.xlsx *.xlsm"),
                ("所有文件", "*.*"),
            ],
        )
        return [Path(item) for item in selected]

    def _add_sources(self) -> None:
        self._add_files(self.source_files, self.source_tree)

    def _add_targets(self) -> None:
        self._add_files(self.target_files, self.target_tree)

    def _add_source_folder(self) -> None:
        self._add_folder(self.source_files, self.source_tree)

    def _add_target_folder(self) -> None:
        self._add_folder(self.target_files, self.target_tree)

    def _add_files(self, collection: list[Path], tree) -> None:
        self._append_files(collection, tree, self._choose_workbooks())

    def _add_folder(self, collection: list[Path], tree) -> None:
        selected = filedialog.askdirectory(title="选择包含 Excel 文件的文件夹")
        if not selected:
            return
        files = workbook_files_in_folder(Path(selected))
        if not files:
            messagebox.showinfo(
                "没有可添加的文件",
                "所选文件夹中没有 .xls、.xlsx 或 .xlsm 文件。",
                parent=self.root,
            )
            return
        if not messagebox.askyesno(
            "确认批量添加",
            f"发现 {len(files)} 个 Excel 文件，是否全部添加？",
            parent=self.root,
        ):
            return
        self._append_files(collection, tree, files)

    def _append_files(self, collection: list[Path], tree, files) -> None:
        for path in files:
            if path not in collection:
                collection.append(path)
                tree.insert("", END, values=(str(path),))
        if not self.output_folder.get() and self.target_files:
            self.output_folder.set(str(self.target_files[0].parent / "映射结果"))

    def _remove_sources(self) -> None:
        self._remove_files(self.source_files, self.source_tree)

    def _remove_targets(self) -> None:
        self._remove_files(self.target_files, self.target_tree)

    def _move_sources_to_targets(self) -> None:
        self._move_selected_files(
            self.source_files,
            self.source_tree,
            self.target_files,
            self.target_tree,
            "source",
        )

    def _move_targets_to_sources(self) -> None:
        self._move_selected_files(
            self.target_files,
            self.target_tree,
            self.source_files,
            self.source_tree,
            "target",
        )

    def _move_selected_files(
        self,
        source_collection: list[Path],
        source_tree,
        target_collection: list[Path],
        target_tree,
        current_role: str,
    ) -> None:
        selected_items = source_tree.selection()
        if not selected_items:
            messagebox.showinfo(
                "请先选择文件",
                "请在列表中选中需要转移的工作簿。",
                parent=self.root,
            )
            return
        selected_paths = {
            Path(source_tree.item(item, "values")[0])
            for item in selected_items
        }
        used_paths = {
            Path(
                rule.source_file
                if current_role == "source"
                else rule.target_file
            )
            for rule in self.rules
        }
        blocked = selected_paths & used_paths
        if blocked:
            messagebox.showinfo(
                "文件正在映射中",
                "以下工作簿已经用于映射，请先删除或修改相关映射：\n"
                + "\n".join(path.name for path in sorted(blocked)),
                parent=self.root,
            )
            return
        for path in selected_paths:
            if path not in target_collection:
                target_collection.append(path)
        source_collection[:] = [
            path for path in source_collection if path not in selected_paths
        ]
        self._refresh_file_tree(source_tree, source_collection)
        self._refresh_file_tree(target_tree, target_collection)
        if not self.output_folder.get() and self.target_files:
            self.output_folder.set(str(self.target_files[0].parent / "映射结果"))

    @staticmethod
    def _remove_files(collection: list[Path], tree) -> None:
        selected = tree.selection()
        paths = {Path(tree.item(item, "values")[0]) for item in selected}
        collection[:] = [path for path in collection if path not in paths]
        for item in selected:
            tree.delete(item)

    def _selected_rule_index(self) -> int | None:
        selection = self.mapping_tree.selection()
        if not selection:
            return None
        return self.mapping_tree.index(selection[0])

    def _add_rule(self) -> None:
        if not self.source_files or not self.target_files:
            messagebox.showinfo(
                "请先添加文件",
                "请先添加至少一个来源工作簿和一个目标工作簿。",
                parent=self.root,
            )
            return
        dialog = MappingDialog(
            self.root,
            self.source_files,
            self.target_files,
            defaults=self.rules[-1] if self.rules else None,
        )
        if dialog.result:
            self.rules.append(dialog.result)
            self._refresh_rules()

    def _edit_rule(self) -> None:
        index = self._selected_rule_index()
        if index is None:
            return
        dialog = MappingDialog(
            self.root,
            self.source_files,
            self.target_files,
            self.rules[index],
        )
        if dialog.result:
            self.rules[index] = dialog.result
            self._refresh_rules()

    def _duplicate_rule(self) -> None:
        index = self._selected_rule_index()
        if index is None:
            return
        self.rules.insert(index + 1, MappingRule(**asdict(self.rules[index])))
        self._refresh_rules()

    def _delete_rule(self) -> None:
        index = self._selected_rule_index()
        if index is not None:
            self.rules.pop(index)
            self._refresh_rules()

    def _clear_rules(self) -> None:
        if self.rules and messagebox.askyesno(
            "清空映射",
            "确定清空全部映射规则吗？",
            parent=self.root,
        ):
            self.rules.clear()
            self._refresh_rules()

    def _clear_all_workflow(self) -> None:
        has_txt_content = (
            hasattr(self, "txt_editor")
            and bool(self.txt_editor.get("1.0", END).strip())
        )
        if not (
            self.source_files
            or self.target_files
            or self.rules
            or has_txt_content
            or self.txt_base_folder.get().strip()
        ):
            self.status.set("当前方案中没有需要清空的内容。")
            return
        if not messagebox.askyesno(
            "清空所有内容",
            "确定清空手动设置和 TXT 方案中的全部内容吗？",
            parent=self.root,
        ):
            return
        self.source_files.clear()
        self.target_files.clear()
        self.rules.clear()
        self._refresh_file_tree(self.source_tree, self.source_files)
        self._refresh_file_tree(self.target_tree, self.target_files)
        self._refresh_rules()
        self._set_txt_text("")
        self.txt_base_folder.set("")
        self._txt_dirty = False
        self.txt_status.set("TXT 方案内容已清空。")
        self.status.set("已清空手动设置和 TXT 方案中的全部内容。")

    def _refresh_rules(self) -> None:
        self.mapping_tree.delete(*self.mapping_tree.get_children())
        for rule in self.rules:
            source = (
                f"{Path(rule.source_file).name} / {rule.source_sheet} / "
                f"{format_cell_addresses(rule.source_cells)}"
            )
            target = (
                f"{Path(rule.target_file).name} / {rule.target_sheet} / "
                f"{format_cell_addresses(rule.target_cells)}"
            )
            self.mapping_tree.insert(
                "",
                END,
                values=(source, target),
            )

    def _preview_expanded(self) -> None:
        try:
            expanded = expand_rules(self.rules)
        except ValueError as exc:
            messagebox.showerror("无法展开映射", str(exc), parent=self.root)
            return
        if not expanded:
            messagebox.showinfo(
                "没有映射",
                "请先添加映射规则。",
                parent=self.root,
            )
            return

        window = Toplevel(self.root)
        window.title(f"映射展开预览（共 {len(expanded)} 项）")
        window.geometry("900x520")
        window.transient(self.root)
        container = ttk.Frame(window, padding=12)
        container.pack(fill=BOTH, expand=True)
        tree = ttk.Treeview(
            container,
            columns=("number", "source", "target"),
            show="headings",
        )
        tree.heading("number", text="序号")
        tree.heading("source", text="来源位置")
        tree.heading("target", text="目标位置")
        tree.column("number", width=60, anchor="center")
        tree.column("source", width=390)
        tree.column("target", width=390)
        scroll = ttk.Scrollbar(
            container,
            orient="vertical",
            command=tree.yview,
        )
        tree.configure(yscrollcommand=scroll.set)
        tree.pack(side=LEFT, fill=BOTH, expand=True)
        scroll.pack(side=RIGHT, fill="y")
        for index, mapping in enumerate(expanded, start=1):
            source = (
                f"{Path(mapping.source_file).name} / "
                f"{mapping.source_sheet} / {mapping.source_cell}"
            )
            target = (
                f"{Path(mapping.target_file).name} / "
                f"{mapping.target_sheet} / {mapping.target_cell}"
            )
            tree.insert("", END, values=(index, source, target))

    def _browse_output(self) -> None:
        selected = filedialog.askdirectory(title="选择副本输出文件夹")
        if selected:
            self.output_folder.set(selected)

    def _copy_txt_example(self) -> None:
        self.root.clipboard_clear()
        self.root.clipboard_append(TXT_EXAMPLE)
        self.status.set("TXT 格式示例已复制到剪贴板。")

    def _browse_txt_base(self) -> None:
        selected = filedialog.askdirectory(title="选择方案基准文件夹")
        if not selected:
            return
        self.txt_base_folder.set(selected)
        if not self.output_folder.get().strip():
            self.output_folder.set(str(Path(selected) / "映射结果"))
        self._txt_dirty = True
        self.txt_status.set("基准文件夹已更改，切换页签时将自动检查。")

    def _txt_editor_modified(self, _event=None) -> None:
        if self._updating_txt_editor:
            self.txt_editor.edit_modified(False)
            return
        if self.txt_editor.edit_modified():
            self._txt_dirty = True
            self.txt_status.set("TXT 内容已修改，切换页签时将自动检查。")
            self.txt_editor.edit_modified(False)

    def _set_txt_text(self, text: str) -> None:
        self._updating_txt_editor = True
        try:
            self.txt_editor.delete("1.0", END)
            self.txt_editor.insert("1.0", text)
            self.txt_editor.edit_modified(False)
        finally:
            self._updating_txt_editor = False

    def _default_txt_base_folder(self) -> Path | None:
        configured = self.txt_base_folder.get().strip()
        if configured:
            return Path(configured).resolve()
        files = self.source_files + self.target_files
        if files:
            return files[0].parent.resolve()
        return None

    def _sync_manual_to_txt(self) -> None:
        if self._txt_dirty:
            self.txt_status.set(
                "已保留尚未同步的 TXT 内容；执行时再进行完整检查。"
            )
            return
        base_folder = self._default_txt_base_folder()
        if base_folder is None:
            self._set_txt_text("")
            self.txt_status.set("请先在“手动设置”中添加工作簿和映射关系。")
            self._txt_dirty = False
            return
        self.txt_base_folder.set(str(base_folder))
        self._set_txt_text(serialize_mapping_project(self.rules, base_folder))
        self._txt_dirty = False
        if self.rules:
            self.txt_status.set(
                f"已从手动设置自动生成，共 {len(self.rules)} 条映射。"
            )
        else:
            self.txt_status.set("尚无映射关系，TXT 内容为空。")

    def _apply_txt_to_manual(
        self,
        show_message: bool = True,
        check_workbooks: bool = True,
    ) -> bool:
        plan_text = self.txt_editor.get("1.0", END).strip()
        if not plan_text:
            self._txt_dirty = False
            self.txt_status.set("TXT 内容为空，未更改手动设置。")
            return True
        base_text = self.txt_base_folder.get().strip()
        if not base_text:
            message = "请先选择方案基准文件夹。"
            self.txt_status.set(message)
            if show_message:
                messagebox.showerror("TXT 方案无法应用", message, parent=self.root)
            return False
        try:
            sources, targets, rules = parse_mapping_project_text(
                plan_text,
                Path(base_text).resolve(),
            )
            warnings: list[str] = []
            if check_workbooks:
                _expanded, warnings = validate_mapping_plan(
                    sources,
                    targets,
                    rules,
                )
        except Exception as exc:
            message = str(exc)
            self.txt_status.set(f"检查失败：{message}")
            if show_message:
                messagebox.showerror(
                    "TXT 方案无法应用",
                    message,
                    parent=self.root,
                )
            return False

        self.source_files = sources
        self.target_files = targets
        self.rules = rules
        self.output_folder.set(str(Path(base_text).resolve() / "映射结果"))
        self._refresh_file_tree(self.source_tree, self.source_files)
        self._refresh_file_tree(self.target_tree, self.target_files)
        self._refresh_rules()
        self._txt_dirty = False
        suffix = f"，发现 {len(warnings)} 个目标单元格已有内容" if warnings else ""
        self.txt_status.set(f"TXT 方案已同步，共 {len(rules)} 条映射{suffix}。")
        return True

    def _workflow_tab_changed(self, _event=None) -> None:
        if (
            self._switching_workflow_tab
            or not hasattr(self, "txt_editor")
        ):
            return
        selected = self.workflow_notebook.index(
            self.workflow_notebook.select()
        )
        if selected == self._active_workflow_tab:
            return
        if selected == 1:
            self._sync_manual_to_txt()
            self._active_workflow_tab = 1
            return
        self._apply_txt_to_manual(
            show_message=False,
            check_workbooks=False,
        )
        self._active_workflow_tab = 0

    def _clear_txt_content(self) -> None:
        if self.txt_editor.get("1.0", END).strip() and not messagebox.askyesno(
            "清空 TXT 内容",
            "确定清空当前 TXT 方案内容吗？",
            parent=self.root,
        ):
            return
        self._set_txt_text("")
        self._txt_dirty = True
        self.txt_status.set("TXT 内容已清空，尚未同步到手动设置。")

    def _show_txt_example(self) -> None:
        window = Toplevel(self.root)
        window.title("TXT 映射方案格式示例")
        window.geometry("820x540")
        window.minsize(680, 430)
        window.transient(self.root)
        container = ttk.Frame(window, padding=12)
        container.pack(fill=BOTH, expand=True)
        ttk.Label(container, text="示例文件内容：").pack(anchor=W)
        example = Text(
            container,
            wrap="none",
            font=("TkFixedFont", 10),
            height=7,
        )
        example.insert("1.0", TXT_EXAMPLE)
        example.configure(state="disabled")
        example.pack(fill=BOTH, expand=True, pady=(3, 10))
        ttk.Label(container, text="填写说明：").pack(anchor=W)
        instructions = Text(
            container,
            wrap="word",
            font=("TkDefaultFont", 10),
            height=10,
        )
        instructions.insert("1.0", TXT_INSTRUCTIONS)
        instructions.configure(state="disabled")
        instructions.pack(fill=BOTH, expand=True, pady=(3, 0))
        actions = ttk.Frame(container)
        actions.pack(fill=X, pady=(8, 0))

        ttk.Button(
            actions,
            text="复制示例",
            command=self._copy_txt_example,
        ).pack(side=LEFT)
        ttk.Button(
            actions,
            text="关闭",
            command=window.destroy,
        ).pack(side=LEFT, padx=(8, 0))

    def _save_project(self) -> None:
        if (
            self.workflow_notebook.index(self.workflow_notebook.select()) == 1
            and not self._apply_txt_to_manual()
        ):
            return
        if not self.rules:
            messagebox.showerror(
                "没有映射关系",
                "请先添加或导入至少一条映射关系。",
                parent=self.root,
            )
            return
        selected = filedialog.asksaveasfilename(
            title="保存 TXT 映射方案",
            defaultextension=".txt",
            filetypes=[("TXT 映射方案", "*.txt"), ("所有文件", "*.*")],
        )
        if selected:
            path = Path(selected)
            save_mapping_project(path, self.rules)
            self.txt_base_folder.set(str(path.parent.resolve()))
            self._set_txt_text(
                serialize_mapping_project(self.rules, path.parent.resolve())
            )
            self._txt_dirty = False
            self.txt_status.set(f"TXT 方案已保存：{path.name}")

    def _load_project(self) -> None:
        selected = filedialog.askopenfilename(
            title="导入 TXT 映射方案",
            filetypes=[("TXT 映射方案", "*.txt"), ("所有文件", "*.*")],
        )
        if not selected:
            return
        try:
            path = Path(selected)
            self.txt_base_folder.set(str(path.parent.resolve()))
            self._set_txt_text(path.read_text(encoding="utf-8-sig"))
            self._txt_dirty = True
            if self._apply_txt_to_manual():
                self.txt_status.set(
                    f"已导入并同步 {path.name}，共 {len(self.rules)} 条映射。"
                )
        except Exception as exc:
            self.txt_status.set(f"导入失败：{exc}")
            messagebox.showerror("导入失败", str(exc), parent=self.root)

    @staticmethod
    def _refresh_file_tree(tree, files: list[Path]) -> None:
        tree.delete(*tree.get_children())
        for path in files:
            tree.insert("", END, values=(str(path),))

    def _precheck(self) -> bool:
        try:
            expanded, warnings = validate_mapping_plan(
                self.source_files,
                self.target_files,
                self.rules,
            )
            message = (
                f"来源工作簿：{len(self.source_files)} 个\n"
                f"目标工作簿：{len(self.target_files)} 个\n"
                f"映射规则：{len(self.rules)} 条\n"
                f"展开后写入：{len(expanded)} 项"
            )
            if warnings:
                preview = "\n".join(warnings[:8])
                message += (
                    f"\n\n有 {len(warnings)} 个目标单元格已有内容，"
                    "将在副本中覆盖：\n"
                    f"{preview}"
                )
                if len(warnings) > 8:
                    message += "\n..."
            messagebox.showinfo("预检查通过", message, parent=self.root)
            return True
        except Exception as exc:
            messagebox.showerror("预检查失败", str(exc), parent=self.root)
            return False

    def _progress(self, current: int, total: int, message: str) -> None:
        maximum = max(total, 1)
        self.progress.configure(maximum=maximum, value=current)
        self.progress_text.set(f"{int(current / maximum * 100)}%")
        self.status.set(message)
        self.root.update_idletasks()

    def _run(self) -> None:
        if (
            self.workflow_notebook.index(self.workflow_notebook.select()) == 1
            and not self._apply_txt_to_manual()
        ):
            return
        if not self.output_folder.get().strip():
            messagebox.showerror(
                "缺少输出位置",
                "请选择副本输出文件夹。",
                parent=self.root,
            )
            return
        try:
            expanded, warnings = validate_mapping_plan(
                self.source_files,
                self.target_files,
                self.rules,
            )
            if warnings and not messagebox.askyesno(
                "确认覆盖副本中的内容",
                f"有 {len(warnings)} 个目标单元格已有内容，"
                "将在新副本中覆盖。\n\n原文件不会修改。是否继续？",
                parent=self.root,
            ):
                return
            self.progress.configure(maximum=max(len(expanded), 1), value=0)
            self.progress_text.set("0%")
            outputs = execute_mapping_plan(
                self.source_files,
                self.target_files,
                self.rules,
                Path(self.output_folder.get()),
                self._progress,
            )
            self.progress.configure(value=self.progress.cget("maximum"))
            self.progress_text.set("100%")
            self.status.set("映射完成；原文件未修改。")
            messagebox.showinfo(
                "映射完成",
                f"已生成 {len(outputs)} 个工作簿副本：\n"
                + "\n".join(str(path) for path in outputs),
                parent=self.root,
            )
        except Exception as exc:
            self.status.set("映射失败。")
            messagebox.showerror("映射失败", str(exc), parent=self.root)


def main() -> None:
    root = Tk()
    MapperApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
