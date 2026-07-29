from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

import export_summary


def make_source(path: Path, marker: int) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "数据1"
    second = workbook.create_sheet("数据2")
    for index, sheet in enumerate((first, second), start=1):
        sheet["I8"] = marker + index + 0.25
        sheet["I8"].number_format = "0"
        sheet["J8"] = marker + index + 0.75
        sheet["J8"].number_format = "0"
    workbook.save(path)


class ExportSummaryTests(unittest.TestCase):
    def test_parse_cell_addresses(self) -> None:
        self.assertEqual(
            export_summary.parse_cell_addresses("i8， J8\nL8;I8"),
            ["I8", "J8", "L8"],
        )
        with self.assertRaises(ValueError):
            export_summary.parse_cell_addresses("I0, BAD")

    def test_cell_grid_address_helpers(self) -> None:
        self.assertEqual(export_summary.column_letters_to_number("A"), 1)
        self.assertEqual(export_summary.column_letters_to_number("AZ"), 52)
        self.assertEqual(export_summary.column_number_to_letters(52), "AZ")
        self.assertEqual(
            sorted(["L8", "I17", "J8", "I8"], key=export_summary.cell_sort_key),
            ["I8", "J8", "L8", "I17"],
        )

    def test_create_combined_summary(self) -> None:
        with TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            first = folder / "甲.xlsx"
            second = folder / "乙.xlsx"
            output = folder / "批量汇总表.xlsx"
            make_source(first, 100)
            make_source(second, 200)

            export_summary.create_new_summary_workbook(
                [first, second],
                ["I8", "J8"],
                output,
                "汇总",
            )

            workbook = load_workbook(output, data_only=False)
            sheet = workbook["汇总"]
            self.assertEqual(sheet.max_row, 5)
            self.assertEqual(sheet["A1"].value, "文件名称")
            self.assertEqual(sheet["B1"].value, "Sheet名称")
            self.assertEqual(sheet["C2"].value, 101.25)
            self.assertEqual(sheet["C2"].number_format, "0")
            workbook.close()

    def test_insert_summary_as_first_sheet(self) -> None:
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "源.xlsx"
            make_source(source, 300)

            export_summary.insert_summary_into_workbook(
                source,
                ["I8", "J8"],
                "汇总",
            )

            workbook = load_workbook(source, data_only=False)
            self.assertEqual(workbook.sheetnames[0], "汇总")
            self.assertEqual(workbook["汇总"]["B2"].value, 301.25)
            self.assertEqual(workbook["汇总"]["B2"].number_format, "0")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
