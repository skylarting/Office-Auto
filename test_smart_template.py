import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from excel_mapper import WorkbookReader
from smart_template import (
    BlockRule,
    FormulaRule,
    SmartTemplatePlan,
    detect_sheet_structure,
    execute_smart_plan,
    match_is_blocked,
    match_pair,
    matches_to_mapping_rules,
    normalize_label,
    resolve_formula_targets,
    SheetPair,
    SmartMatch,
)


def make_book(path: Path, sheet_name: str, shifted: bool = False) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    row_offset = 2 if shifted else 0
    col_offset = 1 if shifted else 0
    sheet.cell(1 + row_offset, 2 + col_offset, "余额")
    sheet.merge_cells(
        start_row=1 + row_offset,
        start_column=2 + col_offset,
        end_row=1 + row_offset,
        end_column=3 + col_offset,
    )
    sheet.cell(2 + row_offset, 2 + col_offset, "人民币")
    sheet.cell(2 + row_offset, 3 + col_offset, "外币")
    sheet.cell(3 + row_offset, 1 + col_offset, "Ⅰ. 资产")
    sheet.cell(4 + row_offset, 1 + col_offset, "1. 现金")
    sheet.cell(4 + row_offset, 2 + col_offset, 10)
    sheet.cell(4 + row_offset, 3 + col_offset, 20)
    sheet.cell(5 + row_offset, 1 + col_offset, "2. 存款")
    sheet.cell(5 + row_offset, 2 + col_offset, 30)
    sheet.cell(5 + row_offset, 3 + col_offset, "=B4+C4")
    book.save(path)


class SmartTemplateTests(unittest.TestCase):
    def test_normalize_label_ignores_presentation_differences(self) -> None:
        self.assertEqual(normalize_label("1. 现 金（人民币）"), normalize_label("1．现金 / 人民币"))

    def test_detect_and_match_shifted_multilevel_headers(self) -> None:
        with TemporaryDirectory() as folder:
            source = Path(folder) / "source.xlsx"
            target = Path(folder) / "target.xlsx"
            make_book(source, "来源")
            make_book(target, "目标", shifted=True)
            pair = SheetPair(str(source), "来源", str(target), "目标", 1.0)
            matches = match_pair(pair)
            addresses = {(item.source_address, item.target_address) for item in matches}
            self.assertIn(("B4", "C6"), addresses)
            self.assertIn(("C4", "D6"), addresses)
            self.assertNotIn(("C5", "D7"), addresses)  # source formula is excluded

    def test_block_rule_has_priority_and_formula_is_preserved(self) -> None:
        with TemporaryDirectory() as folder:
            source = Path(folder) / "source.xlsx"
            target = Path(folder) / "target.xlsx"
            make_book(source, "来源")
            make_book(target, "目标")
            match = SmartMatch(
                str(source), "来源", "B4", ("资产", "现金"), ("余额", "人民币"),
                str(target), "目标", "B4", ("资产", "现金"), ("余额", "人民币"),
                1.0, "自动匹配", True,
            )
            plan = SmartTemplatePlan(
                str(target), [str(source)],
                matches=[match],
                block_rules=[BlockRule("目标", "现金", "人民币", "交叉位置", "测试禁填")],
            )
            rules, skipped = matches_to_mapping_rules(plan)
            self.assertEqual(rules, [])
            self.assertIn("测试禁填", skipped[0])

    def test_value_mapping_can_explicitly_overwrite_target_formula(self) -> None:
        with TemporaryDirectory() as folder:
            source = Path(folder) / "source.xlsx"
            target = Path(folder) / "target.xlsx"
            make_book(source, "来源")
            make_book(target, "目标")
            match = SmartMatch(
                str(source), "来源", "B4", ("资产", "现金"), ("余额", "人民币"),
                str(target), "目标", "C5", ("资产", "存款"), ("余额", "外币"),
                1.0, "人工修改", True,
            )
            plan = SmartTemplatePlan(
                str(target), [str(source)], matches=[match],
                overwrite_target_formulas=True,
            )
            rules, skipped = matches_to_mapping_rules(plan)
            self.assertEqual(len(rules), 1)
            self.assertEqual(skipped, [])

    def test_formula_rule_uses_row_placeholder_and_policy(self) -> None:
        with TemporaryDirectory() as folder:
            target = Path(folder) / "target.xlsx"
            make_book(target, "目标")
            pair = SheetPair(str(target), "目标", str(target), "目标", 1.0)
            plan = SmartTemplatePlan(
                str(target), [], [pair],
                formula_rules=[FormulaRule("目标", "现金", "人民币", "=C{row}+1", "强制覆盖")],
            )
            formulas = resolve_formula_targets(plan)
            self.assertTrue(any(address == "B4" and formula == "=C4+1" for _, _, address, formula, _ in formulas))

    def test_execute_writes_values_to_copy_without_touching_original(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source.xlsx"
            target = root / "target.xlsx"
            output = root / "out"
            make_book(source, "来源")
            make_book(target, "目标", shifted=True)
            pair = SheetPair(str(source), "来源", str(target), "目标", 1.0)
            plan = SmartTemplatePlan(str(target), [str(source)], [pair])
            plan.matches = match_pair(pair)
            outputs, _ = execute_smart_plan(plan, output)
            written = load_workbook(outputs[0], data_only=False)
            original = load_workbook(target, data_only=False)
            try:
                self.assertEqual(written["目标"]["C6"].value, 10)
                self.assertEqual(original["目标"]["C6"].value, 10)
                self.assertEqual(written["目标"]["D7"].value, "=B4+C4")
            finally:
                written.close(); original.close()


if __name__ == "__main__":
    unittest.main()
